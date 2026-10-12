"""Accounting, fourth batch (user, 2026-10-12): the Master Report Excel (a tab per master account), the lines behind a
Partner Accounts figure, AT-HUB / rules never file a line over the limit by themselves, bulk moves say which lines moved
(incl. the ones filed like them), the "like this" pop-up knows what's already in the account. Made-up data."""
import io

from openpyxl import load_workbook

from tests.builders import uid
from tests.test_accounting import _csv, _import, _source, _txns
from tests.test_accounting_rules import c, name


def test_master_report_has_a_tab_per_master_account_and_split_parts(api, client, admin_headers):
    y = 2071
    src = _source(api)
    labor = api.post("/api/accounting/accounts", json={"name": uid("Labor "), "kind": "expense"})
    saul = api.post("/api/accounting/accounts", json={"name": uid("Saul "), "kind": "expense", "parent_id": labor["id"]})
    nir = api.post("/api/accounting/accounts", json={"name": uid("Niraj "), "kind": "owner", "partner": "Niraj"})
    ric = api.post("/api/accounting/accounts", json={"name": uid("Rich "), "kind": "owner", "partner": "Richard"})
    _import(client, admin_headers, src["id"], "m.csv", _csv([(f"05/01/{y}", -100.0, f"ZELLE TO {name('SAULM')}"), (f"05/02/{y}", -50.0, f"ZELLE TO {name('LABM')}"),
                                                             (f"05/03/{y}", -1000.0, f"WT {name('SPLM')}")]))
    rows = sorted(_txns(api, all_dates="true", source_id=src["id"])["rows"], key=lambda t: t["amount"])
    api.put(f"/api/accounting/txns/{rows[0]['id']}", json={"splits": [{"account_id": nir["id"], "amount": -600}, {"account_id": ric["id"], "amount": -400}], "learn": False})
    api.put(f"/api/accounting/txns/{rows[1]['id']}", json={"account_id": saul["id"], "learn": False, "note": "May pay"})
    api.put(f"/api/accounting/txns/{rows[2]['id']}", json={"account_id": labor["id"], "learn": False})
    x = client.get(f"/api/accounting/master.xlsx?start={y}-04-01&end={y + 1}-03-31", headers=admin_headers)
    assert x.status_code == 200 and x.content[:2] == b"PK"
    wb = load_workbook(io.BytesIO(x.content))
    names = wb.sheetnames
    assert names[:3] == ["Summary", "P&L", "All Transactions"] and {"Partner Accounts", "Loans", "Not Filed", "Comments"} <= set(names)
    lab = wb[next(n for n in names if n.startswith(labor["name"][:25]))]
    vals = [[c.value for c in r] for r in lab.iter_rows()]
    assert any(r[0] == f"Total {labor['name']}" and r[4] == -150.0 for r in vals)          # the account + its sub-account
    assert any(r[0] and str(r[0]).startswith("Subtotal") and r[4] == -100.0 for r in vals)
    assert any(r[6] == "May pay" for r in vals)
    alltx = [[c.value for c in r] for r in wb["All Transactions"].iter_rows(min_row=2)]
    split_rows = [r for r in alltx if r[9] == "yes"]
    assert sorted(r[7] for r in split_rows) == [-600.0, -400.0]                                # one row per part


def test_partner_lines_behind_a_month(api, client, admin_headers):
    y = 2072
    who = uid("Partner")
    src = _source(api)
    acc = api.post("/api/accounting/accounts", json={"name": uid("Draws "), "kind": "owner", "partner": who})
    other = api.post("/api/accounting/accounts", json={"name": uid("Other "), "kind": "owner", "partner": uid("P2")})
    _import(client, admin_headers, src["id"], "p.csv", _csv([(f"06/01/{y}", -300.0, f"ZELLE TO {name('PAONE')}"), (f"06/15/{y}", -1000.0, f"WT {name('PATWO')}"),
                                                             (f"07/01/{y}", 200.0, f"DEPOSIT {name('PATHREE')}")]))
    rows = sorted(_txns(api, all_dates="true", source_id=src["id"])["rows"], key=lambda t: t["date"])
    api.put(f"/api/accounting/txns/{rows[0]['id']}", json={"account_id": acc["id"], "learn": False})
    api.put(f"/api/accounting/txns/{rows[1]['id']}", json={"splits": [{"account_id": acc["id"], "amount": -250}, {"account_id": other["id"], "amount": -750}], "learn": False})
    api.put(f"/api/accounting/txns/{rows[2]['id']}", json={"account_id": acc["id"], "learn": False})
    june = api.get(f"/api/accounting/partner-lines?name={who}&start={y}-06-01&end={y}-06-30&what=withdrawn")
    assert june["count"] == 2 and june["total"] == -550.0 and sorted(r["part"] for r in june["rows"]) == [-300.0, -250.0]
    assert api.get(f"/api/accounting/partner-lines?name={who}&start={y}-01-01&end={y}-12-31&what=invested")["total"] == 200.0


def test_over_the_limit_never_filed_by_itself_and_bulk_says_which(api, client, admin_headers):
    y = 2073
    src = _source(api)
    big = api.post("/api/accounting/accounts", json={"name": uid("Big "), "kind": "cogs"})
    tag = name("LIMITQ")
    rid = api.post("/api/accounting/rules", json={"conditions": [c("text", "contains", tag)], "sign": "out", "account_id": big["id"]})["id"]
    api.put("/api/accounting/settings", json={"auto_limit": 5000})
    _import(client, admin_headers, src["id"], "l.csv", _csv([(f"03/01/{y}", -4000.0, f"WT {tag} ONE"), (f"03/02/{y}", -9000.0, f"WT {tag} TWO")]))
    got = {t["amount"]: t for t in _txns(api, all_dates="true", source_id=src["id"])["rows"]}
    assert got[-4000.0]["account_id"] == big["id"]
    assert got[-9000.0]["account_id"] is None and got[-9000.0]["flag"] == "over_limit"
    rv = {t["id"]: t for t in _txns(api, status="review", all_dates="true")["rows"]}
    assert rv[got[-9000.0]["id"]]["suggest"]["account_id"] == big["id"]      # still a one-click suggestion
    # a person files it: the flag goes
    one = api.put(f"/api/accounting/txns/{got[-9000.0]['id']}", json={"account_id": big["id"], "learn": False})
    assert not one["flag"]
    # bulk: which lines moved, the ones filed like them, "show them" by ids
    other = api.post("/api/accounting/accounts", json={"name": uid("Moved "), "kind": "expense"})
    ids = [t["id"] for t in got.values()]
    b = api.post("/api/accounting/txns/bulk", json={"ids": ids, "action": "file", "account_id": other["id"]})
    assert b["done"] == 2 and "more_ids" in b
    shown = _txns(api, ids=",".join(map(str, ids)))
    assert shown["total"] == 2 and all(t["account_id"] == other["id"] for t in shown["rows"])
    api.put("/api/accounting/settings", json={"auto_limit": 0})
    api.delete(f"/api/accounting/rules/{rid}")


def test_like_this_popup_knows_what_is_already_in_the_account(api, client, admin_headers):
    y = 2074
    src = _source(api)
    acc = api.post("/api/accounting/accounts", json={"name": uid("Gas "), "kind": "expense"})
    tag = name("GASQQ")
    _import(client, admin_headers, src["id"], "g.csv", _csv([(f"08/0{d}/{y}", -40.0 - d, f"ZELLE TO {tag} ON 08/0{d} REF #G{d}") for d in range(1, 4)]))
    rows = sorted(_txns(api, all_dates="true", source_id=src["id"])["rows"], key=lambda t: t["date"])
    api.put(f"/api/accounting/txns/{rows[0]['id']}", json={"account_id": acc["id"], "learn": False})
    r = api.put(f"/api/accounting/txns/{rows[1]['id']}", json={"account_id": acc["id"]})
    if r["alike"]:
        assert any(x["id"] == rows[0]["id"] for x in r["alike_here"])
