"""Accounting, fifth batch (user, 2026-10-12): rules with branches, "what it will do" story, combine ticked rules;
Loans first round (timeline events, moves between accounts with book entries, interest, schedule, the lender's
statement); ask a partner to check a line; Reconcile with the full QuickBooks register, add several lines at once,
delete / restore earlier checks, the Fix In QuickBooks list; the Master Report in 10 tabs. Made-up data."""
import io

from openpyxl import Workbook, load_workbook

from tests.builders import uid
from tests.test_accounting import _csv, _import, _source, _txns
from tests.test_accounting_rules import c, name
from tests.test_roles import _user


def test_branch_rule_story_and_combine_ticked(api, client, admin_headers):
    api.put("/api/accounting/settings", json={"auto_limit": 0})
    y = 2080
    src = _source(api)
    jason = api.post("/api/accounting/accounts", json={"name": uid("Jason "), "kind": "expense"})
    sup = api.post("/api/accounting/accounts", json={"name": uid("Supplies "), "kind": "expense"})
    j, p_, s_ = name("JASONB"), name("PESOB"), name("SUPPB")
    _import(client, admin_headers, src["id"], "b.csv", _csv([(f"01/0{d}/{y}", -10.0 * d, desc) for d, desc in
                                                             ((1, f"ZELLE TO {j}"), (2, f"ZELLE {p_} LUNCH"), (3, f"ZELLE TO {j} {s_}"), (4, "ZELLE OTHER"))]))
    rule = {"name": "Jason / Supplies", "sign": "out", "conditions": [c("text", "any_of", f"{j}, {p_}")],
            "branches": [{"conditions": [c("text", "contains", s_)], "account_id": sup["id"]}, {"conditions": [], "account_id": jason["id"]}]}
    story = api.post("/api/accounting/rules/story", json=rule)
    assert story["total"] == 3 and "otherwise" in story["sentence"] and story["examples"]
    rid = api.post("/api/accounting/rules", json=rule)["id"]
    lines = api.get(f"/api/accounting/rules/{rid}/lines")
    api.post(f"/api/accounting/rules/{rid}/apply", json={"ids": [t["id"] for t in lines["groups"]["unfiled"]]})
    got = {round(t["amount"]): t for t in _txns(api, all_dates="true", source_id=src["id"])["rows"]}
    assert got[-10]["account_id"] == jason["id"] and got[-20]["account_id"] == jason["id"] and got[-30]["account_id"] == sup["id"]
    assert got[-40]["account_id"] is None
    why = api.get(f"/api/accounting/txns/{got[-30]['id']}/rules")
    assert next(x for x in why["rules"] if x["id"] == rid)["branch"] == 0
    # combine ticked: two single rules -> one draft; saving replaces them; Undo brings them back
    r1 = api.post("/api/accounting/rules", json={"conditions": [c("text", "contains", j)], "sign": "out", "account_id": jason["id"]})["id"]
    r2 = api.post("/api/accounting/rules", json={"conditions": [c("text", "contains", p_)], "sign": "out", "account_id": jason["id"]})["id"]
    dr = api.post("/api/accounting/rule-combine-draft", json={"rule_ids": [r1, r2]})
    assert dr["same_account"] and dr["lines_before"] == dr["lines_after"]
    saved = api.post("/api/accounting/rules", json={**dr["draft"], "name": "Jason all", "replaces": dr["replaces"]})
    ids = {x["id"] for x in api.get("/api/accounting/rules")}
    assert saved["id"] in ids and r1 not in ids and r2 not in ids
    api.post(f"/api/accounting/changes/{saved['change_id']}/undo")
    ids = {x["id"] for x in api.get("/api/accounting/rules")}
    assert r1 in ids and r2 in ids and saved["id"] not in ids
    # different accounts -> a branch draft
    r3 = api.post("/api/accounting/rules", json={"conditions": [c("text", "contains", s_)], "sign": "out", "account_id": sup["id"]})["id"]
    dr2 = api.post("/api/accounting/rule-combine-draft", json={"rule_ids": [r1, r3]})
    assert not dr2["same_account"] and len(dr2["draft"]["branches"]) == 2
    for x in (rid, r1, r2, r3):
        api.delete(f"/api/accounting/rules/{x}")


def test_loans_timeline_moves_interest_schedule_and_lender_statement(api, client, admin_headers):
    y = 2025  # the schedule only lists months up to today
    src = _source(api)
    l1 = api.post("/api/accounting/accounts", json={"name": uid("Sandy "), "kind": "loan"})
    l2 = api.post("/api/accounting/accounts", json={"name": uid("MT Bank "), "kind": "loan"})
    _import(client, admin_headers, src["id"], "l.csv", _csv([(f"01/05/{y}", 10000.0, f"LOAN IN {name('LOANA')}"), (f"02/05/{y}", -1000.0, f"LOAN PAY {name('LOANB')}")]))
    for t in _txns(api, all_dates="true", source_id=src["id"])["rows"]:
        api.put(f"/api/accounting/txns/{t['id']}", json={"account_id": l1["id"], "learn": False})
    api.put(f"/api/accounting/loans/{l1['id']}/details", json={"lender": "Sandy Spring", "rate": 12, "start": f"{y}-01-05", "payment": 1000, "due_day": 5})
    st = api.get(f"/api/accounting/loans/{l1['id']}")
    assert st["owed"] == 9000.0 and st["details"]["lender"] == "Sandy Spring" and st["schedule"]
    api.post(f"/api/accounting/loans/{l1['id']}/events", json={"date": f"{y}-02-10", "kind": "call", "text": "asked for a lower rate"})
    mv = api.post(f"/api/accounting/loans/{l1['id']}/events", json={"date": f"{y}-03-01", "kind": "moved_out", "amount": 9000, "other_account_id": l2["id"], "text": "refinanced"})
    st1, st2 = api.get(f"/api/accounting/loans/{l1['id']}"), api.get(f"/api/accounting/loans/{l2['id']}")
    assert st1["owed"] == 0.0 and not st1["open"] and any("closed" in h["what"] for h in st1["history"])
    assert st2["owed"] == 9000.0 and any(r["type"] == "event" and r["kind"] == "moved_in" for r in st2["rows"])
    assert api.get(f"/api/accounting/loans/{l1['id']}/interest?to={y}-02-05")["interest"] > 0
    # Undo the move: both loans back
    api.post(f"/api/accounting/changes/{mv['change_id']}/undo")
    assert api.get(f"/api/accounting/loans/{l1['id']}")["owed"] == 9000.0 and api.get(f"/api/accounting/loans/{l2['id']}")["owed"] == 0.0
    # the lender's statement
    wb = Workbook()
    ws = wb.active
    ws.append(["Date", "Description", "Amount"])
    ws.append([f"01/06/{y}", "Advance", 10000])
    ws.append([f"02/05/{y}", "Payment", -1100])
    ws.append([f"02/20/{y}", "Fee", -25])
    buf = io.BytesIO()
    wb.save(buf)
    r = client.post(f"/api/accounting/loans/{l1['id']}/verify", headers=admin_headers, files={"file": ("lender.xlsx", io.BytesIO(buf.getvalue()), "application/octet-stream")})
    assert r.status_code == 200, r.text
    v = r.json()
    assert v["counts"]["matched"] == 1 and v["counts"]["amount_differs"] == 1 and v["counts"]["only_theirs"] == 1
    assert any(x["account_id"] == l1["id"] for x in api.get("/api/accounting/loans-overview")["loans"])
    x = client.get(f"/api/accounting/loans/{l1['id']}/export.xlsx", headers=admin_headers)
    assert x.status_code == 200 and x.content[:2] == b"PK"


def test_ask_a_partner_to_check_a_line(api, client, admin_headers):
    src = _source(api)
    _import(client, admin_headers, src["id"], "a.csv", _csv([("03/03/2082", -500.0, f"ZELLE TO {name('ASKQ')}")]))
    t = _txns(api, all_dates="true", source_id=src["id"])["rows"][0]
    other = _user(client, admin_headers, "admin")  # an admin has Accounting
    users = api.get("/api/accounting/ask/users")
    assert any(u["username"] == "u-admin" for u in users)
    assert api.post("/api/accounting/txns/ask", json={"ids": [t["id"]], "to_user": "u-admin", "text": "did you get this?"})["asked"] == 1
    waiting = client.get("/api/accounting/asks/waiting", headers=other).json()
    w = next(x for x in waiting if x["txn_id"] == t["id"])
    assert w["text"] == "did you get this?"
    todo = client.get("/api/todo/items?view=all", headers=other)
    assert todo.status_code == 200 and "please check" in todo.text
    # the asker can't answer for them; they can
    assert client.post(f"/api/accounting/asks/{w['id']}/answer", headers=admin_headers, json={"status": "yes"}).status_code == 403
    a = client.post(f"/api/accounting/asks/{w['id']}/answer", headers=other, json={"status": "yes", "reply": "got it 3/4"}).json()
    assert a["status"] == "yes"
    line = next(x for x in _txns(api, all_dates="true", source_id=src["id"])["rows"] if x["id"] == t["id"])
    assert line["asks"][0]["status"] == "yes" and line["asks"][0]["reply"] == "got it 3/4"


def test_reconcile_register_add_delete_and_fix_list(api, client, admin_headers):
    y = 2083
    src = _source(api)
    nir = api.post("/api/accounting/accounts", json={"name": uid("NirajX "), "kind": "owner"})
    cogs = api.post("/api/accounting/accounts", json={"name": uid("SuppliersX "), "kind": "cogs"})
    _import(client, admin_headers, src["id"], "r.csv", _csv([(f"08/04/{y}", -6200.0, f"WT FED BNF=Progea {name('PROGQ')} TRN#250804175787")]))
    t = _txns(api, all_dates="true", source_id=src["id"])["rows"][0]
    api.put(f"/api/accounting/txns/{t['id']}", json={"account_id": nir["id"], "learn": False})
    qb = [{"date": f"{y}-08-10", "amount": -300.0, "name": "", "memo": "ZELLE TO SOMEONE", "account": "Partner distributions:Niraj Distributions"}]
    reg = [{"date": f"{y}-08-04", "amount": -6200.0, "name": "Progea", "memo": "Site inspection phase 1", "account": "Cost of goods sold:Suppliers"}]
    body = {"account_ids": [nir["id"]], "start": f"{y}-08-01", "end": f"{y}-08-31", "rows": qb, "register_rows": reg,
            "mapping": {"Partner distributions:Niraj Distributions": nir["id"], "Cost of goods sold:Suppliers": cogs["id"]}, "save": True}
    res = api.post("/api/accounting/reconcile/run", json=body)
    ao = next(x for x in res["athub_only"] if x["id"] == t["id"])
    assert ao["qb_register"]["account"] == "Cost of goods sold:Suppliers" and ao["qb_register"]["mapped_account_id"] == cogs["id"]
    fix = api.post("/api/accounting/reconcile/fix", json=body)["actions"]
    kinds = {a["kind"] for a in fix}
    assert {"move", "delete"} <= kinds
    assert any("Cost of goods sold:Suppliers" in a["what"] for a in fix if a["kind"] == "move")
    x = client.post("/api/accounting/reconcile/fix.xlsx", headers=admin_headers, json=body)
    assert x.status_code == 200 and "To Do" in load_workbook(io.BytesIO(x.content)).sheetnames
    # add the only-in-QuickBooks line, in one go
    add = api.post("/api/accounting/reconcile/add", json={"rows": [{**res["qb_only"][0], "mapped_account_id": nir["id"]}]})
    assert add["added"] == 1
    assert api.post("/api/accounting/reconcile/run", json={**body, "save": False})["counts"]["qb_only"] == 0
    # earlier checks: delete, restore
    rid = res["id"]
    api.post("/api/accounting/reconcile/delete", json={"ids": [rid]})
    assert rid not in {h["id"] for h in api.get("/api/accounting/reconcile/history")}
    api.post("/api/accounting/reconcile/restore", json={"ids": [rid]})
    assert rid in {h["id"] for h in api.get("/api/accounting/reconcile/history")}


def test_master_report_ten_tabs_with_tables(api, client, admin_headers):
    x = client.get("/api/accounting/master.xlsx?start=2083-04-01&end=2084-03-31", headers=admin_headers)
    assert x.status_code == 200
    wb = load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames == ["Summary", "P&L", "All Transactions", "Income", "Cost Of Goods", "Expenses", "Partner Draws", "Loans", "Transfers & Other", "Not Filed"]
    assert wb["Partner Draws"].tables  # a real Excel table (pivot-ready)
