"""Simulate: saved documents, pasted lines -> items and prices, price history, drafting a PO / order (2026-10-06)."""


def test_simulation_documents_and_tools(make, api):
    bolt = make.item(code="SIMB-5811", title="5/8-11 x 2 A325 HDG bolt")
    make.stock(bolt, 50, cost=0.42)
    s = api.post("/api/simulations/", json={"name": "Hudson job"})
    api.put(f"/api/simulations/{s['id']}", json={"doc": {"demand": [{"id": "d1", "multiplier": 8, "lines": []}], "summary": {"profit": 12}}})
    assert api.get(f"/api/simulations/{s['id']}")["doc"]["demand"][0]["multiplier"] == 8
    assert any(x["id"] == s["id"] and x["summary"]["profit"] == 12 for x in api.get("/api/simulations/"))
    rows = api.post("/api/simulations/parse", json={"text": "SIMB-5811\t500\t$0.95\nWIDGET-XYZ 9 pcs @ 2.50\nThanks!", "side": "demand"})
    assert rows[0]["item_id"] == bolt["id"] and rows[0]["qty"] == 500 and rows[0]["price"] == 0.95
    assert rows[1]["item_id"] is None and rows[1]["qty"] == 9 and rows[1]["price"] == 2.5     # not in our database
    ins = api.post("/api/simulations/insights", json={"item_ids": [bolt["id"]]})[str(bolt["id"])]
    assert ins["on_hand"] == 50 and ins["last_cost"] == 0.42 and ins["purchases"]
    v, c = make.vendor(), make.customer()
    po = api.post("/api/simulations/create-po", json={"vendor_id": v["id"], "lines": [{"item_id": bolt["id"], "quantity": 450, "price": 0.4}]})
    assert api.get(f"/api/purchase-orders/{po['id']}")["lines"][0]["quantity"] == 450
    o = api.post("/api/simulations/create-order", json={"customer_id": c["id"], "po_number": "SIM-1",
                                                        "lines": [{"item_id": bolt["id"], "quantity": 500, "price": 0.95}]})
    assert api.get(f"/api/customer-orders/{o['id']}")["status"] == "draft"
    copy = api.post(f"/api/simulations/{s['id']}/copy")
    assert copy["name"].endswith("(copy)")
    api.delete(f"/api/simulations/{copy['id']}")


def test_sticky_notes_page_search(make, api):
    o = make.order(lines=[(make.item(), 5, 1)], po_number="PO-NOTE-77", job_number="JOB-NB")
    api.post("/api/notes", json={"entity_type": "customer_order", "entity_id": o["id"], "text": "Call before shipping"})
    found = api.get("/api/notes/all?q=job-nb")
    assert len(found) == 1 and found[0]["record"]["po_number"] == "PO-NOTE-77" and found[0]["record"]["status"]
    assert api.get("/api/notes/all?q=nothing-like-this") == []


def test_simulation_generic_nuts(make, api):
    """Bulk 77-NUT (not marked generic) stands in for the specific nuts that match its size, thread and grade."""
    bulk = make.item(code="77-NUT", title="7/16-14 A194 GR 2H HVY HEX NUT DOM. HDG")
    bolt = make.item(code="77701", title="BOLT_HH_7/16-14x2_A325_TYPE1_HDG_W/A194-2H HEX NUT")
    nut = make.item(code="77701-NUT", title="7/16-14 A194 2H HVY HEX NUT HDG")
    other = make.item(code="77702-NUT", title="7/16-20 A194 2H HVY HEX NUT HDG")  # different thread; no bolt 77702, still not generic
    r = api.post("/api/simulations/generic", json={"item_ids": [bulk["id"], bolt["id"], nut["id"], other["id"]]})
    assert list(r["generics"]) == [str(bulk["id"])] and "not marked" in r["generics"][str(bulk["id"])]["why"]
    assert r["serves"] == {str(nut["id"]): [{"generic_id": bulk["id"], "match": "exact"}]}  # 77701-NUT is the bolt's own nut, not generic


def test_simulation_insights_last_buy_and_incoming(make, api):
    """The order tracker's defaults: the last purchase (price, PO, vendor) and what's still coming on open POs."""
    it = make.item(code="SIM-LB1", title="5/8 F436 washer")
    v = make.vendor()
    po = api.post("/api/simulations/create-po", json={"vendor_id": v["id"], "lines": [{"item_id": it["id"], "quantity": 300, "price": 0.031}]})
    ins = api.post("/api/simulations/insights", json={"item_ids": [it["id"]]})[str(it["id"])]
    assert ins["last_buy"]["price"] == 0.031 and ins["last_buy"]["doc"] == po["code"] and ins["last_buy"]["vendor_id"] == v["id"]
    assert ins["incoming"] == []  # a draft PO isn't incoming yet


def test_simulation_export_formats(client, admin_headers):
    sheets = [{"title": "Item by item", "columns": [{"name": "Item #"}, {"name": "Need", "fmt": "qty"}, {"name": "Profit", "fmt": "money"}],
               "rows": [["15420", 3150, 4220.85], ["16713", 15000, -12.5], ["x", None, None]]},
              {"title": "Order tracker", "columns": [{"name": "Vendor"}, {"name": "Price", "fmt": "price"}], "rows": [["G&T", 0.02497]]}]
    body = {"name": "Hudson / job 4065715", "facts": [["Profit", "$4,208.35"]], "sheets": sheets}
    for fmt, magic in (("xlsx", b"PK"), ("pdf", b"%PDF"), ("csv", "﻿".encode())):
        r = client.post("/api/simulations/export", json={**body, "fmt": fmt}, headers=admin_headers)
        assert r.status_code == 200 and r.content.startswith(magic), fmt
        assert f".{fmt}" in r.headers["content-disposition"]
    from openpyxl import load_workbook
    import io
    wb = load_workbook(io.BytesIO(client.post("/api/simulations/export", json={**body, "fmt": "xlsx"}, headers=admin_headers).content))
    assert wb.sheetnames == ["Summary", "Item by item", "Order tracker"] and wb["Item by item"]["C2"].value == 4220.85
    assert client.post("/api/simulations/export", json={**body, "fmt": "doc"}, headers=admin_headers).status_code == 400


def test_simulation_changes_detects_new_records(make, api):
    """Opening a simulation: the records it was built from (now), and open POs / confirmed orders / quotes on its items."""
    bolt, other = make.item(code="SIMCH-1"), make.item(code="SIMCH-2")
    mine = make.order(lines=[(bolt, 100, 2.0)])
    src = make.po(lines=[(bolt, 100, 1.0)])
    doc = {"demand": [{"id": "d1", "kind": "order", "ref_id": mine["id"], "lines": [{"id": "l1", "item_id": bolt["id"], "qty": 100, "price": 2.0}]}],
           "sources": [{"id": "s1", "kind": "po", "ref_id": src["id"], "ref_link": f"purchase-orders.html?id={src['id']}",
                        "lines": [{"id": "l2", "item_id": bolt["id"], "qty": 100, "price": 1.0}]}],
           "tracker": {"i1|open": {"status": "draft", "po_id": 99999999, "po_code": "PO-GONE"}}}
    s = api.post("/api/simulations/", json={"name": "Changes", "doc": doc})
    new_po = make.po(lines=[(bolt, 500, 0.9), (other, 5, 1)])          # draft PO by someone else
    new_order = make.order(lines=[(bolt, 40, 2.1)])                     # confirmed
    draft_order = make.order(lines=[(bolt, 7, 2.1)], confirm=False)     # draft: not offered
    r = api.post(f"/api/simulations/{s['id']}/changes")
    assert r["linked"][f"order:{mine['id']}"]["status"] == "confirmed"
    assert r["linked"][f"po:{src['id']}"]["lines"][0]["qty"] == 100
    assert r["linked"]["po:99999999"]["missing"] is True
    po_ids = {p["id"] for p in r["open"]["pos"]}
    assert {new_po["id"], src["id"]} <= po_ids
    assert [l["item_id"] for l in next(p for p in r["open"]["pos"] if p["id"] == new_po["id"])["lines"]] == [bolt["id"]]  # only this simulation's items
    order_ids = {o["id"] for o in r["open"]["orders"]}
    assert new_order["id"] in order_ids and draft_order["id"] not in order_ids
