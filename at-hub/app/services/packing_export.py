"""Packing lists as Excel (.xlsx) or CSV, from the same content the PDF prints (doc_context._packing_list), with real
numbers instead of formatted text so a customer's receiving system or a spreadsheet can use them.

Excel: one shipment -> a "Packing List" sheet (the header facts, then the lines) plus "Boxes" and "Pallets" sheets when
those are asked for; several shipments -> "Lines" / "Boxes" / "Pallets" sheets with the shipment on every row.
CSV holds one table: part=lines (default), boxes or pallets -- shipment, customer and PO # on every row."""
import csv
import io
import re
from typing import List

from sqlalchemy.orm import Session

from app.models import CustomerOrder, Shipment, StockItem
from app.services import doc_context


def _num(text):
    """'1,500' -> 1500.0; '—' / '' -> 0."""
    s = re.sub(r"[^\d.\-]", "", str(text or ""))
    try:
        return float(s) if s not in ("", "-", ".") else 0.0
    except ValueError:
        return 0.0


def _whole(v):
    return int(v) if float(v).is_integer() else v


def _collect(db: Session, shipment: Shipment, opts: dict) -> dict:
    ctx, rows = doc_context.build(db, "packing_list", shipment, {"show_notes": opts.get("notes", True)})
    order = db.get(CustomerOrder, shipment.order_id)
    from app.services.clock import local
    ship_day = local(shipment.ship_date).date() if shipment.ship_date else None
    head = {"Shipment": shipment.code, "Customer": ctx["customer"].get("name", ""), "Customer PO #": order.po_number if order else "",
            "Order #": order.code if order else "", "Job #": order.job_number if order else "", "Ship date": ship_day}
    facts = [("Shipment", shipment.code), ("Customer", ctx["customer"].get("name", "")), ("Ship to", ctx["customer"].get("ship_to", "")),
             ("Customer PO #", head["Customer PO #"]), ("Order #", head["Order #"]), ("Job #", head["Job #"]),
             ("Ship date", ship_day), ("Carrier", ctx["shipment"]["carrier"]), ("Tracking #", ctx["shipment"]["tracking"]),
             ("Lines", _whole(_num(ctx["shipment"]["lines"]))), ("Units", _whole(_num(ctx["shipment"]["units"]))),
             ("Boxes", _whole(_num(ctx["shipment"]["boxes"]))), ("Pallets", _whole(_num(ctx["shipment"]["pallets"]))),
             ("Weight (lbs)", _whole(_num(ctx["shipment"]["weight"]))), ("Notes", ctx["shipment"]["notes"])]
    lines = []
    for r in rows:
        desc, _, note = str(r["description"] or "").partition("\n")  # the PDF prints the line note under the description
        line = {"Line": _whole(_num(r["line_no"])) if r["line_no"] else "", "Part #": r["item_code"], "Description": desc}
        if opts.get("lots"):
            line["Lot #"] = r["lot"]
        line.update({"Qty ordered": _whole(_num(r["ordered"])), "Qty shipped": _whole(_num(r["shipped"])), "Backordered": _whole(_num(r["backorder"]))})
        if opts.get("boxes", True):
            line["Box breakdown"] = r["boxes"].replace("\n", "; ").replace("×", "x") if r["boxes"] != "—" else ""
        if opts.get("pallets"):
            line["Pallet"] = r["pallet"]
        if opts.get("notes", True):
            line["Line notes"] = note.strip()
        if any(x.get("combo_note") for x in rows):  # bolts + nuts sent as assembled units
            line["Packed as"] = r.get("combo_note") or ""
        lines.append(line)
    items = {i.id: i for i in db.query(StockItem).filter(StockItem.id.in_({b.item_id for b in shipment.boxes} or {0})).all()}
    line_no = {l.order_line_id: (l.order_line.line_no or 0) for l in shipment.lines if l.order_line}
    boxes = [{"Box #": b.box_number, "Line": line_no.get(b.order_line_id, ""), "Part #": items[b.item_id].code if b.item_id in items else "",
              "Description": items[b.item_id].title if b.item_id in items else "", "Qty in box": _whole(b.quantity_in_box),
              "Pack size": b.pack_size or "", "Lot #": b.lot_code or "", "Pallet": b.pallet_number or ""}
             for b in sorted(shipment.boxes, key=lambda b: (line_no.get(b.order_line_id, 0), b.box_number))]
    per_pallet = {}
    for b in shipment.boxes:
        if b.pallet_number:
            per_pallet[b.pallet_number] = per_pallet.get(b.pallet_number, 0) + 1
    pallets = [{"Pallet": p["pallet"], "Items": p["items"], "Boxes": per_pallet.get(p["pallet"], ""), "Weight (lbs)": _whole(_num(p["weight"])) if p["weight"] else "",
                "Dimensions (L x W x H in)": p["dimensions"], "Customer PO #": p["po"]} for p in ctx.get("_pallet_rows") or []]
    return {"code": shipment.code, "head": head, "facts": facts, "lines": lines, "boxes": boxes, "pallets": pallets,
            "company": ctx["company"].get("name", "")}


def collect(db: Session, shipments: List[Shipment], opts: dict) -> List[dict]:
    return [_collect(db, s, opts) for s in shipments]


def _table_rows(docs, part: str, with_head: bool):
    out = []
    for d in docs:
        for r in d[part]:
            out.append({**(d["head"] if with_head else {}), **r})
    return out


def to_csv(docs: List[dict], part: str = "lines") -> bytes:
    rows = _table_rows(docs, part if part in ("lines", "boxes", "pallets") else "lines", with_head=True)
    buf = io.StringIO()
    if rows:
        cols = list(dict.fromkeys(k for r in rows for k in r))
        w = csv.DictWriter(buf, fieldnames=cols, extrasaction="ignore", lineterminator="\r\n")
        w.writeheader()
        w.writerows(rows)
    return ("﻿" + buf.getvalue()).encode("utf-8")  # BOM: Excel opens it as UTF-8


def to_xlsx(docs: List[dict], opts: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    head_font, head_fill = Font(bold=True), PatternFill("solid", fgColor="E5E7EB")  # light grey: prints on any printer
    rule = Border(bottom=Side(style="thin", color="1B2430"))

    def table(ws, start_row, rows):
        if not rows:
            ws.cell(row=start_row, column=1, value="Nothing to list.")
            return
        cols = list(dict.fromkeys(k for r in rows for k in r))
        for c, name in enumerate(cols, 1):
            cell = ws.cell(row=start_row, column=c, value=name)
            cell.font, cell.fill, cell.border = head_font, head_fill, rule
        for i, r in enumerate(rows, start_row + 1):
            for c, name in enumerate(cols, 1):
                cell = ws.cell(row=i, column=c, value=r.get(name, ""))
                if hasattr(cell.value, "isoformat"):
                    cell.number_format = "yyyy-mm-dd"
                elif isinstance(cell.value, (int, float)):
                    cell.number_format = "#,##0.##" if isinstance(cell.value, float) else "#,##0"
                if name in ("Description", "Box breakdown", "Items", "Line notes"):
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
        for c, name in enumerate(cols, 1):
            longest = max([len(str(name))] + [min(60, max((len(x) for x in str(r.get(name, "")).split("\n")), default=0)) for r in rows])
            ws.column_dimensions[get_column_letter(c)].width = min(62, max(8, longest + 2))
        ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
        ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(cols))}{start_row + len(rows)}"

    wb = Workbook()
    ws = wb.active
    if len(docs) == 1:
        d = docs[0]
        ws.title = "Packing List"
        ws.cell(row=1, column=1, value=d["company"]).font = Font(bold=True, size=12)
        ws.cell(row=2, column=1, value=f"Packing list {d['code']}").font = Font(bold=True, size=14)
        r = 4
        for k, v in d["facts"]:
            if v in ("", None, 0) and k not in ("Lines", "Units"):
                continue
            ws.cell(row=r, column=1, value=k).font = head_font
            cell = ws.cell(row=r, column=2, value=v)
            cell.alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")
            if hasattr(v, "isoformat"):
                cell.number_format = "yyyy-mm-dd"
            r += 1
        table(ws, r + 1, d["lines"])
        ws.column_dimensions["B"].width = max(ws.column_dimensions["B"].width or 0, 34)
    else:
        ws.title = "Lines"
        table(ws, 1, _table_rows(docs, "lines", True))
        summary = wb.create_sheet("Shipments")
        table(summary, 1, [{**{k: v for k, v in d["facts"] if k != "Notes"}} for d in docs])
    if opts.get("boxes", True):
        table(wb.create_sheet("Boxes"), 1, _table_rows(docs, "boxes", len(docs) > 1))
    if opts.get("pallets"):
        table(wb.create_sheet("Pallets"), 1, _table_rows(docs, "pallets", len(docs) > 1))
    for sheet in wb.worksheets:
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
