"""Any page's tables as one file: Excel (a sheet per table), PDF (each table in turn, landscape) or CSV (every table
one after another, each under its title). The page sends what it shows -- [{title, columns: [{name, fmt}], rows: [[...]]}]
with real numbers -- so nothing is re-worked here. fmt: text | qty | money | price | pct."""
import csv
import io
import re
from typing import List

_FMT_XLSX = {"qty": "#,##0.##", "money": "$#,##0.00;-$#,##0.00", "price": "$#,##0.00000;-$#,##0.00000", "pct": "0.0\"%\""}


def _sheet_name(title: str, used: set) -> str:
    base = re.sub(r"[\[\]:*?/\\]", "", title or "Sheet")[:31] or "Sheet"
    name, n = base, 2
    while name.lower() in used:
        name = f"{base[:28]} {n}"
        n += 1
    used.add(name.lower())
    return name


def to_xlsx(title: str, sheets: List[dict], facts: List[list] = None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    head_font, head_fill = Font(bold=True), PatternFill("solid", fgColor="E5E7EB")  # light grey: prints on any printer
    rule = Border(bottom=Side(style="thin", color="1B2430"))
    wb = Workbook()
    wb.remove(wb.active)
    used = set()
    if facts:
        ws = wb.create_sheet(_sheet_name("Summary", used))
        ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=14)
        for i, (k, v) in enumerate(facts, 3):
            ws.cell(row=i, column=1, value=k).font = head_font
            ws.cell(row=i, column=2, value=v)
        ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 26, 40
    for sh in sheets:
        ws = wb.create_sheet(_sheet_name(sh.get("title"), used))
        cols = sh.get("columns") or []
        for c, col in enumerate(cols, 1):
            cell = ws.cell(row=1, column=c, value=col.get("name", ""))
            cell.font, cell.fill, cell.border = head_font, head_fill, rule
            cell.alignment = Alignment(wrap_text=True, vertical="bottom")
        rows = sh.get("rows") or []
        for r, row in enumerate(rows, 2):
            for c, col in enumerate(cols, 1):
                v = row[c - 1] if c - 1 < len(row) else None
                cell = ws.cell(row=r, column=c, value=v)
                if isinstance(v, (int, float)) and col.get("fmt") in _FMT_XLSX:
                    cell.number_format = _FMT_XLSX[col["fmt"]]
                elif isinstance(v, str) and len(v) > 40:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
        for c, col in enumerate(cols, 1):
            longest = max([len(str(col.get("name", "")))] + [min(50, len(str(row[c - 1]))) for row in rows if c - 1 < len(row) and row[c - 1] is not None])
            ws.column_dimensions[get_column_letter(c)].width = min(52, max(9, longest + 2))
        if cols:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(1, len(rows) + 1)}"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def to_csv(title: str, sheets: List[dict], facts: List[list] = None) -> bytes:
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow([title])
    for k, v in facts or []:
        w.writerow([k, v])
    for sh in sheets:
        w.writerow([])
        w.writerow([sh.get("title", "")])
        w.writerow([c.get("name", "") for c in sh.get("columns") or []])
        for row in sh.get("rows") or []:
            w.writerow(["" if v is None else v for v in row])
    return ("﻿" + out.getvalue()).encode("utf-8")  # BOM: Excel opens it as UTF-8


def _cell_text(v, fmt):
    if v is None or v == "":
        return ""
    if isinstance(v, (int, float)):
        if fmt == "money":
            return f"-${abs(v):,.2f}" if v < 0 else f"${v:,.2f}"
        if fmt == "price":
            s = f"{abs(v):,.5f}".rstrip("0")
            s = s + "0" * max(0, 2 - len(s.split(".")[1])) if "." in s else s + ".00"
            return ("-$" if v < 0 else "$") + s
        if fmt == "pct":
            return f"{v:.1f}%"
        return f"{v:,.0f}" if float(v).is_integer() else f"{v:,.2f}"
    return str(v)


def to_pdf(title: str, sheets: List[dict], facts: List[list] = None, subtitle: str = "") -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import landscape, letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle, KeepTogether
    styles = getSampleStyleSheet()
    cell = ParagraphStyle("cell", parent=styles["Normal"], fontSize=7.2, leading=8.6)
    cell_r = ParagraphStyle("cellr", parent=cell, alignment=2)
    head = ParagraphStyle("head", parent=cell, fontName="Helvetica-Bold")
    head_r = ParagraphStyle("headr", parent=head, alignment=2)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(letter), leftMargin=0.4 * inch, rightMargin=0.4 * inch, topMargin=0.45 * inch,
                            bottomMargin=0.45 * inch, title=title)
    width = landscape(letter)[0] - 0.8 * inch
    story = [Paragraph(title, styles["Title"])]
    if subtitle:
        story.append(Paragraph(subtitle, styles["Normal"]))
    if facts:
        ft = Table([[Paragraph(f"<b>{k}</b>", cell), Paragraph(str(v), cell)] for k, v in facts], colWidths=[1.8 * inch, 3.4 * inch], hAlign="LEFT")
        ft.setStyle(TableStyle([("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#D1D5DB")), ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]))
        story += [Spacer(1, 6), ft]
    esc = lambda s: str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")  # noqa: E731
    for sh in sheets:
        cols, rows = sh.get("columns") or [], sh.get("rows") or []
        if not cols:
            continue
        num = [c.get("fmt") in ("qty", "money", "price", "pct") for c in cols]
        data = [[Paragraph(esc(c.get("name", "")), head_r if num[i] else head) for i, c in enumerate(cols)]]
        for row in rows:
            data.append([Paragraph(esc(_cell_text(row[i] if i < len(row) else None, c.get("fmt"))), cell_r if num[i] else cell) for i, c in enumerate(cols)])
        # widths by content, text columns get the slack
        lens = [max([len(str(c.get("name", "")))] + [len(_cell_text(r[i] if i < len(r) else None, c.get("fmt"))) for r in rows[:300]]) for i, c in enumerate(cols)]
        lens = [min(46, max(5, n)) for n in lens]
        total = sum(lens)
        widths = [width * n / total for n in lens]
        t = Table(data, colWidths=widths, repeatRows=1)
        t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E5E7EB")), ("LINEBELOW", (0, 0), (-1, 0), 0.8, colors.HexColor("#1B2430")),
                               ("LINEBELOW", (0, 1), (-1, -1), 0.25, colors.HexColor("#D1D5DB")), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                               ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
                               ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3)]))
        story += [Spacer(1, 12), KeepTogether([Paragraph(esc(sh.get("title", "")), styles["Heading3"])]), t]
    doc.build(story)
    return buf.getvalue()
