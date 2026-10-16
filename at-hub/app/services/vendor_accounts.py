"""Vendor Accounts: vendors paid partly by wire and partly in cash (user, 2026-10-16 -- Akbar Ali Enterprise).

How such a vendor works: an invoice is paid part by wire, part in cash. Cash is hard to pay right away, so he later adds
old cash to another invoice (a $400 freight invoice billed $1,900) and the wire for it settles both; sometimes he adds more
than was owed. Their shared spreadsheet tracked it with comments and hand formulas, and the totals counted every recovery
twice. Here the account is its own ledger, so it always adds up:

  - a bill is the invoice at its TRUE amount (goods + freight + customs + other), split into wire due / cash due; when his
    paper invoice says more, paper_amount keeps it and the difference is a recovery of old cash, never a cost
  - a payment is one real money movement (a wire with its wire #, or cash handed over), applied to the wire part or the
    cash part of bills (VacctAlloc). A wire applied to a cash part IS the recovery; what isn't applied is a credit
  - balances are worked out from those rows every time: wire owed, cash owed, credit, total

The account is keyed by codes, so a fresh MRPeasy import keeps it (importers/mrpeasy/keep.py) and sync_all() puts the
PO payments back: on every PO an account's bills point at, the import's lump payment ("mrpeasy-po-export") is replaced by
one payment per (payment, bill), and the PO's vendor bill gets the true amount.

The opening history comes from their spreadsheet (plan_sheet / apply_sheet): rows under each SO, wire / cash due and
received, comments that say where a recovery went ("Sent $4000 along with freight invoice IF52"). The same reader
compares a statement the vendor sends back (compare)."""
import io
import json
import re
from collections import defaultdict
from datetime import datetime
from typing import Dict, List, Optional, Tuple

from fastapi import HTTPException
from sqlalchemy.orm import Session

from app.models import (PurchaseOrder, PurchaseOrderLine, PurchaseOrderPayment, VacctAccount, VacctAlloc, VacctBill, VacctPayment,
                        Vendor, VendorBill, VendorPayment)
from app.services.money import cents

TOL = 0.005
BY = "vendor-account"          # created_by on the PO payments / vendor bills this module writes
MRP_TAG = "mrpeasy-po-export"  # services/po_payments_csv.TAG: the import's lump payment per PO


# ---------------------------------------------------------------- codes --------------------------------------------

def _tokens(text) -> List:
    toks = [int(t) if t.isdigit() else t for t in re.findall(r"[A-Z]+|\d+", str(text or "").upper())]
    # a trailing financial year ("25-26", "26-27", "25/26") is not part of the number
    if len(toks) >= 3 and all(isinstance(t, int) for t in toks[-2:]) and 20 <= toks[-2] <= 40 and toks[-1] in (toks[-2], toks[-2] + 1):
        toks = toks[:-2]
    return toks


def code_key(text) -> Optional[Tuple[str, int]]:
    """'IF/47/25-26' -> ('IF', 47); '47' -> ('', 47); 'GR1/14/26-27' -> ('GR1', 14); '032/25-26' -> ('', 32)."""
    toks = _tokens(text)
    if not toks or not isinstance(toks[-1], int):
        return None
    return "".join(str(t) for t in toks[:-1]), toks[-1]


def keys_match(a, b) -> bool:
    """Same number, and the same prefix -- or one side has none (MRPeasy keeps '47' for 'IF/47/25-26')."""
    if not a or not b or a[1] != b[1]:
        return False
    return a[0] == b[0] or not a[0] or not b[0]


# ---------------------------------------------------------------- reading the books ---------------------------------

def _account(db: Session, account_id: int) -> VacctAccount:
    acc = db.get(VacctAccount, account_id)
    if not acc:
        raise HTTPException(status_code=404, detail="Vendor account not found")
    return acc


def _vendor(db: Session, acc: VacctAccount) -> Optional[Vendor]:
    return db.query(Vendor).filter(Vendor.code == acc.vendor_code).first()


def designations(b) -> List[dict]:
    try:
        return [d for d in json.loads(b.recovery_for or "[]") if d.get("bill_id") and (d.get("amount") or 0) > TOL]
    except ValueError:
        return []


def bill_total(b) -> float:
    return cents((b.goods or 0) + (b.freight or 0) + (b.customs or 0) + (b.other or 0))


def _bills(db, acc) -> List[VacctBill]:
    return (db.query(VacctBill).filter(VacctBill.account_id == acc.id)
            .order_by(VacctBill.position.is_(None), VacctBill.position, VacctBill.bill_date.is_(None), VacctBill.bill_date, VacctBill.id).all())


def _payments(db, acc) -> List[VacctPayment]:
    return (db.query(VacctPayment).filter(VacctPayment.account_id == acc.id)
            .order_by(VacctPayment.paid_date.is_(None), VacctPayment.paid_date, VacctPayment.id).all())


def _allocs(db, payment_ids) -> List[VacctAlloc]:
    return db.query(VacctAlloc).filter(VacctAlloc.payment_id.in_(list(payment_ids) or [-1])).order_by(VacctAlloc.id).all()


def match_pos(db: Session, acc: VacctAccount, bills: List[VacctBill]) -> Dict[int, Optional[PurchaseOrder]]:
    """Each bill's PO: picked by hand (po_code), else the vendor's PO with that SO #, else the PO holding that invoice #."""
    vendor = _vendor(db, acc)
    pos = db.query(PurchaseOrder).filter(PurchaseOrder.vendor_id == vendor.id).all() if vendor else []
    by_code = {p.code: p for p in pos}
    so_keys = [(code_key(p.vendor_so_number), p) for p in pos if p.vendor_so_number]
    bill_keys = [(code_key(vb.bill_number), vb.po_id) for vb in
                 (db.query(VendorBill).filter(VendorBill.po_id.in_([p.id for p in pos] or [-1])).all())]
    by_id = {p.id: p for p in pos}
    out = {}
    for b in bills:
        po = by_code.get(b.po_code) if b.po_code else None
        if not po and b.so_number:
            k = code_key(b.so_number)
            hits = [p for kk, p in so_keys if k and kk == k] or [p for kk, p in so_keys if keys_match(k, kk)]
            po = hits[0] if len({p.id for p in hits}) == 1 else None
        if not po and b.bill_number:
            k = code_key(b.bill_number)
            hits = {pid for kk, pid in bill_keys if keys_match(k, kk)}
            po = by_id.get(next(iter(hits))) if len(hits) == 1 else None
        out[b.id] = po
    return out


def view(db: Session, account_id: int) -> dict:
    """Everything the Vendor Account screen shows: bills with what's paid / left per part, payments with where they went,
    and the balances. Nothing is stored that could drift -- it's worked out from the rows."""
    acc = _account(db, account_id)
    vendor = _vendor(db, acc)
    bills, pays = _bills(db, acc), _payments(db, acc)
    ensure_codes(db, acc, pays)
    allocs = _allocs(db, [p.id for p in pays])
    pay_by_id = {p.id: p for p in pays}
    bill_by_id = {b.id: b for b in bills}
    pos = match_pos(db, acc, bills)
    paid = defaultdict(lambda: {"wire": 0.0, "cash": 0.0})
    by_bill, by_pay = defaultdict(list), defaultdict(list)
    for a in allocs:
        paid[a.bill_id][a.part] += a.amount
        p, b = pay_by_id.get(a.payment_id), bill_by_id.get(a.bill_id)
        row = {"id": a.id, "payment_id": a.payment_id, "payment_code": p.code if p else None, "bill_id": a.bill_id, "part": a.part,
               "amount": cents(a.amount), "payment_amount": cents(p.amount) if p else None,
               "bill_number": b.bill_number if b else None, "so_number": b.so_number if b else None,
               "kind": p.kind if p else None, "reference": p.reference if p else None,
               "paid_date": p.paid_date.date().isoformat() if p and p.paid_date else None,
               "recovery": bool(p and p.kind == "wire" and a.part == "cash"),
               "came_with": bill_by_id[p.came_with_bill_id].bill_number if p and p.came_with_bill_id in bill_by_id else None}
        by_bill[a.bill_id].append(row)
        by_pay[a.payment_id].append(row)
    today = datetime.utcnow()
    out_bills = []
    for i, b in enumerate(bills):
        total = bill_total(b)
        wp, cp = cents(paid[b.id]["wire"]), cents(paid[b.id]["cash"])
        po = pos.get(b.id)
        paper = cents(b.paper_amount) if b.paper_amount is not None else None
        out_bills.append({
            "id": b.id, "row": i + 1, "so_number": b.so_number, "so_amount": b.so_amount, "bill_number": b.bill_number,
            "bill_date": b.bill_date.date().isoformat() if b.bill_date else None,
            "goods": cents(b.goods), "freight": cents(b.freight), "customs": cents(b.customs), "other": cents(b.other),
            "total": total, "paper_amount": paper,
            "recovery_in": cents(paper - total) if paper is not None and paper > total + TOL else 0.0,
            "wire_due": cents(b.wire_due), "cash_due": cents(b.cash_due), "wire_paid": wp, "cash_paid": cp,
            "wire_left": cents(b.wire_due - wp), "cash_left": cents(b.cash_due - cp),
            "balance": cents(b.wire_due + b.cash_due - wp - cp),
            "days_open": (today - b.bill_date).days if b.bill_date else None,
            "po_code": po.code if po else None, "po_id": po.id if po else None, "po_picked": b.po_code,
            "note": b.note, "source": b.source, "paid_by": by_bill.get(b.id, []),
        })
    # old cash he put on another invoice (recovery_for): owed through THAT invoice until its wire comes -- this invoice's
    # cash goes down by it, that invoice's balance goes up by it; the total doesn't change
    by_out = {ob["id"]: ob for ob in out_bills}
    for ob in out_bills:
        ob.update({"moved_out": [], "carried_in": [], "balance_own": ob["balance"], "cash_left_own": ob["cash_left"]})
    for b in bills:
        for d in designations(b):
            tgt, src = by_out.get(d["bill_id"]), by_out[b.id]
            if not tgt:
                continue
            via = cents(sum(a.amount for a in allocs if a.bill_id == d["bill_id"] and a.part == "cash" and pay_by_id.get(a.payment_id)
                            and pay_by_id[a.payment_id].came_with_bill_id == b.id))
            room = cents(tgt["cash_left_own"] - sum(x["pending"] for x in tgt["moved_out"]))
            pending = cents(max(0.0, min(d["amount"] - via, room)))
            tgt["moved_out"].append({"to_bill_id": b.id, "to": b.bill_number, "amount": cents(d["amount"]), "paid": via, "pending": pending})
            src["carried_in"].append({"for_bill_id": tgt["id"], "for": tgt["bill_number"], "amount": cents(d["amount"]), "paid": via, "pending": pending})
    for ob in out_bills:
        out_p, in_p = cents(sum(x["pending"] for x in ob["moved_out"])), cents(sum(x["pending"] for x in ob["carried_in"]))
        ob["cash_left"] = cents(ob["cash_left_own"] - out_p)
        ob["carried_open"] = in_p
        ob["moved_open"] = out_p
        ob["balance"] = cents(ob["balance_own"] - out_p + in_p)
    out_pays = []
    for p in pays:
        applied = cents(sum(r["amount"] for r in by_pay.get(p.id, [])))
        out_pays.append({
            "id": p.id, "code": p.code, "review": bool(p.review), "picked_from": p.picked_from,
            "kind": p.kind, "paid_date": p.paid_date.date().isoformat() if p.paid_date else None,
            "amount": cents(p.amount), "reference": p.reference, "note": p.note, "source": p.source,
            "came_with_bill_id": p.came_with_bill_id,
            "came_with": bill_by_id[p.came_with_bill_id].bill_number if p.came_with_bill_id in bill_by_id else None,
            "applied": applied, "credit": cents(p.amount - applied), "parts": by_pay.get(p.id, []),
        })
    t = {k: cents(sum(b[k] for b in out_bills)) for k in ("total", "wire_due", "cash_due", "wire_paid", "cash_paid", "wire_left",
                                                          "cash_left", "balance", "recovery_in")}
    t["carried_open"] = cents(sum(b["carried_open"] for b in out_bills))
    t["wire_left"] = cents(t["wire_left"] + t["carried_open"])  # old cash billed on an invoice is paid with that invoice's wire
    t["paper"] = cents(sum((b["paper_amount"] if b["paper_amount"] is not None else b["total"]) for b in out_bills))
    t["credit"] = cents(sum(p["credit"] for p in out_pays if p["credit"] > 0))
    t["paid"] = cents(sum(p["amount"] for p in out_pays))
    t["wire_sent"] = cents(sum(p["amount"] for p in out_pays if p["kind"] == "wire"))
    t["cash_handed"] = cents(sum(p["amount"] for p in out_pays if p["kind"] == "cash"))
    t["recovered"] = cents(sum(r["amount"] for p in out_pays for r in p["parts"] if r["recovery"]))
    t["owed"] = cents(t["wire_left"] + t["cash_left"] - t["credit"])
    sos = so_view(db, acc, bills, out_bills, pos)
    t["not_invoiced"] = cents(sum(x["not_invoiced"] or 0 for x in sos))
    t["review"] = sum(1 for p in out_pays if p["review"])
    return {
        "account": {"id": acc.id, "vendor_code": acc.vendor_code, "vendor_name": vendor.name if vendor else acc.vendor_name,
                    "vendor_id": vendor.id if vendor else None, "cash_share": acc.cash_share, "note": acc.note},
        "bills": out_bills, "payments": out_pays, "totals": t, "sos": sos,
        "open_cash": sorted([b for b in out_bills if b["cash_left"] > TOL], key=lambda b: (b["bill_date"] or "9999", b["row"])),
    }


def ensure_codes(db: Session, acc: VacctAccount, pays: List[VacctPayment]) -> None:
    """Every payment has a short number (P-001, P-002 ... in date order) -- the statement's Paid By column uses it."""
    missing = [p for p in pays if not p.code]
    if not missing:
        return
    used = [int(p.code[2:]) for p in pays if p.code and p.code[2:].isdigit()]
    n = max(used, default=0)
    for p in sorted(missing, key=lambda x: (x.paid_date or datetime.max, x.id)):
        n += 1
        p.code = f"P-{n:03d}"
    db.commit()


def so_view(db: Session, acc: VacctAccount, bills: List[VacctBill], out_bills: List[dict], pos: dict) -> List[dict]:
    """Per vendor SO: its amount (from the sheet, else the PO's total), what's been invoiced, what's NOT invoiced yet
    (and how that would be paid at the usual cash share), what's still owed on its invoices."""
    po_total = {}
    for po in {p.id: p for p in pos.values() if p}.values():
        po_total[po.id] = cents(sum((l.quantity or 0) * (l.unit_cost or 0)
                                    for l in db.query(PurchaseOrderLine).filter(PurchaseOrderLine.po_id == po.id).all()))
    groups, order = defaultdict(list), []
    for b, ob in zip(bills, out_bills):
        k = b.so_number or f"(no SO) {b.bill_number or b.id}"
        if k not in groups:
            order.append(k)
        groups[k].append((b, ob))
    out = []
    share = (acc.cash_share or 0) / 100.0
    for k in order:
        rows = groups[k]
        po = next((pos.get(b.id) for b, _ in rows if pos.get(b.id)), None)
        so_amt = next((b.so_amount for b, _ in rows if b.so_amount), None)
        source = "sheet" if so_amt else None
        if not so_amt and po and po_total.get(po.id):
            so_amt, source = po_total[po.id], "po"
        goods = cents(sum(ob["goods"] for _, ob in rows))
        charges = cents(sum(ob["freight"] + ob["customs"] + ob["other"] for _, ob in rows))
        remaining = cents(max(0.0, so_amt - goods)) if so_amt else None
        over = cents(goods - so_amt) if so_amt and goods > so_amt + TOL else 0.0
        out.append({"so_number": rows[0][0].so_number, "so_amount": so_amt, "so_amount_from": source, "po_code": po.code if po else None,
                    "po_id": po.id if po else None, "invoices": [ob["bill_number"] for _, ob in rows], "bill_ids": [ob["id"] for _, ob in rows],
                    "goods": goods, "charges": charges, "not_invoiced": remaining, "over": over,
                    "expected_cash": cents((remaining or 0) * share), "expected_wire": cents((remaining or 0) - (remaining or 0) * share),
                    "balance": cents(sum(ob["balance"] for _, ob in rows))})
    return out


def list_accounts(db: Session) -> List[dict]:
    out = []
    for acc in db.query(VacctAccount).order_by(VacctAccount.vendor_name).all():
        v = view(db, acc.id)
        out.append({**v["account"], "totals": v["totals"], "bills": len(v["bills"]), "payments": len(v["payments"])})
    return out


# ---------------------------------------------------------------- changing the books --------------------------------

def create_account(db: Session, vendor_id: int, cash_share: float, by: str) -> VacctAccount:
    vendor = db.get(Vendor, vendor_id)
    if not vendor:
        raise HTTPException(status_code=404, detail="Vendor not found")
    if db.query(VacctAccount).filter(VacctAccount.vendor_code == vendor.code).first():
        raise HTTPException(status_code=400, detail=f"{vendor.name} already has a vendor account")
    acc = VacctAccount(vendor_code=vendor.code, vendor_name=vendor.name, cash_share=_share(cash_share), created_by=by)
    db.add(acc)
    db.commit()
    return acc


def _share(v) -> float:
    v = float(50 if v is None else v)
    if v < 0 or v > 100:
        raise HTTPException(status_code=400, detail="The cash share is a percent between 0 and 100")
    return v


def update_account(db: Session, account_id: int, data: dict) -> None:
    acc = _account(db, account_id)
    if "cash_share" in data:
        acc.cash_share = _share(data["cash_share"])
    if "note" in data:
        acc.note = data["note"]
    db.commit()


def delete_account(db: Session, account_id: int) -> None:
    acc = _account(db, account_id)
    pos = {po.id for po in match_pos(db, acc, _bills(db, acc)).values() if po}
    pays = [p.id for p in _payments(db, acc)]
    db.query(VacctAlloc).filter(VacctAlloc.payment_id.in_(pays or [-1])).delete(synchronize_session=False)
    db.query(VacctPayment).filter(VacctPayment.account_id == acc.id).delete(synchronize_session=False)
    db.query(VacctBill).filter(VacctBill.account_id == acc.id).delete(synchronize_session=False)
    db.query(PurchaseOrderPayment).filter(PurchaseOrderPayment.po_id.in_(pos or [-1]), PurchaseOrderPayment.created_by == BY).delete(synchronize_session=False)
    db.delete(acc)
    db.commit()


def _date(v) -> Optional[datetime]:
    if not v:
        return None
    if isinstance(v, datetime):
        return v
    try:
        return datetime.fromisoformat(str(v)[:10])
    except ValueError:
        raise HTTPException(status_code=400, detail=f"Not a date: {v}")


def _money_in(v, label, allow_negative=False) -> float:
    try:
        x = cents(float(v or 0))
    except (TypeError, ValueError):
        raise HTTPException(status_code=400, detail=f"{label}: not an amount")
    if x < 0 and not allow_negative:
        raise HTTPException(status_code=400, detail=f"{label} can't be negative")
    return x


def save_bill(db: Session, account_id: int, data: dict, by: str, bill_id: int = None) -> VacctBill:
    """A bill at its TRUE amount. Wire / cash split: given, else the account's cash share of the goods in cash and every
    charge (freight, customs, other) by wire. paper_amount (what his invoice says) is kept when it's more."""
    acc = _account(db, account_id)
    b = db.get(VacctBill, bill_id) if bill_id else VacctBill(account_id=acc.id, created_by=by, source="manual")
    if bill_id and (not b or b.account_id != acc.id):
        raise HTTPException(status_code=404, detail="Bill not found")
    for k in ("so_number", "bill_number", "po_code", "note"):
        if k in data:
            setattr(b, k, (str(data[k]).strip() or None) if data[k] is not None else None)
    if "bill_date" in data:
        b.bill_date = _date(data["bill_date"])
    for k, label in (("goods", "Goods"), ("freight", "Freight"), ("customs", "Customs"), ("other", "Other charges")):
        if k in data:
            setattr(b, k, _money_in(data[k], label))
    if "so_amount" in data:
        b.so_amount = _money_in(data["so_amount"], "SO amount") if data["so_amount"] not in (None, "") else None
    if "paper_amount" in data:
        b.paper_amount = _money_in(data["paper_amount"], "Invoice amount") if data["paper_amount"] not in (None, "") else None
    total = bill_total(b)
    if total <= 0:
        raise HTTPException(status_code=400, detail="A bill needs an amount (goods or a charge)")
    if "recovery_for" in data:
        _set_designations(db, acc, b, data["recovery_for"] or [], total)
    elif b.paper_amount is None or b.paper_amount <= total + TOL:
        b.recovery_for = None
    if b.paper_amount is not None and b.paper_amount < total - TOL:
        raise HTTPException(status_code=400, detail=f"The invoice says ${b.paper_amount:,.2f} but the charges add up to ${total:,.2f} -- "
                                                    "the true amount can't be more than the invoice")
    if data.get("cash_due") not in (None, ""):
        b.cash_due = _money_in(data["cash_due"], "Cash due")
        if b.cash_due > total + TOL:
            raise HTTPException(status_code=400, detail="Cash due can't be more than the bill")
    elif not bill_id or any(k in data for k in ("goods", "freight", "customs", "other")):
        b.cash_due = cents((b.goods or 0) * (acc.cash_share or 0) / 100.0)
    b.wire_due = cents(total - b.cash_due)
    if bill_id:
        wp, cp = _paid_parts(db, b.id)
        if wp > b.wire_due + TOL or cp > b.cash_due + TOL:
            raise HTTPException(status_code=400, detail=f"More is already paid on this bill (wire ${wp:,.2f}, cash ${cp:,.2f}) than the new "
                                                        "split allows -- change the payments first")
    if not bill_id:
        last = db.query(VacctBill.position).filter(VacctBill.account_id == acc.id).order_by(VacctBill.position.desc()).first()
        b.position = ((last[0] or 0) + 1) if last else 1
        db.add(b)
    db.flush()
    sync_account(db, acc)
    db.commit()
    return b


def _set_designations(db, acc, b, items, total) -> None:
    """The old cash inside this invoice and the invoices it is for -- no more than the extra on the paper, each no more
    than that invoice's cash."""
    extra = cents((b.paper_amount or 0) - total)
    out, used = [], 0.0
    for d in items:
        amt = _money_in(d.get("amount"), "Old cash amount")
        if amt <= TOL:
            continue
        t = db.get(VacctBill, int(d.get("bill_id") or 0))
        if not t or t.account_id != acc.id or (b.id and t.id == b.id):
            raise HTTPException(status_code=400, detail="Pick another invoice of this vendor for the old cash")
        if amt > cents(t.cash_due) + TOL:
            raise HTTPException(status_code=400, detail=f"{t.bill_number or t.so_number}: its cash is only ${t.cash_due:,.2f}")
        used = cents(used + amt)
        out.append({"bill_id": t.id, "amount": amt})
    if used > extra + TOL:
        raise HTTPException(status_code=400, detail=f"The invoice carries ${max(extra, 0):,.2f} of old cash (what it says less its charges) -- "
                                                    f"you assigned ${used:,.2f}")
    b.recovery_for = json.dumps(out) if out else None


def _paid_parts(db, bill_id) -> Tuple[float, float]:
    w = c = 0.0
    for a in db.query(VacctAlloc).filter(VacctAlloc.bill_id == bill_id).all():
        if a.part == "wire":
            w += a.amount
        else:
            c += a.amount
    return cents(w), cents(c)


def delete_bill(db: Session, account_id: int, bill_id: int) -> None:
    acc = _account(db, account_id)
    b = db.get(VacctBill, bill_id)
    if not b or b.account_id != acc.id:
        raise HTTPException(status_code=404, detail="Bill not found")
    if db.query(VacctAlloc).filter(VacctAlloc.bill_id == b.id).count():
        raise HTTPException(status_code=400, detail="Payments are applied to this bill -- take them off first (edit those payments)")
    db.query(VacctPayment).filter(VacctPayment.came_with_bill_id == b.id).update({"came_with_bill_id": None}, synchronize_session=False)
    for other in _bills(db, acc):
        ds = designations(other)
        if any(d["bill_id"] == b.id for d in ds):
            other.recovery_for = json.dumps([d for d in ds if d["bill_id"] != b.id]) or None
    db.delete(b)
    db.flush()
    sync_account(db, acc)
    db.commit()


def suggest(db: Session, account_id: int, kind: str, amount: float, for_bill_ids: List[int] = None,
            came_with_bill_id: int = None, payment_id: int = None) -> dict:
    """Where a payment would go: a wire pays the wire part of the bills it was sent for (else the oldest open wire), then
    the rest is old cash, oldest first; cash pays cash, oldest first. What's left is a credit."""
    v = view(db, account_id)
    if payment_id:  # editing: what this payment already pays is free again
        for b in v["bills"]:
            for r in b["paid_by"]:
                if r["payment_id"] == payment_id:
                    b[f"{r['part']}_left"] = cents(b[f"{r['part']}_left"] + r["amount"])
    left = cents(amount)
    out = []

    def take(b, part, cap=None):
        nonlocal left
        n = min(left, b[f"{part}_left"] if cap is None else cap)
        if n > TOL:
            out.append({"bill_id": b["id"], "part": part, "amount": cents(n), "bill_number": b["bill_number"], "so_number": b["so_number"],
                        "left_before": b[f"{part}_left"]})
            b[f"{part}_left"] = cents(b[f"{part}_left"] - n)
            left = cents(left - n)

    bills = v["bills"]
    by_id = {b["id"]: b for b in bills}
    oldest = sorted(bills, key=lambda b: (b["bill_date"] or "9999", b["row"]))
    named = [by_id[i] for i in (for_bill_ids or []) if i in by_id]
    if kind == "wire":
        for b in named:  # the invoices it was sent for: their wire part, then their cash
            take(b, "wire")
            take(b, "cash")
        if not named and came_with_bill_id in by_id:
            came = by_id[came_with_bill_id]
            take(came, "wire")  # an inflated invoice: its own charge, then the old cash it carries for the invoices named on it
            for c in came["carried_in"]:
                tgt = by_id.get(c["for_bill_id"])
                if tgt and c["pending"] > TOL:
                    take(tgt, "cash", cap=min(c["pending"], tgt["cash_left_own"]))
        elif not named:
            for b in oldest:
                take(b, "wire")
    else:
        for b in named:
            take(b, "cash")
    for b in oldest:
        take(b, "cash")
    return {"parts": out, "credit": cents(left)}


def save_payment(db: Session, account_id: int, data: dict, by: str, payment_id: int = None) -> VacctPayment:
    acc = _account(db, account_id)
    p = db.get(VacctPayment, payment_id) if payment_id else VacctPayment(account_id=acc.id, created_by=by, source="manual")
    if payment_id and (not p or p.account_id != acc.id):
        raise HTTPException(status_code=404, detail="Payment not found")
    kind = data.get("kind", p.kind or "wire")
    if kind not in ("wire", "cash"):
        raise HTTPException(status_code=400, detail="A payment is a wire or cash")
    amount = _money_in(data.get("amount", p.amount), "Amount")
    if amount <= 0:
        raise HTTPException(status_code=400, detail="Enter the amount paid")
    bills = {b.id: b for b in _bills(db, acc)}
    came = data.get("came_with_bill_id", p.came_with_bill_id)
    if came not in (None, "") and int(came) not in bills:
        raise HTTPException(status_code=400, detail="That invoice isn't in this account")
    p.kind, p.amount, p.came_with_bill_id = kind, amount, int(came) if came not in (None, "") else None
    if "paid_date" in data:
        p.paid_date = _date(data["paid_date"])
    for k in ("reference", "note"):
        if k in data:
            setattr(p, k, (str(data[k]).strip() or None) if data[k] is not None else None)
    parts = data.get("parts")
    if parts is None:
        parts = suggest(db, acc.id, kind, amount, data.get("for_bill_ids"), p.came_with_bill_id, payment_id)["parts"]
    if not payment_id:
        db.add(p)
        db.flush()
        ensure_codes(db, acc, _payments(db, acc))
    p.review = False if payment_id else p.review  # edited by a person = mapped
    _set_parts(db, acc, p, parts, bills)
    sync_account(db, acc)
    db.commit()
    return p


def _set_parts(db, acc, p, parts, bills) -> None:
    db.query(VacctAlloc).filter(VacctAlloc.payment_id == p.id).delete(synchronize_session=False)
    db.flush()
    total = 0.0
    for r in parts:
        bid, part = int(r["bill_id"]), r.get("part", "wire")
        amt = _money_in(r.get("amount"), "Applied amount")
        if amt <= TOL:
            continue
        b = bills.get(bid)
        if not b:
            raise HTTPException(status_code=400, detail="A part points at an invoice that isn't in this account")
        if part not in ("wire", "cash"):
            raise HTTPException(status_code=400, detail="Apply to the wire part or the cash part")
        wp, cp = _paid_parts(db, bid)
        room = cents((b.wire_due - wp) if part == "wire" else (b.cash_due - cp))
        if amt > room + TOL:
            raise HTTPException(status_code=400, detail=f"{b.bill_number or b.so_number}: only ${room:,.2f} of its {part} part is still open "
                                                        f"(you applied ${amt:,.2f}) -- leave the rest as a credit")
        total += amt
        db.add(VacctAlloc(payment_id=p.id, bill_id=bid, part=part, amount=amt))
        db.flush()
    if total > p.amount + TOL:
        raise HTTPException(status_code=400, detail=f"The parts add up to ${total:,.2f}, more than the ${p.amount:,.2f} paid")


def delete_payment(db: Session, account_id: int, payment_id: int) -> None:
    acc = _account(db, account_id)
    p = db.get(VacctPayment, payment_id)
    if not p or p.account_id != acc.id:
        raise HTTPException(status_code=404, detail="Payment not found")
    db.query(VacctAlloc).filter(VacctAlloc.payment_id == p.id).delete(synchronize_session=False)
    db.delete(p)
    db.flush()
    sync_account(db, acc)
    db.commit()


# ---------------------------------------------------------------- the POs follow -----------------------------------

def sync_account(db: Session, acc: VacctAccount) -> dict:
    """On every PO this account's bills point at: the vendor bill carries the true amount, and the PO's payments are
    exactly what the account applied to it (the import's lump payment is replaced). Other POs are left alone."""
    bills = _bills(db, acc)
    pos = match_pos(db, acc, bills)
    touched = {po.id: po for po in pos.values() if po}
    if touched:
        db.query(PurchaseOrderPayment).filter(PurchaseOrderPayment.po_id.in_(list(touched)),
                                              (PurchaseOrderPayment.reference == MRP_TAG) | (PurchaseOrderPayment.created_by == BY)
                                              ).delete(synchronize_session=False)
    vb_for = {}
    for b in bills:
        po = pos.get(b.id)
        if not po:
            continue
        k = code_key(b.bill_number) if b.bill_number else None
        vb = None
        if k:
            vb = next((x for x in db.query(VendorBill).filter(VendorBill.po_id == po.id).all() if keys_match(code_key(x.bill_number), k)), None)
        if not vb:
            vb = VendorBill(po_id=po.id, bill_number=b.bill_number or f"SO {b.so_number or po.vendor_so_number or po.code}",
                            bill_date=b.bill_date, amount=bill_total(b), note="From the Vendor Account", created_by=BY)
            db.add(vb)
            db.flush()
        elif abs((vb.amount or 0) - bill_total(b)) > TOL:
            vb.amount = bill_total(b)
        vb_for[b.id] = vb
    pays = {p.id: p for p in _payments(db, acc)}
    grouped = defaultdict(float)
    for a in _allocs(db, pays):
        if pos.get(a.bill_id):
            grouped[(a.payment_id, a.bill_id)] += a.amount
    n = 0
    for (pid, bid), amt in grouped.items():
        p, po = pays[pid], pos[bid]
        part_txt = "cash part, by wire" if (p.kind == "wire" and any(a.part == "cash" for a in _allocs(db, [pid]) if a.bill_id == bid)) else p.kind
        db.add(PurchaseOrderPayment(po_id=po.id, amount=cents(amt), paid_date=p.paid_date or po.order_date or datetime.utcnow(),
                                    method=p.kind.title(), reference=p.reference or (p.code or ""),
                                    vendor_bill_id=vb_for[bid].id if bid in vb_for else None,
                                    note=f"{p.code or 'Payment'} from the Vendor Account ({part_txt}, ${p.amount:,.2f} in all)", created_by=BY))
        n += 1
    db.flush()
    return {"pos": len(touched), "payments": n}


def sync_all(db: Session) -> None:
    """After a fresh import: every account picks up payments recorded elsewhere and puts its PO payments back
    (importers/mrpeasy/keep.py)."""
    for acc in db.query(VacctAccount).all():
        pickup(db, acc, commit=False)
        sync_account(db, acc)


def _kind_of(method) -> Tuple[str, bool]:
    """(wire | cash, unclear?) from a payment method someone picked on another screen."""
    m = (method or "").strip().lower()
    if "cash" in m:
        return "cash", False
    return "wire", not ("wire" in m or "transfer" in m)


def _map_to_po(db, p: VacctPayment, bills: List[VacctBill], amount: float, first_bill_number: Optional[str]) -> float:
    """Apply `amount` of payment p to a PO's invoices in this account: the invoice it named first, wire part then cash
    part for a wire (cash part first for cash). Returns what was applied (the rest stays a credit)."""
    order = sorted(bills, key=lambda b: (not (first_bill_number and keys_match(code_key(b.bill_number), code_key(first_bill_number))),
                                         b.position or 0, b.id))
    parts = ("wire", "cash") if p.kind == "wire" else ("cash", "wire")
    left, done = cents(amount), 0.0
    for b in order:
        for part in parts:
            wp, cp = _paid_parts(db, b.id)
            room = cents((b.wire_due - wp) if part == "wire" else (b.cash_due - cp))
            n = cents(min(left, room))
            if n > TOL:
                db.add(VacctAlloc(payment_id=p.id, bill_id=b.id, part=part, amount=n))
                db.flush()
                left, done = cents(left - n), cents(done + n)
    return done


def pickup(db: Session, acc: VacctAccount, commit: bool = True) -> int:
    """Payments recorded the old way for this vendor -- Purchase Orders > Vendor Payments, or Record Payment on one of the
    account's POs -- become account payments, mapped to that PO's invoices. They then live only here (the PO shows what
    the account paid, so nothing counts twice). One that isn't clearly a wire / cash, or isn't fully applied, is flagged
    to check."""
    vendor = _vendor(db, acc)
    if not vendor:
        return 0
    bills = _bills(db, acc)
    pos = match_pos(db, acc, bills)
    of_po = defaultdict(list)
    for b in bills:
        if pos.get(b.id):
            of_po[pos[b.id].id].append(b)
    n = 0
    for vp in db.query(VendorPayment).filter(VendorPayment.vendor_id == vendor.id).all():
        apps = db.query(PurchaseOrderPayment).filter(PurchaseOrderPayment.vendor_payment_id == vp.id).all()
        if any(a.po_id not in of_po for a in apps):
            continue  # part of it went to a PO this account doesn't hold: leave it where it is
        kind, unclear = _kind_of(vp.method)
        p = VacctPayment(account_id=acc.id, kind=kind, paid_date=vp.paid_date, amount=cents(vp.amount), reference=vp.reference,
                         note=vp.note, source="picked", picked_from=f"Vendor Payments {vp.code}", created_by=vp.created_by)
        db.add(p)
        db.flush()
        applied = 0.0
        for a in apps:
            vb = db.get(VendorBill, a.vendor_bill_id) if a.vendor_bill_id else None
            applied += _map_to_po(db, p, of_po[a.po_id], a.amount, vb.bill_number if vb else None)
            db.delete(a)
        p.review = unclear or cents(vp.amount - applied) > TOL
        db.delete(vp)
        n += 1
    for po_id, bl in of_po.items():
        for a in db.query(PurchaseOrderPayment).filter(PurchaseOrderPayment.po_id == po_id, PurchaseOrderPayment.vendor_payment_id.is_(None)).all():
            if a.reference == MRP_TAG or a.created_by == BY or (a.reference or "") == BY:
                continue
            kind, unclear = _kind_of(a.method)
            po = pos[bl[0].id]
            p = VacctPayment(account_id=acc.id, kind=kind, paid_date=a.paid_date, amount=cents(a.amount),
                             reference=None if (a.reference or "") == MRP_TAG else a.reference, note=a.note, source="picked",
                             picked_from=f"{po.code} Record Payment", created_by=a.created_by)
            db.add(p)
            db.flush()
            vb = db.get(VendorBill, a.vendor_bill_id) if a.vendor_bill_id else None
            applied = _map_to_po(db, p, bl, a.amount, vb.bill_number if vb else None)
            p.review = unclear or cents(a.amount - applied) > TOL
            db.delete(a)
            n += 1
    if n:
        db.flush()
        ensure_codes(db, acc, _payments(db, acc)) if commit else None
        sync_account(db, acc)
        if commit:
            db.commit()
    return n


# ---------------------------------------------------------------- their spreadsheet ---------------------------------

COLS = {  # our name -> header words (normalised: lower case, single spaces, no dots)
    "so": ("so #", "so no", "so number", "so"),
    "so_amount": ("so amount", "so amt", "so value"),
    "invoice": ("all invoices billed", "invoice #", "invoice no", "invoice number", "invoice"),
    "bill_date": ("invoice date", "bill date"),
    "subtotal": ("invoice sub total", "sub total", "subtotal", "goods"),
    "freight": ("freight",),
    "customs": ("customs / other", "customs", "other charges"),
    "total": ("invoice total (your invoice)", "invoice total", "total"),
    "recovery": ("old cash included", "recovery included", "recovery"),
    "wire_due": ("wire due",),
    "cash_due": ("cash due",),
    "wire_recd": ("wire recd", "wire received", "wire paid"),
    "cash_recd": ("cash recd", "cash received", "cash paid"),
    "cash_adj": ("cash adjustment",),
    "balance": ("balance",),
    "date": ("date recvd", "date received", "date paid", "date"),
    "payments": ("payments", "paid by"),
    "comments": ("comments", "comment", "notes"),
    "their_notes": ("their notes",),
}


def _norm(h) -> str:
    return re.sub(r"\s+", " ", str(h or "").replace(".", "").replace("\n", " ")).strip().lower()


def _cell_num(v) -> float:
    if v is None:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    neg = s.startswith("-") or (s.startswith("(") and s.endswith(")")) or "-$" in s
    s = re.sub(r"[^\d.]", "", s)
    if not s or s == ".":
        return 0.0
    try:
        x = float(s)
    except ValueError:
        return 0.0
    return -x if neg else x


def read_rows(data: bytes, filename: str) -> List[List]:
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), data_only=True)
        ws = wb.worksheets[0]
        return [list(r) for r in ws.iter_rows(values_only=True)]
    import csv
    text = data.decode("utf-8-sig", errors="replace")
    first = text.split("\n", 1)[0]
    delim = "\t" if first.count("\t") > first.count(",") else ","
    return [list(r) for r in csv.reader(io.StringIO(text), delimiter=delim)]


def parse_sheet(data: bytes, filename: str) -> List[dict]:
    """Rows of their sheet (or our statement): one per invoice, the SO carried down to the rows under it."""
    rows = read_rows(data, filename)
    head_i, cols = None, {}
    for i, r in enumerate(rows[:30]):
        names = [_norm(c) for c in r]
        found = {}
        for key, words in COLS.items():
            for w in words:  # the most specific header wins ("wire due" before "date")
                j = next((j for j, n in enumerate(names) if n == w and j not in found.values()), None)
                if j is not None:
                    found[key] = j
                    break
        if len(found) >= 5 and "invoice" in found:
            head_i, cols = i, found
            break
    if head_i is None:
        raise HTTPException(status_code=400, detail="Couldn't find the header row (SO #, Invoice, Invoice Total, Wire Due, Cash Due ...)")
    out, so, so_amount = [], None, None
    for r in rows[head_i + 1:]:
        get = lambda k: r[cols[k]] if k in cols and cols[k] < len(r) else None
        s_so, inv = str(get("so") or "").strip(), str(get("invoice") or "").strip()
        nums = {k: _cell_num(get(k)) for k in ("so_amount", "subtotal", "freight", "customs", "total", "recovery", "wire_due", "cash_due",
                                                "wire_recd", "cash_recd", "cash_adj", "balance")}
        if not s_so and not inv:
            if so and not any(abs(nums[k]) > TOL for k in ("total", "wire_due", "cash_due", "wire_recd", "cash_recd")):
                continue  # an empty row
            if not so:
                continue
            break  # the totals row: the list is over
        if s_so.lower().rstrip(":") in ("total", "totals", "grand total"):
            break  # our own statement's totals row
        if s_so:
            so, so_amount = s_so, nums["so_amount"] or None
        if not inv and not any(abs(nums[k]) > TOL for k in ("total", "wire_due", "cash_due")):
            continue
        total = nums["total"] or (nums["subtotal"] + nums["freight"] + nums["customs"])
        out.append({"row": len(out) + 1, "so": so, "so_amount": so_amount, "invoice": inv or None,
                    "bill_date": get("bill_date"), **{k: cents(v) for k, v in nums.items() if k != "so_amount"}, "total": cents(total),
                    "date": str(get("date") or "").strip(), "comments": str(get("comments") or "").strip(),
                    "their_notes": str(get("their_notes") or "").strip()})
    if not out:
        raise HTTPException(status_code=400, detail="No invoice rows found under the header")
    return out


MONTHS = {m: i for i, m in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}


def _fy_year(so, month) -> Optional[int]:
    """A date with no year ('22-Mar') falls in the SO's financial year ('IF11/25-26': Apr 2025 .. Mar 2026)."""
    m = re.search(r"(\d{2})\s*[-/]\s*(\d{2})\s*$", str(so or ""))
    if not m:
        return None
    a = 2000 + int(m.group(1))
    return a if month >= 4 else a + 1


def parse_date(text, so=None) -> Optional[datetime]:
    t = str(text or "").strip()
    if not t:
        return None
    if isinstance(text, datetime):
        return text
    m = re.match(r"^(\d{1,2})/(\d{1,2})/(\d{2,4})$", t)  # 6/16/2026: month first
    if m:
        y = int(m.group(3)) + (2000 if len(m.group(3)) == 2 else 0)
        return _mk(y, int(m.group(1)), int(m.group(2)))
    m = re.match(r"^(\d{1,2})-(\d{1,2})-(\d{4})$", t)  # 01-10-2025: day first (their sheet is kept in India)
    if m:
        return _mk(int(m.group(3)), int(m.group(2)), int(m.group(1)))
    m = re.match(r"^(\d{1,2})[- ]([A-Za-z]{3})[a-z]*[- ,]*(\d{2,4})?$", t)  # 10-Nov-25, 22-Mar
    if m:
        mon = MONTHS.get(m.group(2).lower()[:3])
        if mon:
            y = m.group(3)
            y = (int(y) + (2000 if len(y) == 2 else 0)) if y else _fy_year(so, mon)
            return _mk(y, mon, int(m.group(1))) if y else None
    m = re.match(r"^([A-Za-z]{3})[a-z]*\.? (\d{1,2}),? (\d{4})$", t)  # Dec 15 2025
    if m and MONTHS.get(m.group(1).lower()[:3]):
        return _mk(int(m.group(3)), MONTHS[m.group(1).lower()[:3]], int(m.group(2)))
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})", t)
    if m:
        return _mk(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    return None


def _mk(y, mo, d) -> Optional[datetime]:
    try:
        return datetime(y, mo, d)
    except ValueError:
        return None


REF_RE = re.compile(r"\b(?:inv(?:oice)?s?\.?\s*#?\s*)((?:[A-Z]{1,3}\d?\s*/?\s*)?\d{1,4})\b|\b([A-Z]{2,3}\d?\s*/?\s*\d{1,4})\b", re.I)
AMOUNT_RE = re.compile(r"\$\s*([\d,]+(?:\.\d+)?)|(?<![\d/.#A-Za-z])([\d,]{3,}(?:\.\d+)?)(?![\d/])")
WIRE_NO_RE = re.compile(r"\b(OW\d{6,})\b", re.I)


def _refs_in(line) -> List[Tuple[str, int]]:
    out = []
    for m in REF_RE.finditer(line):
        raw = m.group(1) or m.group(2)
        k = code_key(raw)
        if not k or k[0] in ("SO", "OW", "PO") or (m.group(2) and not re.search(r"[A-Za-z]", raw)):
            continue
        if m.group(2) and raw[:2].upper() != raw[:2]:  # a word like "on 14" -- only capitals count without "inv"
            continue
        out.append(k)
    return out


def _amounts_in(line) -> List[float]:
    out = []
    for m in AMOUNT_RE.finditer(line):
        raw = (m.group(1) or m.group(2) or "").replace(",", "")
        try:
            x = float(raw)
        except ValueError:
            continue
        if m.group(2) and (1990 <= x <= 2100 and "." not in raw):
            continue  # a year
        if x >= 1:
            out.append(x)
    return out


def plan_sheet(db: Session, account_id: int, data: bytes, filename: str) -> dict:
    """Read their sheet into bills + payments, explaining every recovery found in the comments, and check the result
    ends at the same balance the sheet does. Nothing is saved (apply_sheet does that)."""
    acc = _account(db, account_id)
    rows = parse_sheet(data, filename)
    notes, warnings = [], []
    for r in rows:
        r["key"] = code_key(r["invoice"]) if r["invoice"] else None
        r["goods"] = r["subtotal"] if r["subtotal"] else cents(r["total"] - r["freight"] - r["customs"])
        r["other"] = cents(r["total"] - r["goods"] - r["freight"] - r["customs"]) if r["subtotal"] else 0.0
        r["paper"] = r["total"]
        r["true_wire_due"], r["true_cash_due"] = r["wire_due"], r["cash_due"]
        if abs(r["wire_due"] + r["cash_due"] - r["total"]) > 0.02:
            warnings.append(f"Row {r['row']} {r['invoice'] or r['so']}: wire due + cash due (${r['wire_due'] + r['cash_due']:,.2f}) isn't the "
                            f"invoice total (${r['total']:,.2f}) -- the cash due was set to the difference")
            r["true_cash_due"] = cents(max(0.0, r["total"] - r["wire_due"]))
        r["recovery_out"] = []   # (row it pays, amount): money that came in on this row's wire for another row's cash
        r["inside"] = 0.0        # part of this invoice's own total that was old cash
        r["extra_wire"] = 0.0    # sent with this row's wire on top of the invoice
        r["explained_cash"] = 0.0
        if r.get("recovery"):    # our own statement says it outright
            r["inside"] = r["recovery"]
    by_key = [(r["key"], r) for r in rows if r["key"]]

    def find(k, not_row):
        hits = [r for kk, r in by_key if kk == k and r is not not_row] or [r for kk, r in by_key if keys_match(kk, k) and r is not not_row]
        return hits[0] if len(hits) == 1 else None

    for r in rows:
        if r["cash_recd"] <= TOL or not r["comments"]:
            continue
        for line in r["comments"].split("\n"):
            refs, amts = _refs_in(line), _amounts_in(line)
            if not refs or not amts:
                continue
            x = find(refs[0], r)
            if not x:
                warnings.append(f"Row {r['row']} {r['invoice']}: the comment mentions {''.join(map(str, refs[0]))} -- no such invoice in the sheet")
                continue
            amt = cents(amts[0])
            if amt > r["cash_recd"] - r["explained_cash"] + TOL:
                warnings.append(f"Row {r['row']} {r['invoice']}: '{line.strip()}' is more than the cash received there -- not used")
                continue
            own = x["paper"] - x["inside"] - amt
            fits_so = x["so_amount"] is None or own <= x["so_amount"] + max(1.0, 0.01 * (x["so_amount"] or 0))
            if x["wire_recd"] >= amt - TOL and own > TOL and fits_so and not x.get("recovery"):
                x["inside"] = cents(x["inside"] + amt)
                x.setdefault("for_rows", []).append((r["row"], amt))
                notes.append(f"Row {x['row']} {x['invoice']}: the invoice (${x['paper']:,.2f}) includes ${amt:,.2f} of old cash for row "
                             f"{r['row']} {r['invoice']} -- its real charge is ${cents(x['paper'] - x['inside']):,.2f}")
            elif not x.get("recovery"):
                x["extra_wire"] = cents(x["extra_wire"] + amt)
                notes.append(f"Row {x['row']} {x['invoice']}: ${amt:,.2f} was sent with its wire on top of the invoice, for row "
                             f"{r['row']} {r['invoice']} cash")
            x["recovery_out"].append((r, amt))
            r["explained_cash"] = cents(r["explained_cash"] + amt)
    # the real charge of an invoice with old cash inside: less goods, less wire due
    for r in rows:
        if r["inside"] > TOL:
            r["goods"] = cents(max(0.0, r["goods"] - r["inside"]))
            r["true_wire_due"] = cents(max(0.0, r["true_wire_due"] - r["inside"]))
    bills = []
    for r in rows:
        bills.append({"row": r["row"], "so_number": r["so"], "so_amount": r["so_amount"], "bill_number": r["invoice"],
                      "bill_date": parse_date(r["bill_date"], r["so"]).date().isoformat() if parse_date(r["bill_date"], r["so"]) else None,
                      "goods": r["goods"], "freight": r["freight"], "customs": r["customs"], "other": r["other"],
                      "paper_amount": r["paper"] if r["inside"] > TOL else None, "for_rows": r.get("for_rows", []),
                      "wire_due": r["true_wire_due"], "cash_due": r["true_cash_due"], "note": r["comments"] or None})
    payments = []
    for r in rows:
        dates = [d for d in (parse_date(x.strip(), r["so"]) for x in r["date"].split("\n")) if d]
        refs = WIRE_NO_RE.findall(r["comments"])
        own_wire = cents(r["wire_recd"] - r["inside"])
        wires = []
        if r["wire_recd"] > TOL or r["extra_wire"] > TOL:
            split = _split_wires(r, own_wire)
            if split:
                wires = split
            else:
                in_comment = next((d for d in (_date_in(x) for x in r["comments"].split("\n") if re.search(r"wire|OW\d", x, re.I)) if d), None)
                wires = [{"amount": cents(r["wire_recd"] + r["extra_wire"]), "date": dates[0] if dates else in_comment,
                          "reference": refs[0] if len(set(refs)) == 1 else (", ".join(dict.fromkeys(refs)) or None)}]
        # what each wire pays: this row's wire part first, then the old cash it carried
        todo = [("wire", r, min(own_wire, r["true_wire_due"]))] + [("cash", t, a) for t, a in r["recovery_out"]]
        for w in wires:
            w.update({"kind": "wire", "row": r["row"], "came_with_row": r["row"] if r["recovery_out"] else None, "parts": [],
                      "note": r["comments"] or None})
            room = w["amount"]
            for i, (part, target, amt) in enumerate(todo):
                n = cents(min(room, amt))
                if n > TOL:
                    w["parts"].append({"row": target["row"], "part": part, "amount": n})
                    room = cents(room - n)
                    todo[i] = (part, target, cents(amt - n))
            payments.append(w)
        cash = cents(r["cash_recd"] - r["explained_cash"])
        if cash > TOL:
            payments.append({"kind": "cash", "row": r["row"], "amount": cash, "date": dates[0] if dates else None, "reference": None,
                             "note": r["comments"] or "Cash (from the sheet)", "parts": [{"row": r["row"], "part": "cash", "amount": cash}]})
        if abs(r["cash_adj"]) > TOL:
            payments.append({"kind": "cash", "row": r["row"], "amount": abs(r["cash_adj"]), "date": dates[0] if dates else None,
                             "reference": None, "note": "Cash adjustment (from the sheet)",
                             "parts": [{"row": r["row"], "part": "cash", "amount": abs(r["cash_adj"])}]})
    payments = _merge_wires(payments, rows, notes)
    # parts can't be more than a bill's part: the over-collection stays a credit on the payment
    room = {(b["row"], "wire"): b["wire_due"] for b in bills}
    room.update({(b["row"], "cash"): b["cash_due"] for b in bills})
    for p in payments:
        for part in p["parts"]:
            k = (part["row"], part["part"])
            n = cents(min(part["amount"], max(0.0, room[k])))
            if n < part["amount"] - TOL:
                notes.append(f"Row {part['row']}: ${cents(part['amount'] - n):,.2f} more than its {part['part']} part was paid -- kept as a "
                             f"credit with the vendor")
            part["amount"], room[k] = n, cents(room[k] - n)
        p["parts"] = [x for x in p["parts"] if x["amount"] > TOL]
        p["credit"] = cents(p["amount"] - sum(x["amount"] for x in p["parts"]))
    # the check: same balance as the sheet
    sheet_owed = cents(sum(r["total"] for r in rows) - sum(r["wire_recd"] + r["cash_recd"] + r["cash_adj"] for r in rows))
    ours_wire = cents(sum(room[(b["row"], "wire")] for b in bills))
    ours_cash = cents(sum(room[(b["row"], "cash")] for b in bills))
    credit = cents(sum(p["credit"] for p in payments))
    ours = cents(ours_wire + ours_cash - credit)
    sheet_wire = cents(sum(r["wire_due"] - r["wire_recd"] for r in rows))
    sheet_cash = cents(sum(r["cash_due"] - r["cash_recd"] - r["cash_adj"] for r in rows))
    double = cents(sum(r["inside"] for r in rows))
    pos = match_pos(db, acc, [VacctBill(id=-b["row"], so_number=b["so_number"], bill_number=b["bill_number"]) for b in bills])
    for b in bills:
        po = pos.get(-b["row"])
        b["po_code"] = po.code if po else None
        b["total"] = cents(b["goods"] + b["freight"] + b["customs"] + b["other"])
    for p in payments:
        p["date"] = p["date"].date().isoformat() if p.get("date") else None
    missing = [b for b in bills if not b["po_code"]]
    return {
        "rows": len(rows), "bills": bills, "payments": payments, "notes": notes, "warnings": warnings,
        "check": {"sheet_owed": sheet_owed, "sheet_wire": sheet_wire, "sheet_cash": sheet_cash, "owed": ours, "wire_owed": ours_wire,
                  "cash_owed": ours_cash, "credit": credit, "difference": cents(ours - sheet_owed), "matches": abs(ours - sheet_owed) <= 0.05,
                  "double_counted": double, "sheet_billed": cents(sum(r["total"] for r in rows)),
                  "true_billed": cents(sum(b["total"] for b in bills))},
        "no_po": [f"Row {b['row']} {b['bill_number'] or b['so_number']}" for b in missing],
        "replaces": db.query(VacctBill).filter(VacctBill.account_id == acc.id).count() + db.query(VacctPayment).filter(VacctPayment.account_id == acc.id).count(),
    }


def _split_wires(r, own_wire) -> List[dict]:
    """'Wire # OW...043 for $6000' + 'Wire # OW...909 for $9256' on two lines, with two dates: two wires (only when they add up)."""
    lines = [x for x in r["comments"].split("\n") if re.search(r"wire", x, re.I) and not NOT_A_WIRE.search(x)]
    if len(lines) < 2 or r["extra_wire"] > TOL or r["inside"] > TOL:
        return []
    amts = []
    for x in lines:
        a = _amounts_in(x)
        if not a:
            return []
        amts.append(a[0])
    if abs(sum(amts) - r["wire_recd"]) > 0.02:
        return []
    dates = [d for d in (parse_date(x.strip(), r["so"]) for x in r["date"].split("\n")) if d]
    out = []
    for i, (x, a) in enumerate(zip(lines, amts)):
        ref = WIRE_NO_RE.findall(x)
        when = dates[i] if i < len(dates) else (_date_in(x) or (dates[-1] if dates else None))
        out.append({"amount": cents(a), "date": when, "reference": ref[0] if ref else None})
    return out


NOT_A_WIRE = re.compile(r"cover|applied|toward|total", re.I)  # "wire sent ... to cover $5426": about another row's money
DATE_IN_RE = re.compile(r"(?<![\d/])(\d{1,2}/\d{1,2}/\d{2,4})(?![\d/])")


def _date_in(text) -> Optional[datetime]:
    """A date written in a comment line ('Wire sent 10/16/2026', 'OW... - 09/09/26')."""
    m = DATE_IN_RE.search(text or "")
    return parse_date(m.group(1)) if m else None


def _merge_wires(payments: List[dict], rows: List[dict], notes: List[str]) -> List[dict]:
    """One real wire often paid several rows: the same wire # in their comments, or the same total written on each
    ('Wire sent for 12454 covering 3 invoices', '$23630 ($18408 + $5222)'). Those become ONE payment with every part."""
    by_row = {r["row"]: r for r in rows}
    wires = [p for p in payments if p["kind"] == "wire"]
    groups = []
    taken = set()
    refs = defaultdict(list)
    for p in wires:
        if p["reference"] and "," not in p["reference"]:
            refs[p["reference"]].append(p)
    for ref, ps in refs.items():
        if len(ps) > 1:
            groups.append((f"wire {ref}", ps))
            taken.update(id(p) for p in ps)

    def totals_of(p):
        out = set()
        for line in by_row[p["row"]]["comments"].split("\n"):
            if re.search(r"wire", line, re.I):
                out.update(a for a in _amounts_in(line) if a > p["amount"] + TOL)
        return out
    for p in wires:
        if id(p) in taken:
            continue
        for total in sorted(totals_of(p), reverse=True):
            ps = [q for q in wires if id(q) not in taken and total in totals_of(q)
                  and (not p["date"] or not q["date"] or p["date"] == q["date"])]
            if len(ps) > 1 and abs(sum(q["amount"] for q in ps) - total) <= 0.02:
                groups.append((f"one wire of ${total:,.2f}", ps))
                taken.update(id(q) for q in ps)
                break
    for label, ps in groups:
        ps.sort(key=lambda q: q["row"])
        keep = ps[0]
        for q in ps[1:]:
            keep["parts"] += q["parts"]
            keep["amount"] = cents(keep["amount"] + q["amount"])
            keep["date"] = keep["date"] or q["date"]
            keep["reference"] = keep["reference"] or q["reference"]
            if q.get("came_with_row") and not keep.get("came_with_row"):
                keep["came_with_row"] = q["came_with_row"]
            payments.remove(q)
        rows_txt = ", ".join(f"row {q['row']}" for q in ps)
        keep["note"] = f"One wire for {rows_txt}." + (f"\n{keep['note']}" if keep.get("note") else "")
        notes.append(f"{label[:1].upper() + label[1:]}: {rows_txt} were paid together (${keep['amount']:,.2f}) -- recorded as one payment")
    return payments


def apply_sheet(db: Session, account_id: int, data: bytes, filename: str, by: str) -> dict:
    """Load the sheet as this account's history -- everything already in the account is replaced."""
    acc = _account(db, account_id)
    plan = plan_sheet(db, account_id, data, filename)
    # a payment the sheet has no date for takes the date MRPeasy's payment on that PO had
    mrp_dates = {}
    for b in plan["bills"]:
        po = db.query(PurchaseOrder).filter(PurchaseOrder.code == b["po_code"]).first() if b["po_code"] else None
        lump = (db.query(PurchaseOrderPayment).filter(PurchaseOrderPayment.po_id == po.id, PurchaseOrderPayment.reference == MRP_TAG)
                .order_by(PurchaseOrderPayment.paid_date.desc()).first()) if po else None
        if lump and lump.paid_date:
            mrp_dates[b["row"]] = lump.paid_date.date().isoformat()
    for p in plan["payments"]:
        if not p["date"] and p["parts"] and mrp_dates.get(p["parts"][0]["row"]):
            p["date"] = mrp_dates[p["parts"][0]["row"]]
            p["note"] = ((p["note"] or "") + "\n(date from MRPeasy -- the sheet has none)").strip()
    old = [p.id for p in _payments(db, acc)]
    db.query(VacctAlloc).filter(VacctAlloc.payment_id.in_(old or [-1])).delete(synchronize_session=False)
    db.query(VacctPayment).filter(VacctPayment.account_id == acc.id).delete(synchronize_session=False)
    db.query(VacctBill).filter(VacctBill.account_id == acc.id).delete(synchronize_session=False)
    db.flush()
    bill_id = {}
    for b in plan["bills"]:
        row = VacctBill(account_id=acc.id, so_number=b["so_number"], so_amount=b["so_amount"], bill_number=b["bill_number"],
                        bill_date=_date(b["bill_date"]), goods=b["goods"], freight=b["freight"], customs=b["customs"], other=b["other"],
                        paper_amount=b["paper_amount"], wire_due=b["wire_due"], cash_due=b["cash_due"], note=b["note"], source="sheet",
                        position=b["row"], created_by=by)
        db.add(row)
        db.flush()
        bill_id[b["row"]] = row.id
    by_row = {b["row"]: b for b in plan["bills"]}
    for b in plan["bills"]:  # which invoice the old cash inside an invoice was for
        ds = [{"bill_id": bill_id[r], "amount": cents(min(a, by_row[r]["cash_due"]))} for r, a in b.get("for_rows", []) if r in bill_id]
        if ds:
            db.get(VacctBill, bill_id[b["row"]]).recovery_for = json.dumps(ds)
    for p in plan["payments"]:
        row = VacctPayment(account_id=acc.id, kind=p["kind"], paid_date=_date(p["date"]), amount=p["amount"], reference=p["reference"],
                           came_with_bill_id=bill_id.get(p.get("came_with_row")), note=p["note"], source="sheet", created_by=by)
        db.add(row)
        db.flush()
        for part in p["parts"]:
            db.add(VacctAlloc(payment_id=row.id, bill_id=bill_id[part["row"]], part=part["part"], amount=part["amount"]))
    db.flush()
    synced = sync_account(db, acc)
    db.commit()
    return {**plan["check"], "bills": len(plan["bills"]), "payments": len(plan["payments"]), "pos": synced["pos"]}


def compare(db: Session, account_id: int, data: bytes, filename: str) -> dict:
    """A statement the vendor sent back (or their own sheet) against the account: invoice by invoice, what differs."""
    v = view(db, account_id)
    theirs = parse_sheet(data, filename)
    ours = list(v["bills"])
    used, rows = set(), []
    for t in theirs:
        k = code_key(t["invoice"]) if t["invoice"] else None
        o = next((b for b in ours if b["id"] not in used and k and code_key(b["bill_number"]) == k), None) or \
            next((b for b in ours if b["id"] not in used and k and keys_match(code_key(b["bill_number"]), k)), None) or \
            (next((b for b in ours if b["id"] not in used and not b["bill_number"] and code_key(b["so_number"]) == code_key(t["so"])), None)
             if not t["invoice"] else None)
        their_bal = cents(t["total"] - t["recovery"] - t["wire_recd"] - t["cash_recd"] - t["cash_adj"])
        item = {"invoice": t["invoice"] or f"SO {t['so']}", "their_total": t["total"], "their_balance": their_bal,
                "their_notes": t["their_notes"] or None}
        if not o:
            rows.append({**item, "status": "only_theirs"})
            continue
        used.add(o["id"])
        our_paper = o["paper_amount"] if o["paper_amount"] is not None else o["total"]
        diffs = []
        if abs(our_paper - t["total"]) > 0.02:
            diffs.append(f"invoice ${t['total']:,.2f} on theirs, ${our_paper:,.2f} on ours")
        # an over-collection is a credit on ours, a negative balance on theirs
        credit = cents(sum(p["credit"] for p in v["payments"] if p["credit"] > TOL and any(r["bill_id"] == o["id"] for r in p["parts"])))
        own = o["balance_own"]
        if abs(own - their_bal) > 0.02 and abs(own - credit - their_bal) > 1.0 and abs(o["balance"] - their_bal) > 0.02:
            diffs.append(f"balance ${their_bal:,.2f} on theirs, ${own:,.2f} on ours")
        rows.append({**item, "our_total": our_paper, "our_balance": own, "bill_id": o["id"],
                     "status": "differs" if diffs else ("note" if t["their_notes"] else "same"), "differences": diffs})
    for o in ours:
        if o["id"] not in used:
            rows.append({"invoice": o["bill_number"] or f"SO {o['so_number']}", "our_total": o["total"], "our_balance": o["balance_own"],
                         "bill_id": o["id"], "status": "only_ours"})
    their_owed = cents(sum(t["total"] - t["recovery"] - t["wire_recd"] - t["cash_recd"] - t["cash_adj"] for t in theirs))
    counts = defaultdict(int)
    for r in rows:
        counts[r["status"]] += 1
    # a credit can sit on their rows (a negative balance) or apart (our statement's summary): either way is the same
    gross = cents(v["totals"]["owed"] + v["totals"]["credit"])
    diff = min((cents(their_owed - v["totals"]["owed"]), cents(their_owed - gross)), key=abs)
    return {"rows": rows, "counts": dict(counts), "their_owed": their_owed, "our_owed": v["totals"]["owed"], "difference": diff}


# ---------------------------------------------------------------- the statement -------------------------------------

def _paid_by(b: dict) -> str:
    return "; ".join(f"{r['payment_code'] or ''} {'wire' if r['kind'] == 'wire' else 'cash'} ${r['amount']:,.2f}"
                     f"{' old cash' + (' via ' + r['came_with'] if r['came_with'] else '') if r['recovery'] else ''}".strip()
                     for r in b["paid_by"])


def _moved_note(b: dict) -> str:
    out = [f"${x['pending']:,.2f} of this invoice's cash is billed on {x['to']}" for x in b["moved_out"] if x["pending"] > TOL]
    out += [f"includes ${x['pending']:,.2f} old cash for {x['for']}" for x in b["carried_in"] if x["pending"] > TOL]
    return "; ".join(out)


def _stmt_rows(v: dict, open_only: bool) -> List[dict]:
    return [b for b in v["bills"] if not open_only or abs(b["balance"]) > TOL]


STMT_HEAD = ["SO #", "SO Amount", "Invoice", "Invoice Date", "Invoice Sub Total", "Freight", "Customs / Other", "Invoice Total (Your Invoice)",
             "Old Cash Included", "Wire Due", "Cash Due", "Wire Recd", "Cash Recd", "Old Cash Moved", "Balance", "Paid By", "Comments", "Their Notes"]


def statement_xlsx(db: Session, account_id: int, company_name: str, open_only: bool = False) -> bytes:
    """Their layout, worked out right, with live formulas (Balance, totals, the summary) and a Their Notes column. Paid By
    names the payments (P-001 ...); the Payments sheet maps each payment to every invoice it paid; SO Summary shows what
    each SO still has to invoice. They can change it however they like and send it back."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter as L
    v = view(db, account_id)
    t = v["totals"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Statement"
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="DCE6F1")
    pay_fill = PatternFill("solid", fgColor="F2F2F2")
    thin = Side(style="thin", color="B7B7B7")
    money = '"$"#,##0.00;[Red]-"$"#,##0.00'

    def head(sheet, cols):
        sheet.append(cols)
        for c in sheet[sheet.max_row]:
            c.font, c.fill, c.alignment = bold, head_fill, Alignment(wrap_text=True, vertical="center")
            c.border = Border(bottom=thin)
    ws.append([company_name])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Statement of account -- {v['account']['vendor_name']} ({v['account']['vendor_code']})"])
    ws["A2"].font = Font(bold=True, size=12)
    ws.append([f"As of {datetime.utcnow():%B %d, %Y}" + ("  ·  open invoices only" if open_only else "")])
    ws.append(["Each invoice at its true amount. When an invoice of yours included old cash, it shows under Old Cash Included and "
               "that money paid the cash of the invoice it was for. Paid By names the payments (sheet Payments shows each one in full). "
               "Write anything you see differently under Their Notes."])
    ws.append([])
    head(ws, STMT_HEAD)
    first = ws.max_row + 1
    last_so = None
    for b in _stmt_rows(v, open_only):
        r = ws.max_row + 1
        paper = b["paper_amount"] if b["paper_amount"] is not None else b["total"]
        ws.append([b["so_number"], b["so_amount"] if b["so_number"] != last_so else None, b["bill_number"],
                   datetime.fromisoformat(b["bill_date"]) if b["bill_date"] else None, b["goods"], b["freight"],
                   cents(b["customs"] + b["other"]), paper, b["recovery_in"] or None, b["wire_due"], b["cash_due"], b["wire_paid"],
                   b["cash_paid"], cents(b["carried_open"] - b["moved_open"]) or None, f"=J{r}+K{r}-L{r}-M{r}+N{r}", _paid_by(b) or None,
                   _moved_note(b) or b["note"], None])
        last_so = b["so_number"]
        ws.cell(r, 4).number_format = "mm/dd/yyyy"
    last = ws.max_row
    ws.append([])
    tr = ws.max_row + 1
    ws.append(["Totals"] + [None] * 3 + [f"=SUM({L(c)}{first}:{L(c)}{last})" for c in range(5, 16)])
    for c in ws[tr]:
        c.font = bold
        c.border = Border(top=thin)
    for row in ws.iter_rows(min_row=first, max_row=tr, min_col=2, max_col=15):
        for c in row:
            if c.column != 3 and c.column != 4:
                c.number_format = money
    ws.append([])
    s0 = ws.max_row + 1
    nr = f"N{first}:N{last}"
    summary = [("Wire Balance", f'=J{tr}-L{tr}+SUMIF({nr},">0")'), ("Cash Balance", f'=K{tr}-M{tr}+SUMIF({nr},"<0")'),
               ("Credit With You (Paid Ahead)", t["credit"]), ("Total Balance", f"=O{s0}+O{s0 + 1}-O{s0 + 2}")]
    for i, (label, val) in enumerate(summary):
        ws.cell(s0 + i, 12, label).font = bold
        c = ws.cell(s0 + i, 15, val)
        c.number_format, c.font = money, bold
    extra = [("Old cash collected through invoices", t["recovered"]), ("Ordered but not invoiced yet (not owed yet)", t["not_invoiced"])]
    for i, (label, val) in enumerate(extra):
        ws.cell(s0 + len(summary) + 1 + i, 12, label).font = Font(italic=True)
        ws.cell(s0 + len(summary) + 1 + i, 15, val).number_format = money
    for i, w in enumerate([16, 12, 16, 12, 14, 11, 12, 16, 12, 12, 12, 12, 12, 13, 12, 46, 40, 30], 1):
        ws.column_dimensions[L(i)].width = w
    ws.freeze_panes = ws.cell(first, 4)

    # every payment, and every invoice it paid
    ps = wb.create_sheet("Payments")
    head(ps, ["Payment", "Date", "Type", "Wire #", "Payment Amount", "Invoice", "SO #", "Part", "Applied", "Note", "Their Notes"])
    for p in v["payments"]:
        lines = [(r["bill_number"] or "", r["so_number"] or "", ("Cash (old cash)" if r["recovery"] else r["part"].title()), r["amount"])
                 for r in p["parts"]]
        if p["credit"] > TOL:
            lines.append(("Not applied yet (credit with you)", "", "", p["credit"]))
        start = ps.max_row + 1
        for i, (inv, so, part, amt) in enumerate(lines or [("", "", "", 0)]):
            ps.append([p["code"], datetime.fromisoformat(p["paid_date"]) if p["paid_date"] else None, p["kind"].title(), p["reference"],
                       p["amount"], inv, so, part, amt, p["note"] if i == 0 else None, None] if i == 0 else
                      [None, None, None, None, None, inv, so, part, amt, None, None])
            ps.cell(ps.max_row, 2).number_format = "mm/dd/yyyy"
            for col in (5, 9):
                ps.cell(ps.max_row, col).number_format = money
        ps.append([f"{p['code']} total", None, None, None, None, f"{len(p['parts'])} invoice part(s)", None, None,
                   f"=SUM(I{start}:I{ps.max_row})", None, None])
        for c in ps[ps.max_row]:
            c.font, c.fill = bold, pay_fill
        ps.cell(ps.max_row, 9).number_format = money
    for i, w in enumerate([12, 12, 8, 22, 14, 34, 16, 16, 13, 50, 30], 1):
        ps.column_dimensions[L(i)].width = w
    ps.freeze_panes = "A2"

    # what each SO still has to invoice
    so = wb.create_sheet("SO Summary")
    head(so, ["SO #", "PO", "SO Amount", "Invoices", "Invoiced (Goods)", "Freight / Customs / Other", "Not Invoiced Yet",
              "Expected Wire", "Expected Cash", "Owed On Its Invoices"])
    f0 = so.max_row + 1
    for x in v["sos"]:
        r = so.max_row + 1
        so.append([x["so_number"], x["po_code"], x["so_amount"], ", ".join(i or "" for i in x["invoices"]), x["goods"], x["charges"],
                   f"=MAX(0,C{r}-E{r})" if x["so_amount"] else None, x["expected_wire"] or None, x["expected_cash"] or None, x["balance"]])
        for col in (3, 5, 6, 7, 8, 9, 10):
            so.cell(r, col).number_format = money
    tr2 = so.max_row + 1
    so.append(["Totals", None] + [f"=SUM({L(c)}{f0}:{L(c)}{tr2 - 1})" if c != 4 else None for c in range(3, 11)])
    for c in so[tr2]:
        c.font = bold
        c.number_format = money
    for i, w in enumerate([18, 12, 13, 40, 14, 14, 14, 13, 13, 14], 1):
        so.column_dimensions[L(i)].width = w
    so.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def statement_pdf(db: Session, account_id: int, open_only: bool = False) -> bytes:
    from reportlab.lib.pagesizes import landscape, letter
    from reportlab.lib.units import inch
    from reportlab.platypus import PageBreak, SimpleDocTemplate, Spacer
    from app.services import pdf as P
    from app.services.crud import get_company_profile
    v = view(db, account_id)
    t = v["totals"]
    company = get_company_profile(db)
    story = P._header(company, "Vendor Statement", v["account"]["vendor_name"] or v["account"]["vendor_code"])
    for f in story:  # the header is made for a portrait page: keep it on the left of this wide one
        f.hAlign = "LEFT"
    story.append(P.p(f"As of {datetime.utcnow():%B %d, %Y}" + (" · open invoices only" if open_only else "") +
                     ". Each invoice at its true amount; old cash included in an invoice is shown separately and paid the cash of "
                     "the invoice it was for. Paid By names the payments listed after the invoices.", "small"))
    story.append(Spacer(1, 8))
    story.append(P._meta_table([("Wire balance", P.money(t["wire_left"])), ("Cash balance", P.money(t["cash_left"])),
                                ("Credit with you", P.money(t["credit"]) if t["credit"] else None),
                                ("Total balance", P.money(t["owed"])),
                                ("Not invoiced yet", P.money(t["not_invoiced"]) if t["not_invoiced"] else None)]))
    story[-1].hAlign = "LEFT"
    story.append(Spacer(1, 12))
    head = ["SO #", "Invoice", "Invoice Total", "Old Cash Incl.", "Wire Due", "Cash Due", "Wire Recd", "Cash Recd", "Balance", "Paid By"]
    rows = []
    for b in _stmt_rows(v, open_only):
        paper = b["paper_amount"] if b["paper_amount"] is not None else b["total"]
        rows.append([b["so_number"] or "", b["bill_number"] or "", P.money(paper),
                     P.money(b["recovery_in"]) if b["recovery_in"] else "", P.money(b["wire_due"]), P.money(b["cash_due"]),
                     P.money(b["wire_paid"]), P.money(b["cash_paid"]), P.money(b["balance"]),
                     ", ".join(dict.fromkeys(r["payment_code"] for r in b["paid_by"] if r["payment_code"]))])
    rows.append(["Totals", "", P.money(t["paper"]), P.money(t["recovery_in"]) if t["recovery_in"] else "", P.money(t["wire_due"]),
                 P.money(t["cash_due"]), P.money(t["wire_paid"]), P.money(t["cash_paid"]), P.money(t["balance"]), ""])
    story.append(P._data_table(head, rows, [w * inch for w in (1.2, 1.2, 0.95, 0.9, 0.9, 0.9, 0.9, 0.9, 0.9, 1.4)], right_cols=range(2, 9)))
    if v["payments"]:
        story += [PageBreak(), P.p("Payments -- what each one paid", "title"), Spacer(1, 6)]
        prow = []
        for p in v["payments"]:
            first = True
            for r in p["parts"]:
                prow.append([p["code"] if first else "", (p["paid_date"] or "") if first else "", (p["kind"].title() + (" " + p["reference"] if p["reference"] else "")) if first else "",
                             P.money(p["amount"]) if first else "", f"{r['bill_number'] or ''} (SO {r['so_number'] or '-'})",
                             "old cash" if r["recovery"] else r["part"], P.money(r["amount"])])
                first = False
            if p["credit"] > TOL:
                prow.append([p["code"] if first else "", "", "", P.money(p["amount"]) if first else "", "not applied yet (credit)", "", P.money(p["credit"])])
        story.append(P._data_table(["Payment", "Date", "Type / Wire #", "Amount", "Invoice (SO)", "Part", "Applied"], prow,
                                   [w * inch for w in (0.7, 0.9, 2.2, 1.0, 2.6, 0.8, 1.0)], right_cols=(3, 6)))
    open_sos = [x for x in v["sos"] if x["not_invoiced"]]
    if open_sos:
        story += [Spacer(1, 14), P.p("Ordered but not invoiced yet", "title"), Spacer(1, 6)]
        story.append(P._data_table(["SO #", "SO Amount", "Invoiced", "Not Invoiced Yet", "Expected Wire", "Expected Cash"],
                                   [[x["so_number"] or "", P.money(x["so_amount"]), P.money(x["goods"]), P.money(x["not_invoiced"]),
                                     P.money(x["expected_wire"]), P.money(x["expected_cash"])] for x in open_sos],
                                   [w * inch for w in (1.6, 1.2, 1.2, 1.3, 1.2, 1.2)], right_cols=range(1, 6)))
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(letter), title="Vendor Statement", leftMargin=0.5 * inch, rightMargin=0.5 * inch,
                            topMargin=0.5 * inch, bottomMargin=0.6 * inch)
    doc.build(story, canvasmaker=P._numbered_canvas(P._footer_text(company)))
    return buf.getvalue()
