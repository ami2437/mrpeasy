"""Accounting, the rest (routes/accounting.py): bank-style search with live totals / groups / stats, comments, send a
line back to Review, the Ignored list, Check My Books, Partner Accounts, Payees, 1099, Recurring, scratch sheets
(Projections, Loans quick calculations) and reconciling against QuickBooks. Built on accounting.py."""
import csv
import io
import itertools
import json
import re
import statistics
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from typing import Dict, List, Optional, Tuple

from fastapi import HTTPException
from sqlalchemy import or_
from sqlalchemy.orm import Session, selectinload

from app.models import (AcctAccount, AcctCheckOk, AcctComment, AcctIgnored, AcctReconcile, AcctRule, AcctSheet, AcctSource, AcctSplit,
                        AcctTxn, AppSetting)
from app.services import accounting as A
from app.services import acct_mapper, acct_rules
from app.services.money import cents


def _setting(db: Session, key: str, default):
    row = db.get(AppSetting, key)
    return json.loads(row.value) if row and row.value else default


def _save_setting(db: Session, key: str, value, by: str) -> None:
    row = db.get(AppSetting, key)
    if not row:
        row = AppSetting(key=key)
        db.add(row)
    row.value, row.updated_by = json.dumps(value), by


def period_key(start: Optional[date], end: Optional[date]) -> Optional[str]:
    return f"{start.isoformat()}..{end.isoformat()}" if start and end else None


def prev_range(start: date, end: date) -> Tuple[date, date]:
    """The period just before, the same length (FY -> the FY before)."""
    days = (end - start).days + 1
    return start - timedelta(days=days), start - timedelta(days=1)


# ================================================================ search
_NUM = r"\$?\d[\d,]*(?:\.\d{1,2})?"


def _num(s: str) -> float:
    return float(s.replace("$", "").replace(",", ""))


def parse_q(q: str) -> dict:
    """A bank-style search box: amounts ("1072.50", "$1,072.50", "1072"), ranges (">500", "<25", "500-1000"), money
    out / in ("-300", "+50"), words (all of them, any order) and "quoted phrases"."""
    out = {"words": [], "phrases": [], "amount": [], "loose_num": []}
    q = q or ""
    out["phrases"] = [p.strip().lower() for p in re.findall(r'"([^"]+)"', q) if p.strip()]
    for tok in re.sub(r'"[^"]*"', " ", q).split():
        m = None
        if re.fullmatch(rf"[<>]=?{_NUM}", tok):
            op = tok[0] + ("=" if tok[1] == "=" else "")
            out["amount"].append(("gt" if op[0] == ">" else "lt", _num(tok.lstrip("<>="))))
        elif (m := re.fullmatch(rf"({_NUM})\s*-\s*({_NUM})", tok)):
            a, b = sorted((_num(m.group(1)), _num(m.group(2))))
            out["amount"].append(("between", (a, b)))
        elif re.fullmatch(rf"[-+]{_NUM}", tok):
            v = _num(tok[1:])
            out["amount"].append(("signed", (-1 if tok[0] == "-" else 1, v, "." in tok)))
        elif re.fullmatch(rf"\$[\d,]+(?:\.\d{{1,2}})?|\d[\d,]*\.\d{{1,2}}|\d{{1,3}}(?:,\d{{3}})+", tok):
            out["amount"].append(("eq", (_num(tok), "." in tok)))
        elif re.fullmatch(r"\d+", tok):
            out["loose_num"].append(tok)  # an amount in whole dollars, or a check / reference number
        else:
            out["words"].append(tok.lower())
    return out


def _amount_ok(cond, amount: float) -> bool:
    a = abs(amount)
    kind, v = cond
    if kind == "gt":
        return a > v - 0.005
    if kind == "lt":
        return a < v + 0.005
    if kind == "between":
        return v[0] - 0.005 <= a <= v[1] + 0.005
    if kind == "eq":
        val, exact = v
        return abs(a - val) < 0.005 if exact else int(a) == int(val)
    if kind == "signed":
        sign, val, exact = v
        return (amount < 0 if sign < 0 else amount > 0) and (abs(a - val) < 0.005 if exact else int(a) == int(val))
    return True


def _text_of(t: AcctTxn, accounts) -> str:
    acc = accounts.get(t.account_id)
    parts = [t.description, t.bank_description or "", t.payee or "", t.note or "", A.path(acc, accounts) if acc else ""]
    parts += [A.path(accounts[x.account_id], accounts) for x in t.splits if x.account_id in accounts]
    return " ".join(parts).lower()


def search(db: Session, p: dict) -> dict:
    """Transactions for the lists (Review, Transactions, drill-downs): every filter, totals / stats / groups for ALL
    that match, the page asked for."""
    accounts, sources = A.maps(db)
    qy = db.query(AcctTxn).options(selectinload(AcctTxn.splits))
    dfrom, dto = A.parse_day(p.get("date_from")), A.parse_day(p.get("date_to"))
    if not p.get("all_dates"):
        dfrom = dfrom or A.parse_day(p.get("start"))
        dto = dto or A.parse_day(p.get("end"))
    if dfrom:
        qy = qy.filter(AcctTxn.date >= A._dt(dfrom))
    if dto:
        qy = qy.filter(AcctTxn.date <= A._dt(dto))
    if p.get("source_id"):
        qy = qy.filter(AcctTxn.source_id == int(p["source_id"]))
    if p.get("ids"):  # "show them": the lines a bulk move just moved
        qy = qy.filter(AcctTxn.id.in_([int(x) for x in str(p["ids"]).split(",") if x.strip().isdigit()]))
    only = None  # filtered to an account (or a type): a split line counts only its part(s) in it
    if p.get("account_id"):
        ids = [int(p["account_id"])] + A.descendants(int(p["account_id"]), accounts)
        only = set(ids)
        qy = qy.filter(or_(AcctTxn.account_id.in_(ids), AcctTxn.splits.any(AcctSplit.account_id.in_(ids))))
    if p.get("kind"):
        ids = [a.id for a in accounts.values() if a.kind == p["kind"]]
        only = set(ids) if only is None else only & set(ids)
        qy = qy.filter(or_(AcctTxn.account_id.in_(ids), AcctTxn.splits.any(AcctSplit.account_id.in_(ids))))
    st = p.get("status")
    if st == "unfiled":
        qy = qy.filter(AcctTxn.account_id.is_(None), ~AcctTxn.splits.any())
    elif st == "review":  # not filed, or filed but flagged (rules disagree, a rule's "check it", split in QuickBooks, Reconcile)
        qy = qy.filter(or_((AcctTxn.account_id.is_(None)) & ~AcctTxn.splits.any(), AcctTxn.flag.isnot(None)))
    elif st == "flagged":
        qy = qy.filter(AcctTxn.flag.isnot(None))
    elif st == "auto":
        qy = qy.filter(AcctTxn.reviewed.is_(False))
    elif st == "expected":
        qy = qy.filter(AcctTxn.expected.is_(True))
    elif st == "typed":
        qy = qy.filter(AcctTxn.origin == "manual")
    elif st == "split":
        qy = qy.filter(AcctTxn.splits.any())
    elif st == "sent_back":
        qy = qy.filter(AcctTxn.sent_back.is_(True))
    if p.get("how"):
        qy = qy.filter(AcctTxn.how == p["how"])
    if p.get("sign") == "in":
        qy = qy.filter(AcctTxn.amount > 0)
    elif p.get("sign") == "out":
        qy = qy.filter(AcctTxn.amount < 0)
    rows = qy.all()
    # amount from / to / exact (always the size of the amount; in / out is "sign")
    amin, amax, amt = (A.parse_money(p.get(k)) if p.get(k) not in (None, "") else None for k in ("amin", "amax", "amt"))
    if amin is not None:
        rows = [t for t in rows if abs(t.amount) >= amin - 0.005]
    if amax is not None:
        rows = [t for t in rows if abs(t.amount) <= amax + 0.005]
    if amt is not None:
        rows = [t for t in rows if abs(abs(t.amount) - amt) < 0.005]
    if p.get("has_note"):
        rows = [t for t in rows if (t.note or "").strip()]
    q = parse_q(p.get("q") or "")
    if any(q.values()):
        keep = []
        for t in rows:
            text = _text_of(t, accounts)
            if any(w not in text for w in q["words"]) or any(ph not in text for ph in q["phrases"]):
                continue
            if not all(_amount_ok(c, t.amount) for c in q["amount"]):
                continue
            if not all(int(abs(t.amount)) == int(n) or re.search(rf"\b{n}\b", text) for n in q["loose_num"]):
                continue
            keep.append(t)
        rows = keep
    # sort
    sort = p.get("sort") or "date_desc"
    if st == "review" and not p.get("sort"):
        sort = "flag_first"
    acc_name = lambda t: (A.path(accounts[t.account_id], accounts) if t.account_id in accounts else ("~split" if t.splits else "~~")).lower()
    key = {"date_desc": lambda t: (t.date, t.id), "date_asc": lambda t: (t.date, t.id), "amount_desc": lambda t: (abs(t.amount), t.id),
           "amount_asc": lambda t: (abs(t.amount), t.id), "payee_asc": lambda t: ((t.payee or "~").lower(), t.date), "payee_desc": lambda t: ((t.payee or "").lower(), t.date),
           "account_asc": lambda t: (acc_name(t), t.date), "account_desc": lambda t: (acc_name(t), t.date),
           "flag_first": lambda t: (0 if t.flag else 1, -t.date.toordinal(), -t.id)}.get(sort, lambda t: (t.date, t.id))
    rows.sort(key=key, reverse=sort in ("date_desc", "amount_desc", "payee_desc", "account_desc"))
    # totals / stats / groups for everything that matched -- a split line, when the list is one account's, by its part there
    def part(t):
        if only is None or not t.splits:
            return t.amount
        return cents(sum(x.amount for x in t.splits if x.account_id in only))
    amts = [part(t) for t in rows]
    ins = [a_ for a_ in amts if a_ > 0]
    outs = [a_ for a_ in amts if a_ < 0]
    tot = {"count": len(rows), "money_in": cents(sum(ins)), "money_out": cents(sum(outs)), "net": cents(sum(ins) + sum(outs))}
    stats = {}
    if rows:
        ds = [t.date.date() for t in rows]
        months = max(1, (max(ds).year - min(ds).year) * 12 + max(ds).month - min(ds).month + 1)
        stats = {"largest_in": cents(max(ins)) if ins else None, "largest_out": cents(min(outs)) if outs else None,
                 "average": cents(tot["net"] / len(rows)), "first": min(ds).isoformat(), "last": max(ds).isoformat(),
                 "per_month": cents(tot["net"] / months), "months": months}
    groups = []
    gmode = p.get("group")
    if gmode in ("account", "payee", "month", "source"):
        g = defaultdict(lambda: {"count": 0, "in": 0.0, "out": 0.0})
        for t in rows:
            if gmode == "account":  # a split line: each part under its own account
                ps = [(x.account_id, x.amount) for x in t.splits if only is None or x.account_id in only] if t.splits else [(t.account_id, t.amount)]
                for aid, v in ps:
                    a = accounts.get(aid)
                    x = g[A.path(a, accounts) if a else "Not filed"]
                    x["count"] += 1
                    x["in" if v > 0 else "out"] += v
                continue
            if gmode == "payee":
                k = t.payee or acct_mapper.payee_label(t.bank_description or t.description)
            elif gmode == "month":
                k = t.date.strftime("%Y-%m")
            else:
                k = sources[t.source_id].name if t.source_id in sources else "?"
            x = g[k]
            x["count"] += 1
            v = part(t)
            x["in" if v > 0 else "out"] += v
        groups = [{"key": k, "count": v["count"], "in": cents(v["in"]), "out": cents(v["out"]), "net": cents(v["in"] + v["out"])} for k, v in g.items()]
        groups.sort(key=(lambda x: x["key"]) if gmode == "month" else (lambda x: -abs(x["net"])))
    # running balance (one bank account, by date): the bank's balance after each line
    balance = {}
    if p.get("source_id") and sort in ("date_desc", "date_asc"):
        s = sources.get(int(p["source_id"]))
        if s and s.kind != "book":
            run = s.opening_balance or 0.0
            q2 = db.query(AcctTxn.id, AcctTxn.amount, AcctTxn.date).filter(AcctTxn.source_id == s.id, AcctTxn.expected.is_(False))
            if s.opening_date:
                q2 = q2.filter(AcctTxn.date >= s.opening_date)
            for tid, amt_, _d in sorted(q2.all(), key=lambda x: (x[2], x[0])):
                run += amt_
                balance[tid] = cents(run)
    off, lim = int(p.get("offset") or 0), min(int(p.get("limit") or 300), 5000)
    page = rows[off:off + lim]
    return {"rows_obj": page, "all_rows": rows, "total": len(rows), **tot, "stats": stats, "groups": groups,
            "in_account": {t.id: part(t) for t in rows if only is not None and t.splits},
            "balance": {k: v for k, v in balance.items() if k in {t.id for t in page}}, "has_balance": bool(balance)}


def found_xlsx(db: Session, rows: List[AcctTxn], in_account: Optional[Dict[int, float]] = None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    accounts, sources = A.maps(db)
    wb = Workbook()
    ws = wb.active
    ws.title = "Found"
    in_account = in_account or {}
    ws.append(["Date", "Amount", "Description", "Payee", "Account", "Bank Account", "Note"] + (["In This Account"] if in_account else []))
    for c in ws[1]:
        c.font = Font(bold=True)
    for t in rows:
        a = accounts.get(t.account_id)
        acc = A.path(a, accounts) if a else (", ".join(f"{A.path(accounts[x.account_id], accounts)} {x.amount:.2f}" for x in t.splits if x.account_id in accounts) or "Not filed")
        ws.append([t.date.date(), t.amount, t.bank_description or t.description, t.payee or "", acc, sources[t.source_id].name if t.source_id in sources else "", t.note or ""]
                  + ([in_account.get(t.id, t.amount)] if in_account else []))
        ws.cell(ws._current_row, 2).number_format = "#,##0.00;[Red]-#,##0.00"
    ws.append([])
    ws.append(["Total", cents(sum(t.amount for t in rows)), f"{len(rows)} line(s)"] + (["", "", "", "", cents(sum(in_account.get(t.id, t.amount) for t in rows))] if in_account else []))
    ws.cell(ws._current_row, 1).font = Font(bold=True)
    for col, w in zip("ABCDEFG", (12, 13, 70, 26, 28, 20, 30)):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ================================================================ comments
def comment_dict(c: AcctComment) -> dict:
    return {"id": c.id, "section": c.section, "period_key": c.period_key, "ref_type": c.ref_type, "ref_id": c.ref_id, "text": c.text,
            "created_by": c.created_by, "created_at": c.created_at, "updated_by": c.updated_by, "updated_at": c.updated_at}


def comments(db: Session, section: Optional[str] = None, pkey: Optional[str] = None, all_periods: bool = False) -> List[dict]:
    q = db.query(AcctComment)
    if section:
        q = q.filter(AcctComment.section == section)
    if pkey and not all_periods:
        q = q.filter(or_(AcctComment.period_key == pkey, AcctComment.period_key.is_(None)))
    return [comment_dict(c) for c in q.order_by(AcctComment.created_at).all()]


# ================================================================ send back to Review
def send_back(db: Session, ids: List[int], reason: Optional[str], by: str) -> int:
    n = 0
    for t in db.query(AcctTxn).options(selectinload(AcctTxn.splits)).filter(AcctTxn.id.in_(ids)).all():
        if not t.account_id and not t.splits:
            continue
        was = {"account_id": t.account_id} if t.account_id else {"splits": [{"account_id": x.account_id, "amount": x.amount, "note": x.note} for x in t.splits]}
        t.sent_back_from = json.dumps(was)
        t.splits.clear()
        t.account_id, t.how, t.rule_id, t.reviewed = None, None, None, True
        t.sent_back, t.sent_back_reason, t.sent_back_by, t.sent_back_at = True, (reason or "").strip() or None, by, datetime.utcnow()
        t.flag = t.flag_info = None
        t.updated_by = by
        n += 1
    db.commit()
    return n


def flag_lines(db: Session, ids: List[int], reason: Optional[str], by: str, kind: str = "reconcile") -> int:
    """Keep where they're filed but put them in Review with a red flag and the reason (Reconcile: "QuickBooks says ...")."""
    n = 0
    for t in db.query(AcctTxn).filter(AcctTxn.id.in_(ids)).all():
        t.flag, t.flag_info, t.updated_by = kind, json.dumps({"reason": (reason or "").strip() or "Flagged for a look", "by": by}), by
        n += 1
    db.flush()
    return n


def clear_flags(db: Session, ids: List[int], by: str) -> int:
    """"Looks right": out of Review (the line keeps its account; a rule's check is done)."""
    n = 0
    for t in db.query(AcctTxn).options(selectinload(AcctTxn.splits)).filter(AcctTxn.id.in_(ids)).all():
        if t.flag == "conflict" and not t.account_id and not t.splits:
            continue  # rules disagree and it isn't filed: pick an account (or a rule) first
        if t.flag:
            t.flag = t.flag_info = None
            t.reviewed, t.updated_by = True, by
            if t.how in ("auto", "rule") and t.account_id:
                t.how = "person" if t.how == "auto" else t.how
            n += 1
    db.flush()
    return n


def undo_send_back(db: Session, ids: List[int], by: str) -> int:
    n = 0
    for t in db.query(AcctTxn).options(selectinload(AcctTxn.splits)).filter(AcctTxn.id.in_(ids)).all():
        if not t.sent_back_from or t.account_id or t.splits:
            continue
        was = json.loads(t.sent_back_from)
        if was.get("account_id"):
            t.account_id = was["account_id"]
        for x in was.get("splits") or []:
            t.splits.append(AcctSplit(account_id=x["account_id"], amount=x["amount"], note=x.get("note")))
        t.how, t.reviewed, t.sent_back, t.updated_by = "person", True, None, by
        n += 1
    db.commit()
    return n


# ================================================================ Ignored (duplicates / estimates)
def ignored(db: Session, import_id: Optional[int] = None, status: Optional[str] = None) -> dict:
    accounts, sources = A.maps(db)
    q = db.query(AcctIgnored)
    if import_id:
        q = q.filter(AcctIgnored.import_id == import_id)
    if status:
        q = q.filter(AcctIgnored.status == status)
    rows = q.order_by(AcctIgnored.status.desc(), AcctIgnored.date.desc()).all()
    dup_ids = {r.dup_of_txn_id for r in rows if r.dup_of_txn_id}
    dups = {t.id: A.txn_dict(t, accounts, sources) for t in db.query(AcctTxn).options(selectinload(AcctTxn.splits)).filter(AcctTxn.id.in_(dup_ids)).all()} if dup_ids else {}
    out = [{"id": r.id, "import_id": r.import_id, "source": sources[r.source_id].name if r.source_id in sources else "", "date": r.date.date().isoformat() if r.date else None,
            "amount": r.amount, "description": r.description, "extra": json.loads(r.extra) if r.extra else {}, "reason": r.reason, "status": r.status,
            "dup_of": dups.get(r.dup_of_txn_id), "restored_txn_id": r.restored_txn_id, "decided_by": r.decided_by, "created_at": r.created_at} for r in rows]
    c = Counter(r["status"] for r in out)
    return {"rows": out, "possible": c["possible"], "ignored": c["ignored"], "restored": c["restored"],
            "total_ignored": cents(sum(r["amount"] or 0 for r in out if r["status"] == "ignored"))}


def decide_ignored(db: Session, ids: List[int], action: str, by: str) -> int:
    """ignore: it IS a duplicate (stays out) | keep / restore: it isn't -- it becomes a normal line, in Review."""
    n = 0
    for r in db.query(AcctIgnored).filter(AcctIgnored.id.in_(ids)).all():
        if action == "ignore":
            if r.status == "possible":
                r.status, r.decided_by = "ignored", by
                n += 1
        elif action in ("keep", "restore"):
            if r.status == "restored":
                continue
            src = db.get(AcctSource, r.source_id) or A.book_source(db)
            t = AcctTxn(source_id=src.id, date=r.date, amount=r.amount or 0.0, description=r.description or "",
                        payee=acct_mapper.payee_label(r.description or ""), origin="import", import_id=r.import_id,
                        dedupe=f"{r.dedupe or ''}|kept{r.id}", created_by=by)
            db.add(t)
            db.flush()
            A.autofill(db, [t])
            r.status, r.restored_txn_id, r.decided_by = "restored", t.id, by
            n += 1
        else:
            raise HTTPException(status_code=400, detail="ignore, keep or restore")
    db.commit()
    return n


# ================================================================ Check My Books
CHECKS = {
    "odd_one_out": ("Odd One Out", "high"), "rule_says": ("A Rule Says Another Account", "high"),
    "rules_disagree": ("Two Rules Disagree", "low"), "direction": ("Money Going The Wrong Way", "med"),
    "unusual": ("Unusual Amount", "med"), "duplicate": ("Possible Duplicate", "high"), "split": ("Split Doesn't Add Up", "high"),
    "inactive": ("In A Switched-Off Account", "low"), "loan_negative": ("Loan Repaid More Than Borrowed", "med"),
    "transfer_net": ("Transfers Don't Net Out", "low"), "book_no_note": ("Book Entry Without A Note", "low"),
    "auto_old": ("AT-HUB's Pick Not Looked At (30+ Days)", "low"),
    "qb_split": ("Split In QuickBooks, Not Here", "high"), "two_partners": ("Large Payment Naming Another Partner -- Split?", "low"),
    "flagged": ("Waiting In Review With A Flag", "low"),
}


def check_books(db: Session, start: Optional[date], end: Optional[date]) -> dict:
    accounts, sources = A.maps(db)
    oks = {(o.txn_id, o.check) for o in db.query(AcctCheckOk).all()}
    every = db.query(AcctTxn).options(selectinload(AcctTxn.splits)).all()
    inp = [t for t in every if (not start or t.date.date() >= start) and (not end or t.date.date() <= end)]
    findings = []
    inp_ids = {t.id for t in inp}

    def add(check, t, why, suggest=None, also=None, ref=None):
        if t is not None and (t.id, check) in oks:
            return
        label, sev = CHECKS[check]
        findings.append({"key": f"{check}:{t.id if t is not None else ref}", "check": check, "label": label, "severity": sev, "why": why,
                         "txn": A.txn_dict(t, accounts, sources) if t is not None else None,
                         "also": [A.txn_dict(x, accounts, sources) for x in (also or [])],
                         "suggest": {"account_id": suggest, "account": A.path(accounts[suggest], accounts)} if suggest in accounts else None})

    # odd one out: a payee whose lines nearly all go to one account, this one went elsewhere
    groups = defaultdict(list)
    for t in every:
        if t.account_id and not t.splits:
            groups[(acct_mapper.payee(t.bank_description or t.description), t.amount >= 0)].append(t)
    for _k, ts in groups.items():
        if len(ts) < 4:
            continue
        c = Counter(t.account_id for t in ts)
        top, n = c.most_common(1)[0]
        if n / len(ts) < 0.8:
            continue
        for t in ts:
            if t.account_id != top and c[t.account_id] <= max(1, len(ts) // 10) and t.id in inp_ids and top in accounts:
                add("odd_one_out", t, f"{n} of {len(ts)} lines from {acct_mapper.payee_label(t.bank_description or t.description)} are in "
                                      f"{A.path(accounts[top], accounts)}; this one is in {A.path(accounts[t.account_id], accounts) if t.account_id in accounts else '?'}", top)
    # rules
    specs = [s for s in A.rule_specs(db) if s["enabled"] and s["action"] != "suggest"]
    for t in inp:
        if not t.account_id or t.splits:
            continue
        line = acct_rules.line_of(t)
        hits = [s for s in specs if acct_rules.matches(s, line)]
        if hits and hits[0]["account_id"] and t.account_id != hits[0]["account_id"] and not hits[0]["splits"] and t.rule_id != hits[0]["id"]:
            add("rule_says", t, f"Rule \"{hits[0]['name'] or acct_rules.describe(hits[0])}\" puts lines like this in "
                                f"{A.path(accounts[hits[0]['account_id']], accounts) if hits[0]['account_id'] in accounts else '?'}", hits[0]["account_id"])
        if len({s["account_id"] for s in hits}) > 1:
            add("rules_disagree", t, "These rules all match it and point to different accounts (the first one in the order wins): "
                + "; ".join(f"{s['name'] or acct_rules.describe(s)}" for s in hits[:3]))
    # direction / switched off / splits / book entries / old auto picks
    today = A.today()
    for t in inp:
        a = accounts.get(t.account_id)
        if a:
            if a.kind == "income" and t.amount < 0:
                add("direction", t, f"Money out in {A.path(a, accounts)}, an income account")
            elif a.kind in ("cogs", "expense") and t.amount > 0:
                add("direction", t, f"Money in in {A.path(a, accounts)}, a {A.KINDS[a.kind][0].lower()} account -- a refund? Then it's fine")
            if not a.is_active:
                add("inactive", t, f"{A.path(a, accounts)} is switched off")
        if t.splits and abs(sum(x.amount for x in t.splits) - t.amount) > 0.01:
            add("split", t, f"The parts add up to {sum(x.amount for x in t.splits):,.2f}, the line is {t.amount:,.2f}")
        if t.source_id in sources and sources[t.source_id].kind == "book" and not (t.note or "").strip():
            add("book_no_note", t, "A book entry (not on a statement) with no note to say why")
        if t.how in ("auto", "rule") and not t.reviewed and t.created_at and (today - t.created_at.date()).days > 30:
            add("auto_old", t, "AT-HUB filed it by itself over a month ago and nobody has looked")
        if t.flag == "qb_split" and not t.splits:
            add("qb_split", t, "QuickBooks has this line split over several accounts (its register shows no single account) -- split it the same way")
        elif t.flag and t.flag != "qb_split":
            info = json.loads(t.flag_info) if t.flag_info else {}
            add("flagged", t, info.get("reason") or "Flagged")
    # a partner account's line whose wording names another partner too (Zelle to Niraj "MAALIKS SHARE" ...): split?
    names = defaultdict(set)  # partner -> words that name them (their name, their accounts' names)
    for a in accounts.values():
        pn = _partner_of(a, accounts) if a.kind in ("owner", "loan") else None
        if pn:
            names[pn.lower()] |= {w for w in re.findall(r"[a-z]{4,}", (pn + " " + a.name).lower())}
    for t in inp:
        if t.splits or t.account_id not in accounts or accounts[t.account_id].kind != "owner":
            continue
        mine = (_partner_of(accounts[t.account_id], accounts) or "").lower()
        text_ = set(re.findall(r"[a-z]{4,}", ((t.bank_description or t.description) + " " + (t.note or "")).lower()))
        others = [p_ for p_, ws in names.items() if p_ != mine and ws & text_]
        if others and mine and abs(t.amount) >= 20000:  # a partner's share is often paid through another partner -- only the big ones
            add("two_partners", t, f"In {A.path(accounts[t.account_id], accounts)} but the wording also names {', '.join(o.title() for o in others)} -- one payment for two partners?")
    # unusual amounts per account
    per_acc = defaultdict(list)
    for t in every:
        if t.account_id and not t.splits:
            per_acc[t.account_id].append(abs(t.amount))
    med = {k: statistics.median(v) for k, v in per_acc.items() if len(v) >= 5}
    for t in inp:
        if t.account_id in med and abs(t.amount) > 500 and abs(t.amount) > 5 * med[t.account_id]:
            add("unusual", t, f"{abs(t.amount):,.2f} is over 5x the usual {med[t.account_id]:,.2f} for {A.path(accounts[t.account_id], accounts)}")
    # possible duplicates already in the books
    by_amt = defaultdict(list)
    for t in inp:
        by_amt[(t.source_id, round(t.amount, 2))].append(t)
    for _k, ts in by_amt.items():
        ts.sort(key=lambda x: (x.date, x.id))
        for i, t in enumerate(ts):
            for u in ts[i + 1:]:
                if (u.date - t.date).days > 3:
                    break
                if A.same_line(t.bank_description or t.description, u.bank_description or u.description, (u.date - t.date).days):
                    add("duplicate", u, f"Looks like the {t.date.date().isoformat()} line of the same amount", also=[t])
    # account-level: loans repaid more than borrowed, transfers that don't net out
    for l_ in A.loans(db, end):
        if l_["owed"] < -0.5:
            add("loan_negative", None, f"{l_['name']}: repaid {abs(l_['owed']):,.2f} more than borrowed (owed below zero)", ref=f"loan{l_['account_id']}")
    tnet = defaultdict(float)
    for t in inp:
        a = accounts.get(t.account_id)
        if a and a.kind == "transfer":
            tnet[a.id] += t.amount
    for aid, v in tnet.items():
        if abs(v) > 1:
            add("transfer_net", None, f"{A.path(accounts[aid], accounts)} nets to {v:,.2f} in this period (transfers usually come back to ~0)", ref=f"transfer{aid}")
    order = {"high": 0, "med": 1, "low": 2}
    findings.sort(key=lambda f: (order[f["severity"]], f["check"], f["txn"]["date"] if f["txn"] else ""))
    counts = Counter(f["check"] for f in findings)
    return {"findings": findings, "counts": [{"check": k, "label": CHECKS[k][0], "severity": CHECKS[k][1], "count": counts[k]} for k in CHECKS if counts[k]],
            "total": len(findings)}


# ================================================================ rules: conflicts, health, combine
def _rule_lines(db: Session):
    return [(t, acct_rules.line_of(t)) for t in db.query(AcctTxn).options(selectinload(AcctTxn.splits)).all()]


def rule_conflicts(db: Session) -> List[dict]:
    """Pairs of switched-on rules that match the same lines and say different things -- those lines wait in Review
    (unless one of the two is marked to win)."""
    accounts, _ = A.maps(db)
    specs = [x for x in A.rule_specs(db) if x["enabled"]]
    pairs = {}
    for t, line in _rule_lines(db):
        hits = [x for x in specs if acct_rules.matches(x, line)]
        for a, b in itertools.combinations(hits, 2):
            if acct_rules.outcome(a) == acct_rules.outcome(b):
                continue
            k = (a["id"], b["id"])
            g = pairs.setdefault(k, {"a": a, "b": b, "ids": [], "total": 0.0, "sample": []})
            g["ids"].append(t.id)
            g["total"] += t.amount
            if len(g["sample"]) < 3:
                g["sample"].append({"date": t.date.date().isoformat(), "amount": t.amount, "description": t.bank_description or t.description})
    def side(x):
        return {"id": x["id"], "name": x["name"] or acct_rules.describe(x), "wins": x["wins"],
                "account": A.path(accounts[x["account_id"]], accounts) if x["account_id"] in accounts else "?", "split": bool(x["splits"])}
    out = []
    for g in pairs.values():
        win = next((x for x in (g["a"], g["b"]) if x["wins"]), None)
        out.append({"a": side(g["a"]), "b": side(g["b"]), "lines": len(g["ids"]), "ids": g["ids"], "total": cents(g["total"]), "sample": g["sample"],
                    "result": f"{win['name'] or acct_rules.describe(win)} wins" if win else "they wait in Review"})
    out.sort(key=lambda x: -x["lines"])
    return out


def rule_overlaps(db: Session, spec_: dict) -> List[dict]:
    """While a rule is being built: the other switched-on rules that also match its lines and say something else."""
    accounts, _ = A.maps(db)
    others = [x for x in A.rule_specs(db) if x["enabled"] and x["id"] != spec_.get("id")]
    hits = defaultdict(list)
    for t, line in _rule_lines(db):
        if not acct_rules.matches(spec_, line):
            continue
        for x in others:
            if acct_rules.outcome(x) != acct_rules.outcome(spec_) and acct_rules.matches(x, line):
                hits[x["id"]].append(t)
    by_id = {x["id"]: x for x in others}
    return [{"id": k, "name": by_id[k]["name"] or acct_rules.describe(by_id[k]), "wins": by_id[k]["wins"],
             "account": A.path(accounts[by_id[k]["account_id"]], accounts) if by_id[k]["account_id"] in accounts else "?",
             "lines": len(v), "sample": [t.bank_description or t.description for t in v[:2]]} for k, v in sorted(hits.items(), key=lambda x: -len(x[1]))]


def rule_health(db: Session) -> List[dict]:
    """Per rule: lines it matches (all time / last 12 months), lines it filed, lines a person moved elsewhere (a wrong
    rule?), another rule that already covers every line it matches (redundant), how many lines it fights over."""
    accounts, _ = A.maps(db)
    specs = A.rule_specs(db)
    lines = _rule_lines(db)
    since = A.today() - timedelta(days=365)
    matched = {x["id"]: set() for x in specs}
    recent_ = defaultdict(int)
    for t, line in lines:
        for x in specs:
            if acct_rules.matches(x, line):
                matched[x["id"]].add(t.id)
                if t.date.date() >= since:
                    recent_[x["id"]] += 1
    tx = {t.id: t for t, _l in lines}
    filed = Counter(t.rule_id for t, _l in lines if t.rule_id)
    conflicts = Counter()
    for c in rule_conflicts(db):
        conflicts[c["a"]["id"]] += c["lines"]
        conflicts[c["b"]["id"]] += c["lines"]
    out = []
    for x in specs:
        ids = matched[x["id"]]
        moved = [i for i in ids if tx[i].how == "person" and not x["splits"] and (tx[i].account_id or tx[i].splits) and tx[i].account_id != x["account_id"]]
        by_person = [i for i in ids if tx[i].how == "person"]
        covered = next((y for y in specs if y["id"] != x["id"] and y["enabled"] and ids and ids <= matched[y["id"]]
                        and acct_rules.outcome(y) == acct_rules.outcome(x)), None)
        issues = []
        if not x["enabled"]:
            issues.append("switched off")
        if not recent_[x["id"]]:
            issues.append("no lines in the last 12 months" if ids else "matches no line at all")
        if len(moved) >= 2 and len(moved) >= 0.5 * max(1, len(by_person)):
            issues.append(f"a person moved {len(moved)} of its lines elsewhere -- is it right?")
        if covered:
            issues.append(f"rule \"{covered['name'] or acct_rules.describe(covered)}\" already covers every line it matches")
        if conflicts[x["id"]]:
            issues.append(f"disagrees with another rule on {conflicts[x['id']]} line(s)")
        out.append({"id": x["id"], "name": x["name"] or acct_rules.describe(x), "enabled": x["enabled"], "matches": len(ids), "recent": recent_[x["id"]],
                    "filed": filed[x["id"]], "moved_by_person": len(moved), "issues": issues,
                    "account": A.path(accounts[x["account_id"]], accounts) if x["account_id"] in accounts else ""})
    return out


_COMBINE_FIELDS = ("text", "description", "payee")


def _combinable(x: dict) -> bool:
    cs = x["conditions"]
    return len(cs) == 1 and cs[0].get("field") in _COMBINE_FIELDS and cs[0].get("op") in ("contains", "any_of") and isinstance(cs[0].get("value"), str)


def _parts(x: dict) -> List[str]:
    c = x["conditions"][0]
    return [v.strip() for v in c["value"].split(",") if v.strip()] if c["op"] == "any_of" else [c["value"].strip()]


def combine_suggestions(db: Session) -> List[dict]:
    """Rules that do exactly the same thing (same account / split, money in / out, bank accounts, files or suggests) and
    each look for one set of words: one rule with "contains any of" does the job -- with how many lines each way."""
    accounts, _ = A.maps(db)
    groups = defaultdict(list)
    for x in A.rule_specs(db):
        if x["enabled"] and _combinable(x):
            k = (acct_rules.outcome(x), x["sign"], tuple(sorted(x["source_ids"])), x["review"], x["wins"], x["set_payee"], x["set_note"],
                 x["conditions"][0]["field"])
            groups[k].append(x)
    lines = _rule_lines(db)
    out = []
    for k, rs in groups.items():
        if len(rs) < 2:
            continue
        vals, seen = [], set()
        for x in rs:
            for v in _parts(x):
                if v.upper() not in seen:
                    seen.add(v.upper())
                    vals.append(v)
        combined = {**rs[0], "id": None, "name": "", "conditions": [{"field": k[-1], "op": "any_of", "value": ", ".join(vals)}]}
        old = {t.id for t, line in lines if any(acct_rules.matches(x, line) for x in rs)}
        new = {t.id for t, line in lines if acct_rules.matches(combined, line)}
        out.append({"rule_ids": [x["id"] for x in rs], "rules": [x["name"] or acct_rules.describe(x) for x in rs],
                    "account": A.path(accounts[rs[0]["account_id"]], accounts) if rs[0]["account_id"] in accounts else "?",
                    "combined": acct_rules.describe(combined), "conditions": combined["conditions"], "lines_before": len(old), "lines_after": len(new),
                    "same_lines": old == new, "only_before": len(old - new), "only_after": len(new - old)})
    out.sort(key=lambda g: -len(g["rule_ids"]))
    return out


def combine_rules(db: Session, rule_ids: List[int], name: Optional[str], by: str) -> int:
    rows = [db.get(AcctRule, i) for i in rule_ids]
    if len(rows) < 2 or any(r is None for r in rows):
        raise HTTPException(status_code=400, detail="Pick two or more rules")
    specs = [acct_rules.spec(r) for r in rows]
    if not all(_combinable(x) for x in specs) or len({(acct_rules.outcome(x), x["sign"], tuple(sorted(x["source_ids"])), x["conditions"][0]["field"]) for x in specs}) != 1:
        raise HTTPException(status_code=400, detail="These rules don't do the same thing -- they can't be one rule")
    vals, seen = [], set()
    for x in specs:
        for v in _parts(x):
            if v.upper() not in seen:
                seen.add(v.upper())
                vals.append(v)
    first = min(rows, key=lambda r: (r.priority if r.priority is not None else 10 ** 6, r.id))
    c = AcctRule(name=(name or "").strip() or first.name, match_mode="all",
                 conditions=json.dumps([{"field": specs[0]["conditions"][0]["field"], "op": "any_of", "value": ", ".join(vals)}]),
                 source_ids=first.source_ids, splits=first.splits, set_payee=first.set_payee, set_note=first.set_note, action=first.action,
                 review=first.review, wins=first.wins, enabled=True, priority=first.priority, contains=", ".join(vals)[:200], sign=first.sign,
                 account_id=first.account_id, hits=sum(r.hits or 0 for r in rows), created_by=by)
    db.add(c)
    db.flush()
    db.query(AcctTxn).filter(AcctTxn.rule_id.in_(rule_ids)).update({AcctTxn.rule_id: c.id}, synchronize_session=False)
    for r in rows:
        db.delete(r)
    db.flush()
    return c.id


# ================================================================ Partner Accounts / Payees / 1099 / Recurring
def _partner_of(a: AcctAccount, accounts) -> Optional[str]:
    for x in [a] + A.ancestors(a, accounts):
        if x.partner:
            return x.partner
    return None


def partner_accounts(db: Session, start: date, end: date, fy_key: str) -> dict:
    """Per partner: withdrawn (money out on their partner accounts), invested (money in on them), loans they made the
    company (loan accounts tagged with them), net, by fiscal year (all time) and by month (the period)."""
    s = A.settings(db)
    accounts = {a.id: a for a in db.query(AcctAccount).all()}
    names = [p_["name"] for p_ in s["partners"]]
    for a in accounts.values():
        if a.kind in ("owner", "loan"):
            pn = _partner_of(a, accounts) or (a.name if a.kind == "owner" and not a.parent_id else None)
            if pn and pn.strip().lower() not in {n.lower() for n in names}:
                names.append(pn)
    m = s["fy_start_month"]
    pt = A.partners(db, start, end, fy_key)
    share = {x["name"].lower(): x for x in pt["partners"]}
    out = []
    for name in names:
        key = name.strip().lower()
        accs = [a for a in accounts.values() if a.kind in ("owner", "loan") and ((_partner_of(a, accounts) or (a.name if a.kind == "owner" else "")).strip().lower() == key)]
        ids = {a.id for a in accs}
        kind_of = {a.id: a.kind for a in accs}
        lines = [l for l in A.lines(db) if l["account_id"] in ids]
        def tot(ls):
            w = sum(l["amount"] for l in ls if kind_of[l["account_id"]] == "owner" and l["amount"] < 0)
            i = sum(l["amount"] for l in ls if kind_of[l["account_id"]] == "owner" and l["amount"] > 0)
            lent = sum(l["amount"] for l in ls if kind_of[l["account_id"]] == "loan" and l["amount"] > 0)
            repaid = sum(l["amount"] for l in ls if kind_of[l["account_id"]] == "loan" and l["amount"] < 0)
            return {"withdrawn": cents(w), "invested": cents(i), "net": cents(w + i), "lent": cents(lent), "repaid": cents(repaid), "loan_owed": cents(lent + repaid), "count": len(ls)}
        period = [l for l in lines if start <= l["date"] <= end]
        by_fy = defaultdict(list)
        for l in lines:
            by_fy[A.fy_label(A.fy_start(l["date"], m), m)].append(l)
        by_month = defaultdict(list)
        for l in period:
            by_month[l["date"].strftime("%Y-%m")].append(l)
        sh = share.get(key)
        out.append({"name": name, "account_ids": sorted(ids), "accounts": [A.path(a, accounts) for a in accs],
                    "period": tot(period), "all_time": tot(lines),
                    "by_fy": [{"fy": k, **tot(v)} for k, v in sorted(by_fy.items())],
                    "by_month": [{"month": k, **tot(v)} for k, v in sorted(by_month.items())],
                    "share": sh["share"] if sh else None, "share_pct": sh["share_pct"] if sh else None,
                    "remaining": sh["remaining"] if sh else None, "adjust": sh["adjust"] if sh else 0})
    return {"partners": out, "to_share": pt["total"], "year": fy_key}


def partner_lines(db: Session, name: str, start: Optional[date], end: Optional[date], what: str = "net") -> dict:
    """The lines behind one figure of Partner Accounts (a month / a year / the period): withdrawn, invested, lent,
    repaid or all -- a split line by the partner's part."""
    accounts, sources = A.maps(db)
    key = name.strip().lower()
    accs = {a.id: a for a in accounts.values() if a.kind in ("owner", "loan")
            and ((_partner_of(a, accounts) or (a.name if a.kind == "owner" else "")).strip().lower() == key)}
    def keep(kind, amt):
        return {"withdrawn": kind == "owner" and amt < 0, "invested": kind == "owner" and amt > 0, "lent": kind == "loan" and amt > 0,
                "repaid": kind == "loan" and amt < 0, "net": kind == "owner", "loans": kind == "loan"}.get(what, True)
    parts = defaultdict(float)
    for l in A.lines(db, start, end):
        a = accs.get(l["account_id"])
        if a and keep(a.kind, l["amount"]):
            parts[(l["txn_id"], l["account_id"])] += l["amount"]
    ids = {k[0] for k in parts}
    tx = {t.id: t for t in db.query(AcctTxn).options(selectinload(AcctTxn.splits)).filter(AcctTxn.id.in_(ids)).all()} if ids else {}
    rows = []
    for (tid, aid), amt in parts.items():
        t = tx.get(tid)
        if t:
            d = A.txn_dict(t, accounts, sources)
            d["part"], d["part_account"] = cents(amt), A.path(accounts[aid], accounts)
            rows.append(d)
    rows.sort(key=lambda d: d["date"])
    return {"rows": rows, "total": cents(sum(r["part"] for r in rows)), "count": len(rows)}


def payees(db: Session, start: date, end: date) -> List[dict]:
    accounts, _ = A.maps(db)
    g = defaultdict(lambda: {"count": 0, "in": 0.0, "out": 0.0, "last": None, "first": None, "accounts": Counter(), "months": set()})
    for t in db.query(AcctTxn).filter(AcctTxn.date >= A._dt(start), AcctTxn.date <= A._dt(end)).all():
        a = accounts.get(t.account_id)
        if a and a.kind == "pending":
            continue
        k = t.payee or acct_mapper.payee_label(t.bank_description or t.description)
        x = g[k]
        x["count"] += 1
        x["in" if t.amount > 0 else "out"] += t.amount
        d = t.date.date()
        x["last"] = max(x["last"], d) if x["last"] else d
        x["first"] = min(x["first"], d) if x["first"] else d
        x["months"].add(d.strftime("%Y-%m"))
        x["accounts"][A.path(a, accounts) if a else "Not filed"] += 1
    out = [{"payee": k, "count": v["count"], "in": cents(v["in"]), "out": cents(v["out"]), "net": cents(v["in"] + v["out"]),
            "first": v["first"].isoformat(), "last": v["last"].isoformat(), "months": len(v["months"]),
            "accounts": [{"account": a, "count": c} for a, c in v["accounts"].most_common(3)]} for k, v in g.items()]
    out.sort(key=lambda x: -abs(x["net"]))
    return out


TAX_KEY = "acct_1099"


def tax1099(db: Session, year: int, threshold: float = 600.0) -> dict:
    """People / businesses paid >= $600 in a calendar year (money out filed in cost of goods / expenses / other) -- the
    1099 list for the CPA. Each can be marked needs a 1099 / doesn't (a card company, a corporation)."""
    accounts, _ = A.maps(db)
    marks = _setting(db, TAX_KEY, {}).get(str(year), {})
    g = defaultdict(lambda: {"total": 0.0, "count": 0, "accounts": Counter(), "first": None, "last": None})
    for t in db.query(AcctTxn).options(selectinload(AcctTxn.splits)).filter(AcctTxn.date >= datetime(year, 1, 1), AcctTxn.date < datetime(year + 1, 1, 1), AcctTxn.amount < 0).all():
        parts = [(x.account_id, x.amount) for x in t.splits] or [(t.account_id, t.amount)]
        for aid, amt in parts:
            a = accounts.get(aid)
            if not a or a.kind not in ("cogs", "expense", "other"):
                continue
            k = t.payee or acct_mapper.payee_label(t.bank_description or t.description)
            x = g[k]
            x["total"] += -amt
            x["count"] += 1
            x["accounts"][A.path(a, accounts)] += 1
            d = t.date.date()
            x["first"] = min(x["first"], d) if x["first"] else d
            x["last"] = max(x["last"], d) if x["last"] else d
    rows = [{"payee": k, "total": cents(v["total"]), "count": v["count"], "first": v["first"].isoformat(), "last": v["last"].isoformat(),
             "accounts": [a for a, _c in v["accounts"].most_common(3)], "mark": marks.get(k)} for k, v in g.items() if v["total"] >= threshold - 0.005]
    rows.sort(key=lambda x: -x["total"])
    return {"year": year, "threshold": threshold, "rows": rows, "needs": cents(sum(r["total"] for r in rows if r["mark"] == "yes"))}


def mark1099(db: Session, year: int, payee: str, mark: Optional[str], by: str) -> None:
    allm = _setting(db, TAX_KEY, {})
    y = allm.setdefault(str(year), {})
    if mark in ("yes", "no"):
        y[payee] = mark
    else:
        y.pop(payee, None)
    _save_setting(db, TAX_KEY, allm, by)
    db.commit()


def recurring(db: Session) -> List[dict]:
    """Payments that keep coming back (rent, loan payments, payroll, card payments): how often, how much, when the next
    one is due, and ones that look missed."""
    accounts, _ = A.maps(db)
    today = A.today()
    since = today - timedelta(days=460)
    g = defaultdict(list)
    for t in db.query(AcctTxn).filter(AcctTxn.date >= A._dt(since), AcctTxn.expected.is_(False)).all():
        g[(acct_mapper.payee(t.bank_description or t.description), t.amount >= 0)].append(t)
    out = []
    for (k, money_in), ts in g.items():
        if len(ts) < 3 or k.startswith("text:") and len(k) < 10:
            continue
        ts.sort(key=lambda t: t.date)
        gaps = [(b.date - a.date).days for a, b in zip(ts, ts[1:]) if (b.date - a.date).days > 0]
        if len(gaps) < 2:
            continue
        gap = statistics.median(gaps)
        how = "weekly" if 5 <= gap <= 9 else "every 2 weeks" if 12 <= gap <= 16 else "monthly" if 26 <= gap <= 35 else "quarterly" if 84 <= gap <= 96 else None
        if not how:
            continue
        regular = sum(1 for x in gaps if abs(x - gap) <= max(3, gap * 0.25)) / len(gaps)
        if regular < 0.6:
            continue
        amts = [abs(t.amount) for t in ts]
        medamt = statistics.median(amts)
        last = ts[-1].date.date()
        nxt = last + timedelta(days=round(gap))
        acc = Counter(t.account_id for t in ts if t.account_id).most_common(1)
        out.append({"key": k, "payee": ts[-1].payee or acct_mapper.payee_label(ts[-1].bank_description or ts[-1].description),
                    "money": "in" if money_in else "out", "how_often": how, "every_days": round(gap), "amount": cents(medamt),
                    "same_amount": (max(amts) - min(amts)) <= max(1.0, medamt * 0.1), "times": len(ts), "last": last.isoformat(),
                    "next": nxt.isoformat(), "missed": today > nxt + timedelta(days=max(5, round(gap * 0.3))),
                    "account": A.path(accounts[acc[0][0]], accounts) if acc and acc[0][0] in accounts else None,
                    "per_month": cents(medamt * 30.4 / gap * (1 if money_in else -1))})
    out.sort(key=lambda x: x["next"])
    return out


# ================================================================ scratch sheets (Projections, Loans quick calculations)
def sheet_dict(s: AcctSheet) -> dict:
    return {"id": s.id, "kind": s.kind, "name": s.name, "owner": s.owner, "doc": json.loads(s.doc or "{}"), "created_by": s.created_by,
            "created_at": s.created_at, "updated_by": s.updated_by, "updated_at": s.updated_at}


def books_snapshot(db: Session, months: int = 6) -> dict:
    """Today's figures from the books, for "Fill from the books" on a projection."""
    sources = [A.source_dict(db, s) for s in db.query(AcctSource).filter(AcctSource.kind != "book", AcctSource.is_active.is_(True)).all()]
    oi = A.open_items(db)
    groups = defaultdict(float)
    for r in oi["items"]:
        if r["status"] == "open":
            groups[(r["side"], r["group"])] += r["amount"]
    today = A.today()
    end = today.replace(day=1) - timedelta(days=1)
    start = date(end.year, end.month, 1)
    for _ in range(months - 1):
        start = date(start.year - (start.month == 1), (start.month - 2) % 12 + 1, 1)
    p = A.pnl(db, start, end, "month")
    k = {x["kind"]: x["total"] for x in p["sections"]}
    return {"banks": [{"name": s["name"], "balance": s["balance"] if s["opening_date"] else None, "has_opening": bool(s["opening_date"])} for s in sources],
            "loans": [{"name": l_["name"], "owed": l_["owed"]} for l_ in A.loans(db) if abs(l_["owed"]) > 0.004],
            "ar": [{"name": g, "amount": cents(v)} for (side, g), v in groups.items() if side == "ar"],
            "ap": [{"name": g, "amount": cents(v)} for (side, g), v in groups.items() if side == "ap"],
            "monthly": {"months": months, "from": start.isoformat(), "to": end.isoformat(), "income": cents(k.get("income", 0) / months),
                        "cogs": cents(k.get("cogs", 0) / months), "expense": cents(k.get("expense", 0) / months)},
            "recurring": [r for r in recurring(db) if r["money"] == "out"][:20],
            "partners": A.settings(db)["partners"]}


_REF = re.compile(r"\bR(\d+)\b")


def sheet_xlsx(name: str, doc: dict) -> bytes:
    """A scenario as an Excel sheet: Group | Name | Amount (formulas kept: R3 -> C4) | Month | Note."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = (name or "Projection")[:31]
    ws.append(["Group", "Name", "Amount", "Month", "Note"])
    for c in ws[1]:
        c.font = Font(bold=True)
    def to_xl(f: str):
        f = str(f or "")
        if not f.startswith("="):
            try:
                return float(f.replace(",", "").replace("$", ""))
            except ValueError:
                return f
        f = re.sub(r'GROUP\("([^"]+)"\)', lambda m: f'SUMIF(A:A,"{m.group(1)}",C:C)', f, flags=re.I)
        return _REF.sub(lambda m: f"C{int(m.group(1)) + 1}", f)
    for r in doc.get("rows", []):
        ws.append([r.get("group", ""), r.get("name", ""), to_xl(r.get("a", "")), r.get("month", ""), r.get("note", "")])
        ws.cell(ws._current_row, 3).number_format = "#,##0.00;[Red]-#,##0.00"
    for col, w in zip("ABCDE", (22, 34, 16, 12, 40)):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ================================================================ reconcile against QuickBooks
QB_MAP_KEY = "acct_qb_map"
QB_HEADS = {"date": ("date", "transaction date", "txn date"), "amount": ("amount",), "debit": ("debit",), "credit": ("credit",),
            "name": ("name", "payee", "vendor", "customer"), "memo": ("memo/description", "memo", "description"),
            "account": ("account", "account full name", "account name"), "type": ("transaction type", "type"), "num": ("num", "no.", "ref no.")}


def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", (s or "").strip().lower())


def parse_qb(data: Optional[bytes], filename: Optional[str], text: Optional[str]) -> dict:
    """QuickBooks lines from a CSV / Excel export or text pasted from Excel. Understands "Transaction Detail by Account"
    (account names as group rows, "Total for ..." rows) as well as a plain list with an Account column."""
    if data is not None:
        rows = A.read_table(data, filename or "qb.csv")
    else:
        txt = (text or "").strip()
        if not txt:
            raise HTTPException(status_code=400, detail="Paste the QuickBooks lines or pick a file")
        delim = "\t" if "\t" in txt.splitlines()[0] or "\t" in txt else ","
        rows = [r for r in csv.reader(io.StringIO(txt), delimiter=delim) if any((c or "").strip() for c in r)]
    head_i, cols = None, {}
    for i, r in enumerate(rows[:30]):
        cells = [_norm(str(c or "")) for c in r]
        found = {}
        for key, names in QB_HEADS.items():
            for j, c in enumerate(cells):
                if c in names and j not in found.values() and key not in found:
                    found[key] = j
        if "date" in found and ("amount" in found or "debit" in found):
            head_i, cols = i, found
            break
    if head_i is None:
        raise HTTPException(status_code=400, detail="Couldn't find the header row (Date, Amount, Name / Memo, Account) -- copy it with the lines")
    out, current, skipped = [], None, 0
    for r in rows[head_i + 1:]:
        cell = lambda k: (r[cols[k]] if k in cols and cols[k] < len(r) else None)
        d = A.parse_day(cell("date"))
        if cols.get("amount") is not None:
            amt = A.parse_money(cell("amount"))
        else:
            dr, cr = A.parse_money(cell("debit")), A.parse_money(cell("credit"))
            amt = None if dr is None and cr is None else (dr or 0) - (cr or 0)
        first = str(r[0] or "").strip() if r else ""
        if d is None and amt is None:
            if first and not first.lower().startswith("total") and not any(str(c or "").strip() for c in r[1:]):
                current = first  # "Partner Distributions:Anuj" group row
            continue
        if d is None or amt is None:
            if first.lower().startswith("total"):
                continue
            skipped += 1
            continue
        acct = str(cell("account") or "").strip() or current or ""
        out.append({"i": len(out), "date": d.isoformat(), "amount": cents(amt), "name": str(cell("name") or "").strip(), "memo": str(cell("memo") or "").strip(),
                    "account": acct, "type": str(cell("type") or "").strip(), "num": str(cell("num") or "").strip()})
    return {"rows": out, "skipped": skipped, "columns": {k: v for k, v in cols.items()}}


def qb_map(db: Session) -> Dict[str, int]:
    return {k: int(v) for k, v in _setting(db, QB_MAP_KEY, {}).items()}


def guess_map(name: str, accounts: Dict[int, AcctAccount]) -> Optional[int]:
    """'Partner Distributions:Anuj' -> AT-HUB's Anuj; the last part of a QuickBooks name against account names / paths."""
    parts = [_norm(x) for x in re.split(r"[:›>]", name or "") if x.strip()]
    if not parts:
        return None
    by_path = {_norm(A.path(a, accounts)).replace(" › ", ":"): a.id for a in accounts.values()}
    full = ":".join(parts)
    if full in by_path:
        return by_path[full]
    last = parts[-1]
    hits = [a.id for a in accounts.values() if _norm(a.name) == last]
    if len(hits) == 1:
        return hits[0]
    hits = [a.id for a in accounts.values() if a.partner and _norm(a.partner) == last and a.kind == "owner" and not a.parent_id]
    return hits[0] if len(hits) == 1 else None


REC_KEY = "acct_rec_decisions"


def rec_key(r: dict) -> str:
    """A QuickBooks line's identity across checks: date | amount | its words."""
    return f"{r['date']}|{float(r['amount']):.2f}|{_norm((r.get('name') or '') + ' ' + (r.get('memo') or ''))[:60]}"


def rec_decisions(db: Session) -> dict:
    return _setting(db, REC_KEY, {})


def rec_decide(db: Session, items: List[dict], by: str) -> int:
    """Remember what a person said about a QuickBooks / AT-HUB line: athub_right (AT-HUB's account is right), ignore (not
    a line for AT-HUB), ok (an AT-HUB-only line that's fine), pair (this QuickBooks line IS that AT-HUB line); None clears."""
    d = rec_decisions(db)
    n = 0
    for it in items:
        k = it.get("key") or (f"athub:{it['txn_id']}" if it.get("txn_id") and it.get("decision") == "ok" else None)
        if not k:
            continue
        if not it.get("decision"):
            d.pop(k, None)
        else:
            d[k] = {"d": it["decision"], "txn_id": it.get("txn_id"), "by": by, "at": datetime.utcnow().isoformat(timespec="seconds")}
        n += 1
    _save_setting(db, REC_KEY, d, by)
    db.commit()
    return n


def _words(s: str) -> set:
    return set(re.findall(r"[a-z0-9]+", (s or "").lower())) - {"to", "on", "ref", "from", "the", "of", "and"}


def reconcile(db: Session, account_ids: List[int], start: date, end: date, qb_rows: List[dict], mapping: Dict[str, int], tol_days: int = 5) -> dict:
    """Each QuickBooks line against AT-HUB's lines (any account -- that's how a line in the wrong account is found).

    Matched by size of amount + date within tol_days (closest date, most shared words); each match says WHY (same bank
    reference / same start of the text / same day and amount / amount N days apart) and how sure (high / medium / low).
    Buckets: same account, different account, QuickBooks account not mapped, check these matches (weak ones), split in
    QuickBooks (2-4 QuickBooks lines that add up to one AT-HUB line), amount differs, date differs, only in QuickBooks,
    only in AT-HUB (in the accounts picked), decided (a person said so earlier: AT-HUB is right / ignore / OK).
    A line already split in AT-HUB is matched part by part. A person's manual pairs come first."""
    accounts, sources = A.maps(db)
    picked = set()
    for aid in account_ids:
        picked |= {aid, *A.descendants(aid, accounts)}
    pool = (db.query(AcctTxn).options(selectinload(AcctTxn.splits))
            .filter(AcctTxn.date >= A._dt(start - timedelta(days=62)), AcctTxn.date <= A._dt(end + timedelta(days=62))).all())
    pool_by_id = {t.id: t for t in pool}
    decisions = rec_decisions(db)
    guessed = {}
    for nm in {r["account"] for r in qb_rows}:
        if _norm(nm) not in mapping:
            g_ = guess_map(nm, accounts)
            if g_:
                guessed[_norm(nm)] = g_
    eff = {**guessed, **mapping}

    def in_tree(acc_id, root):
        return acc_id == root or root in [x.id for x in A.ancestors(accounts[acc_id], accounts)] if acc_id in accounts else False

    # targets: a whole line, or one part of a split line
    targets = []
    for t in pool:
        targets.append({"key": f"{t.id}", "t": t, "part": None, "amount": t.amount})
        for k, x in enumerate(t.splits):
            targets.append({"key": f"{t.id}:{k}", "t": t, "part": x, "amount": x.amount})
    by_amt = defaultdict(list)
    for g in targets:
        by_amt[round(abs(g["amount"]), 2)].append(g)
    used_t, used_line = set(), set()  # target keys / AT-HUB lines matched as a whole or by parts

    def free(g):
        if g["key"] in used_t:
            return False
        if g["part"] is None:
            return g["t"].id not in used_line and not any(k.startswith(f"{g['t'].id}:") for k in used_t)
        return f"{g['t'].id}" not in used_t

    def why(r, t, days):
        qtext = (r.get("name") or "") + " " + (r.get("memo") or "")
        desc = t.bank_description or t.description
        rr, tr = A._plain_refs(qtext.upper()), A._plain_refs(desc.upper())
        shared = len(_words(qtext) & _words(desc + " " + (t.payee or "") + " " + (t.note or "")))
        if rr and tr and rr & tr:
            return "same bank reference", "high", "same"
        if r.get("memo") and A.starts_same(r["memo"], desc):
            return "same start of the text", "high", None
        if rr and tr and not rr & tr:
            return f"amount{'' if not days else f', {days} day(s) apart'} -- different bank references", "low", "different"
        if days == 0:
            return "same day and amount", ("high" if shared >= 2 else "medium"), None
        if shared >= 2:
            return f"amount + {shared} shared words, {days} day(s) apart", "medium", None
        return f"amount only, {days} day(s) apart", "low", None

    out = {k: [] for k in ("same", "different", "unmapped", "weak", "split", "amount_differs", "date_differs", "qb_only", "athub_only", "decided")}

    def row_of(r):
        rr = dict(r)
        rr["key"] = rec_key(r)
        mapped = eff.get(_norm(r["account"]))
        rr["mapped_account_id"] = mapped
        rr["mapped_account"] = A.path(accounts[mapped], accounts) if mapped in accounts else None
        rr["mapped_by"] = "you" if _norm(r["account"]) in mapping else ("guess" if mapped else None)
        dec = decisions.get(rr["key"])
        rr["decision"] = dec["d"] if dec else None
        return rr

    def place(rr, g, reason, conf, refs, days):
        t = g["t"]
        used_t.add(g["key"])
        if g["part"] is None:
            used_line.add(t.id)
        rr["athub"] = A.txn_dict(t, accounts, sources)
        rr["athub_part"] = ({"account_id": g["part"].account_id, "account": A.path(accounts[g["part"].account_id], accounts) if g["part"].account_id in accounts else "?",
                             "amount": g["part"].amount} if g["part"] is not None else None)
        rr["days"], rr["reason"], rr["confidence"], rr["refs"] = days, reason, conf, refs
        mapped = rr["mapped_account_id"]
        acc_id = g["part"].account_id if g["part"] is not None else t.account_id
        if acc_id:
            agree = in_tree(acc_id, mapped) if mapped else None
        else:
            agree = any(in_tree(x.account_id, mapped) and abs(abs(x.amount) - abs(rr["amount"])) < 0.01 for x in t.splits) if mapped else None
        rr["agree"] = agree
        if rr["decision"] in ("athub_right", "ignore"):
            out["decided"].append(rr)
        elif not mapped:
            out["unmapped"].append(rr)
        elif conf == "low":
            out["weak"].append(rr)
        else:
            out["same" if agree else "different"].append(rr)

    rows = [row_of(r) for r in qb_rows]
    left = []
    for rr in rows:  # a person's pairs first
        dec = decisions.get(rr["key"])
        if dec and dec["d"] == "pair" and dec.get("txn_id") in pool_by_id:
            g = next((g for g in targets if g["t"].id == dec["txn_id"] and g["part"] is None), None)
            if g and free(g):
                place(rr, g, "paired by " + (dec.get("by") or "a person"), "high", None, abs((g["t"].date.date() - A.parse_day(rr["date"])).days))
                continue
        left.append(rr)
    rest = []
    for rr in left:
        d = A.parse_day(rr["date"])
        words = _words((rr.get("name") or "") + " " + (rr.get("memo") or ""))
        cands = []
        for g in by_amt.get(round(abs(rr["amount"]), 2), []):
            if not free(g):
                continue
            t = g["t"]
            days = abs((t.date.date() - d).days)
            if days > tol_days:
                continue
            tw = _words((t.bank_description or t.description) + " " + (t.payee or "") + " " + (t.note or ""))
            in_picked = (t.account_id in picked) or any(x.account_id in picked for x in t.splits)
            reason, conf, refs = why(rr, t, days)
            rank = {"high": 0, "medium": 1, "low": 2}[conf]
            cands.append((rank, -(1 if in_picked else 0), days, -len(words & tw), t.id, g, reason, conf, refs))
        if not cands:
            rest.append(rr)
            continue
        cands.sort(key=lambda x: x[:5])
        c = cands[0]
        place(rr, c[5], c[6], c[7], c[8], c[2])
    # split in QuickBooks: 2-4 unmatched QuickBooks lines that add up to one AT-HUB line
    unmatched = list(rest)
    for t in sorted(pool, key=lambda t: -abs(t.amount)):
        if t.id in used_line or any(k.startswith(f"{t.id}:") for k in used_t) or not (start - timedelta(days=tol_days) <= t.date.date() <= end + timedelta(days=tol_days)):
            continue
        near = [r for r in unmatched if abs((A.parse_day(r["date"]) - t.date.date()).days) <= tol_days and (r["amount"] < 0) == (t.amount < 0) and abs(r["amount"]) < abs(t.amount)]
        if len(near) < 2:
            continue
        near = sorted(near, key=lambda r: -abs(r["amount"]))[:14]
        hit = None
        for n_ in (2, 3, 4):
            for combo in itertools.combinations(near, n_):
                if abs(sum(r["amount"] for r in combo) - t.amount) < 0.01:
                    hit = combo
                    break
            if hit:
                break
        if not hit:
            continue
        for r in hit:
            unmatched.remove(r)
        used_line.add(t.id)
        parts = [{"key": r["key"], "date": r["date"], "amount": r["amount"], "name": r.get("name"), "memo": r.get("memo"), "account": r["account"],
                  "mapped_account_id": r["mapped_account_id"], "mapped_account": r["mapped_account"]} for r in hit]
        same = bool(t.splits) and sorted((x.account_id, round(x.amount, 2)) for x in t.splits) == sorted((p_["mapped_account_id"], round(p_["amount"], 2)) for p_ in parts)
        out["split"].append({"athub": A.txn_dict(t, accounts, sources), "parts": parts, "already_same": same,
                             "can_split": all(p_["mapped_account_id"] for p_ in parts), "date": t.date.date().isoformat(), "amount": t.amount})
    # amount differs / date differs / only in QuickBooks (with the closest AT-HUB lines)
    free_lines = [t for t in pool if t.id not in used_line and not any(k.startswith(f"{t.id}:") for k in used_t)]
    for rr in unmatched:
        d = A.parse_day(rr["date"])
        qtext = (rr.get("name") or "") + " " + (rr.get("memo") or "")
        qrefs = A._plain_refs(qtext.upper())
        def alike(t):
            desc = t.bank_description or t.description
            return bool(qrefs and qrefs & A._plain_refs(desc.upper())) or bool(rr.get("memo") and A.starts_same(rr["memo"], desc))
        if rr["decision"] in ("ignore", "athub_right"):
            out["decided"].append(rr)
            continue
        amt = next((t for t in free_lines if abs((t.date.date() - d).days) <= max(1, tol_days) and abs(abs(t.amount) - abs(rr["amount"])) > 0.005 and alike(t)), None)
        if amt:
            rr["athub"] = A.txn_dict(amt, accounts, sources)
            rr["difference"] = cents(abs(rr["amount"]) - abs(amt.amount))
            out["amount_differs"].append(rr)
            continue
        dd = next((t for t in sorted(free_lines, key=lambda t: abs((t.date.date() - d).days))
                   if tol_days < abs((t.date.date() - d).days) <= 60 and abs(abs(t.amount) - abs(rr["amount"])) < 0.005 and alike(t)), None)
        if dd:
            rr["athub"] = A.txn_dict(dd, accounts, sources)
            rr["days"] = abs((dd.date.date() - d).days)
            out["date_differs"].append(rr)
            continue
        cl = []
        for t in pool:
            days = abs((t.date.date() - d).days)
            same_amt = abs(abs(t.amount) - abs(rr["amount"])) < 0.005
            sim = len(_words(qtext) & _words(t.bank_description or t.description))
            if (same_amt and days <= 60) or (sim >= 2 and days <= 10):
                cl.append(((0 if same_amt else 1), days, -sim, t))
        cl.sort(key=lambda x: x[:3])
        rr["closest"] = [{**A.txn_dict(x[3], accounts, sources), "days": x[1], "same_amount": x[0] == 0, "taken": x[3].id in used_line} for x in cl[:3]]
        out["qb_only"].append(rr)
    for t in free_lines:
        if not (start <= t.date.date() <= end):
            continue
        if t.account_id in picked or any(x.account_id in picked for x in t.splits):
            d_ = A.txn_dict(t, accounts, sources)
            if (decisions.get(f"athub:{t.id}") or {}).get("d") == "ok":
                d_["decision"] = "ok"
                out["decided"].append({"athub_only": True, "athub": d_, "decision": "ok", "date": d_["date"], "amount": t.amount})
                continue
            d_["closest"] = [{"key": r["key"], "date": r["date"], "amount": r["amount"], "name": r.get("name"), "memo": r.get("memo"), "account": r["account"]}
                             for r in sorted((r for r in out["qb_only"] if abs(abs(r["amount"]) - abs(t.amount)) < 0.005 or abs((A.parse_day(r["date"]) - t.date.date()).days) <= 3),
                                             key=lambda r: (abs(abs(r["amount"]) - abs(t.amount)) > 0.005, abs((A.parse_day(r["date"]) - t.date.date()).days)))[:3]]
            out["athub_only"].append(d_)
    # totals per picked account: AT-HUB vs QuickBooks; month by month
    totals = []
    in_period = [t for t in pool if start <= t.date.date() <= end]
    for aid in account_ids:
        if aid not in accounts:
            continue
        tree = {aid, *A.descendants(aid, accounts)}
        at_total = 0.0
        for t in in_period:
            parts = [(x.account_id, x.amount) for x in t.splits] or [(t.account_id, t.amount)]
            at_total += sum(a_ for acc, a_ in parts if acc in tree)
        qb_total = sum(r["amount"] for r in qb_rows if eff.get(_norm(r["account"])) in tree)
        totals.append({"account_id": aid, "account": A.path(accounts[aid], accounts), "athub": cents(at_total), "qb": cents(qb_total),
                       "athub_abs": cents(abs(at_total)), "qb_abs": cents(abs(qb_total)), "difference": cents(abs(at_total) - abs(qb_total))})
    months = defaultdict(lambda: {"qb": 0.0, "athub": 0.0, "qb_n": 0, "athub_n": 0})
    for r in qb_rows:
        if eff.get(_norm(r["account"])) in picked:
            m_ = months[r["date"][:7]]
            m_["qb"] += abs(r["amount"])
            m_["qb_n"] += 1
    for t in in_period:
        parts = [(x.account_id, x.amount) for x in t.splits] or [(t.account_id, t.amount)]
        v = sum(abs(a_) for acc, a_ in parts if acc in picked)
        if v:
            m_ = months[t.date.strftime("%Y-%m")]
            m_["athub"] += v
            m_["athub_n"] += 1
    grid, run, first_bad = [], 0.0, None
    for k in sorted(months):
        m_ = months[k]
        diff = cents(m_["athub"] - m_["qb"])
        run = cents(run + diff)
        if abs(diff) > 0.009 and not first_bad:
            first_bad = k
        grid.append({"month": k, "qb": cents(m_["qb"]), "athub": cents(m_["athub"]), "difference": diff, "running": run, "qb_lines": m_["qb_n"], "athub_lines": m_["athub_n"]})
    names = Counter(r["account"] for r in qb_rows)
    qb_sum = defaultdict(float)
    for r in qb_rows:
        qb_sum[r["account"]] += r["amount"]
    mapping_rows = []
    for n_, c_ in names.most_common():
        aid = eff.get(_norm(n_))
        at_sum = 0.0
        if aid in accounts:
            tree = {aid, *A.descendants(aid, accounts)}
            for t in in_period:
                parts = [(x.account_id, x.amount) for x in t.splits] or [(t.account_id, t.amount)]
                at_sum += sum(a_ for acc, a_ in parts if acc in tree)
        mapping_rows.append({"name": n_, "count": c_, "qb_total": cents(qb_sum[n_]), "account_id": aid,
                             "account": A.path(accounts[aid], accounts) if aid in accounts else None, "athub_total": cents(at_sum) if aid in accounts else None,
                             "difference": cents(abs(at_sum) - abs(qb_sum[n_])) if aid in accounts else None,
                             "by": "you" if _norm(n_) in mapping else ("guess" if aid else None)})
    unmapped_names = Counter(r["account"] for r in rows if not r["mapped_account_id"])
    return {**out, "totals": totals, "months": grid, "first_bad_month": first_bad, "mapping": mapping_rows,
            "counts": {k: len(v) for k, v in out.items()},
            "unmapped_names": [{"name": n_, "count": c_, "guess": guess_map(n_, accounts),
                                "guess_account": A.path(accounts[guess_map(n_, accounts)], accounts) if guess_map(n_, accounts) in accounts else None}
                               for n_, c_ in unmapped_names.most_common()]}


def reconcile_xlsx(res: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    wb.remove(wb.active)
    labels = {"different": "Different Account", "weak": "Check These Matches", "split": "Split In QuickBooks", "amount_differs": "Amount Differs",
              "date_differs": "Date Differs", "unmapped": "Not Mapped", "qb_only": "Only In QuickBooks", "athub_only": "Only In AT-HUB",
              "same": "Same Account", "decided": "Decided"}
    head = ["QB Date", "QB Name", "QB Memo", "QB Account", "QB Amount", "AT-HUB Date", "AT-HUB Description", "AT-HUB Account", "AT-HUB Amount", "Days", "Why Matched", "How Sure"]
    for k, title in labels.items():
        ws = wb.create_sheet(title[:31])
        ws.append(head)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in res.get(k, []):
            if k == "athub_only":
                ws.append(["", "", "", "", "", r["date"], r["description"], r.get("account") or "", r["amount"]])
                continue
            if k == "split":
                for p_ in r["parts"]:
                    ws.append([p_["date"], p_.get("name") or "", p_.get("memo") or "", p_["account"], p_["amount"], r["athub"]["date"], r["athub"]["description"],
                               r["athub"].get("account") or "split", r["athub"]["amount"], "", "parts add up to the AT-HUB line", ""])
                continue
            a = r.get("athub") or {}
            ws.append([r.get("date"), r.get("name") or "", r.get("memo") or "", r.get("account") or "", r.get("amount"), a.get("date", ""), a.get("description", ""),
                       (r.get("athub_part") or {}).get("account") or a.get("account") or "", a.get("amount", ""), r.get("days", ""), r.get("reason", ""), r.get("confidence", "")])
        for col, w in zip("ABCDEFGHIJKL", (11, 22, 40, 28, 12, 11, 50, 28, 12, 6, 34, 9)):
            ws.column_dimensions[col].width = w
    ws = wb.create_sheet("Month By Month")
    ws.append(["Month", "QuickBooks", "AT-HUB", "Difference", "Running"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for m_ in res.get("months", []):
        ws.append([m_["month"], m_["qb"], m_["athub"], m_["difference"], m_["running"]])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def drop_estimate_accounts(db: Session) -> dict:
    """One-off (user, 2026-10-11): the old year-end AR / AP estimate accounts (kind "pending") and their lines go -- they
    were 2024 guesses, already left out of every total. Rules pointing at them go too."""
    accs = db.query(AcctAccount).filter(AcctAccount.kind == "pending").all()
    ids = [a.id for a in accs]
    if not ids:
        return {"accounts": 0, "lines": 0}
    lines = db.query(AcctTxn).filter(AcctTxn.account_id.in_(ids)).all()
    total = cents(sum(t.amount for t in lines))
    for t in lines:
        db.delete(t)
    db.query(AcctSplit).filter(AcctSplit.account_id.in_(ids)).delete(synchronize_session=False)
    from app.models import AcctRule
    db.query(AcctRule).filter(AcctRule.account_id.in_(ids)).delete(synchronize_session=False)
    db.query(AcctAccount).filter(AcctAccount.parent_id.in_(ids)).update({AcctAccount.parent_id: None}, synchronize_session=False)
    for a in accs:
        db.delete(a)
    db.commit()
    return {"accounts": len(ids), "lines": len(lines), "total": total, "names": [a.name for a in accs]}
