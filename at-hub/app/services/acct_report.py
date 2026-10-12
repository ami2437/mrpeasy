"""The Master Report (user, 2026-10-12): one Excel workbook with everything for a period -- Summary, P&L by month, every
transaction, then one tab per master (top-level) account with every line in it or its sub-accounts (subtotal per
sub-account, total at the foot), Partner Accounts by month, Loans, Not Filed, Comments. Split lines: one row per part.
The user's Auto_Analyzed-style "Output Excel" (accounting.export_xlsx) stays as it is."""
import io
import json
import re
from collections import defaultdict
from datetime import date
from typing import Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session, selectinload

from app.models import AcctAccount, AcctComment, AcctTxn
from app.services import accounting as A
from app.services.money import cents

MONEY = "#,##0.00;[Red]-#,##0.00"
BOLD = Font(bold=True)
HEAD_FILL = PatternFill("solid", fgColor="DCE6F1")
SUB_FILL = PatternFill("solid", fgColor="F2F2F2")
TOTAL_FILL = PatternFill("solid", fgColor="FFF2CC")
KIND_ORDER = ["income", "cogs", "expense", "loan", "owner", "transfer", "other", "pending"]


def _sheet_name(name: str, used: set) -> str:
    base = re.sub(r"[\[\]:*?/\\]", " ", name).strip()[:31] or "Account"
    n, out = 2, base
    while out.lower() in used:
        suffix = f" ({n})"
        out = base[:31 - len(suffix)] + suffix
        n += 1
    used.add(out.lower())
    return out


def _head(ws, cols: List[str], widths: List[int]):
    ws.append(cols)
    for c in ws[ws._current_row]:
        c.font, c.fill = BOLD, HEAD_FILL
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(ws._current_row + 1, 1)


def _row(ws, values, money_cols=(), bold=False, fill=None):
    ws.append(values)
    r = ws._current_row
    for c in money_cols:
        ws.cell(r, c).number_format = MONEY
    if bold or fill:
        for c in ws[r]:
            if bold:
                c.font = BOLD
            if fill:
                c.fill = fill


def master_xlsx(db: Session, start: date, end: date, fy_key: str) -> bytes:
    accounts, sources = A.maps(db)

    def top(aid):
        a = accounts.get(aid)
        if not a:
            return None
        up = A.ancestors(a, accounts)
        return up[-1] if up else a
    txns = (db.query(AcctTxn).options(selectinload(AcctTxn.splits)).filter(AcctTxn.date >= A._dt(start), AcctTxn.date <= A._dt(end))
            .order_by(AcctTxn.date, AcctTxn.id).all())
    parts = []  # one row per line, or per part of a split line
    for t in txns:
        for aid, amt, note in ([(x.account_id, x.amount, x.note) for x in t.splits] or [(t.account_id, t.amount, None)]):
            a, m = accounts.get(aid), top(aid)
            parts.append({"t": t, "account_id": aid, "amount": amt, "part_note": note, "account": A.path(a, accounts) if a else "",
                          "master": m, "kind": a.kind if a else None})
    wb = Workbook()
    used = set()
    # ---- Summary
    ws = wb.active
    ws.title = _sheet_name("Summary", used)
    ws.append(["Accounting Master Report"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"{start.isoformat()} to {end.isoformat()}  ({fy_key})"])
    ws.append([f"{len(txns)} transaction(s), {len(parts)} row(s) with split parts"])
    ws.append([])
    _head(ws, ["Master Account", "Type", "Lines", "Money In", "Money Out", "Net"], [36, 22, 9, 16, 16, 16])
    ws.freeze_panes = None
    by_master = defaultdict(lambda: {"n": 0, "in": 0.0, "out": 0.0})
    for p_ in parts:
        k = p_["master"].id if p_["master"] else 0
        x = by_master[k]
        x["n"] += 1
        x["in" if p_["amount"] > 0 else "out"] += p_["amount"]
    masters = sorted([m for m in {p_["master"].id: p_["master"] for p_ in parts if p_["master"]}.values()],
                     key=lambda a: (KIND_ORDER.index(a.kind) if a.kind in KIND_ORDER else 99, a.name.lower()))
    kind_tot = defaultdict(float)
    for m in masters:
        x = by_master[m.id]
        _row(ws, [m.name, A.KINDS.get(m.kind, A.KINDS["other"])[0], x["n"], cents(x["in"]), cents(x["out"]), cents(x["in"] + x["out"])], (4, 5, 6))
        kind_tot[m.kind] += x["in"] + x["out"]
    if by_master[0]["n"]:
        x = by_master[0]
        _row(ws, ["Not Filed", "", x["n"], cents(x["in"]), cents(x["out"]), cents(x["in"] + x["out"])], (4, 5, 6))
    ws.append([])
    inc, cogs, exp = (kind_tot.get(k, 0.0) for k in ("income", "cogs", "expense"))
    for label, v in (("Income", inc), ("Cost Of Goods", cogs), ("Gross Profit", inc + cogs), ("Expenses", exp), ("Profit", inc + cogs + exp)):
        _row(ws, [label, "", "", "", "", cents(v)], (6,), bold=label in ("Gross Profit", "Profit"), fill=TOTAL_FILL if label == "Profit" else None)
    try:
        pt = A.partners(db, start, end, fy_key)
        ws.append([])
        ws.append(["Partners' Split (" + fy_key + ")"])
        ws.cell(ws._current_row, 1).font = BOLD
        _row(ws, ["To Share", "", "", "", "", pt["total"]], (6,), bold=True)
        for p_ in pt["partners"]:
            _row(ws, [p_["name"], f"{p_['share_pct'] * 100:.1f}%", "", p_["share"], p_["taken"], p_["remaining"]], (4, 5, 6))
        ws.append(["", "", "", "Share", "Taken", "Remaining"])
    except Exception:  # the summary never stops the report
        pass
    # ---- P&L by month
    p = A.pnl(db, start, end, "month")
    ws = wb.create_sheet(_sheet_name("P&L", used))
    _head(ws, ["Account"] + p["columns"] + ["Total"], [34] + [12] * len(p["columns"]) + [14])
    n_cols = len(p["columns"]) + 2
    for sec in p["sections"]:
        if not sec["rows"]:
            continue
        _row(ws, [sec["label"]], bold=True, fill=SUB_FILL)
        for r in sec["rows"]:
            _row(ws, ["   " * r["depth"] + r["name"]] + r["cells"] + [r["total"]], range(2, n_cols + 1))
        _row(ws, [f"Total {sec['label']}"] + sec["totals"] + [sec["total"]], range(2, n_cols + 1), bold=True)
    _row(ws, ["Gross Profit"] + p["gross"] + [p["gross_total"]], range(2, n_cols + 1), bold=True)
    _row(ws, ["Profit"] + p["profit"] + [p["profit_total"]], range(2, n_cols + 1), bold=True, fill=TOTAL_FILL)
    if p["unfiled_total"]:
        _row(ws, ["Not Filed"] + p["unfiled"] + [p["unfiled_total"]], range(2, n_cols + 1))
    # ---- every transaction
    ws = wb.create_sheet(_sheet_name("All Transactions", used))
    cols = ["Date", "Bank Account", "Description", "Payee", "Master Account", "Account", "Type", "Amount", "Line Amount", "Split",
            "Filed By", "Note", "QuickBooks Memo", "Flag"]
    _head(ws, cols, [11, 18, 60, 24, 22, 28, 14, 14, 14, 6, 10, 30, 30, 14])
    for p_ in parts:
        t = p_["t"]
        _row(ws, [t.date.date(), sources[t.source_id].name if t.source_id in sources else "", t.bank_description or t.description, t.payee or "",
                  p_["master"].name if p_["master"] else "Not Filed", p_["account"], A.KINDS.get(p_["kind"], ("",))[0] if p_["kind"] else "",
                  cents(p_["amount"]), t.amount, "yes" if t.splits else "", t.how or "", " · ".join(x for x in (t.note, p_["part_note"]) if x),
                  t.qb_memo or "", t.flag or ""], (8, 9))
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ws._current_row}"
    _row(ws, ["Total", "", f"{len(txns)} line(s)", "", "", "", "", cents(sum(p_["amount"] for p_ in parts))], (8,), bold=True, fill=TOTAL_FILL)
    # ---- one tab per master account
    for m in masters:
        ws = wb.create_sheet(_sheet_name(m.name, used))
        mine = [p_ for p_ in parts if p_["master"] and p_["master"].id == m.id]
        ws.append([f"{m.name} -- {A.KINDS.get(m.kind, A.KINDS['other'])[0]}", "", f"{start.isoformat()} to {end.isoformat()}"])
        ws.cell(ws._current_row, 1).font = Font(bold=True, size=13)
        _head(ws, ["Date", "Sub-Account", "Description", "Payee", "Amount", "Line Amount", "Note", "Filed By", "Bank Account"], [11, 26, 60, 24, 14, 14, 30, 10, 18])
        groups = defaultdict(list)
        for p_ in mine:
            groups[p_["account"]].append(p_)
        multi = len(groups) > 1
        for acc_name in sorted(groups, key=lambda k: (k != m.name, k.lower())):
            rows = groups[acc_name]
            if multi:
                _row(ws, [acc_name if acc_name != m.name else f"{m.name} (itself)"], bold=True, fill=SUB_FILL)
            for p_ in rows:
                t = p_["t"]
                _row(ws, [t.date.date(), p_["account"], t.bank_description or t.description, t.payee or "", cents(p_["amount"]),
                          t.amount if t.splits else "", " · ".join(x for x in (t.note, p_["part_note"]) if x), t.how or "",
                          sources[t.source_id].name if t.source_id in sources else ""], (5, 6))
            if multi:
                _row(ws, [f"Subtotal {acc_name}", "", f"{len(rows)} line(s)", "", cents(sum(p_["amount"] for p_ in rows))], (5,), bold=True)
        _row(ws, [f"Total {m.name}", "", f"{len(mine)} line(s)", "", cents(sum(p_["amount"] for p_ in mine))], (5,), bold=True, fill=TOTAL_FILL)
    # ---- Partner Accounts by month
    from app.services import acct_more as M
    pa = M.partner_accounts(db, start, end, fy_key)
    ws = wb.create_sheet(_sheet_name("Partner Accounts", used))
    _head(ws, ["Partner", "Month", "Withdrawn", "Invested", "Net", "Lent To Company", "Repaid To Them", "Lines"], [18, 10, 14, 14, 14, 16, 16, 8])
    for p_ in pa["partners"]:
        for x in p_["by_month"]:
            _row(ws, [p_["name"], x["month"], x["withdrawn"], x["invested"], x["net"], x["lent"], x["repaid"], x["count"]], (3, 4, 5, 6, 7))
        t_ = p_["period"]
        _row(ws, [f"Total {p_['name']}", "", t_["withdrawn"], t_["invested"], t_["net"], t_["lent"], t_["repaid"], t_["count"]], (3, 4, 5, 6, 7), bold=True, fill=SUB_FILL)
    # ---- Loans
    ws = wb.create_sheet(_sheet_name("Loans", used))
    _head(ws, ["Lender", "Owed At Start", "Borrowed", "Repaid", "Still Owed", "Lines"], [30, 14, 14, 14, 14, 8])
    for l_ in A.loans(db, end):
        _row(ws, [l_["name"], l_["opening"], l_["borrowed"], l_["repaid"], l_["owed"], l_["count"]], (2, 3, 4, 5))
    # ---- Not Filed
    ws = wb.create_sheet(_sheet_name("Not Filed", used))
    _head(ws, ["Date", "Description", "Payee", "Amount", "Flag", "Note"], [11, 60, 24, 14, 16, 30])
    for p_ in parts:
        if not p_["account_id"]:
            t = p_["t"]
            _row(ws, [t.date.date(), t.bank_description or t.description, t.payee or "", t.amount, t.flag or "", t.note or ""], (4,))
    # ---- Comments
    from app.services.acct_more import period_key
    ws = wb.create_sheet(_sheet_name("Comments", used))
    _head(ws, ["Section", "Comment", "By", "When", "Period"], [16, 80, 18, 18, 24])
    pk = period_key(start, end)
    for c in db.query(AcctComment).order_by(AcctComment.section, AcctComment.created_at).all():
        if c.period_key in (None, pk):
            _row(ws, [c.section, c.text, c.created_by or "", c.created_at.strftime("%Y-%m-%d %H:%M") if c.created_at else "", c.period_key or "every period"])
            ws.cell(ws._current_row, 2).alignment = Alignment(wrap_text=True)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
