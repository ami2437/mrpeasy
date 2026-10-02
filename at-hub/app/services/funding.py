"""Bulk funding upload for customer invoices -- the AT-HUB version of the portal's
"bulk payment fields" Excel upload (its custom_570/571/572 = disbursement date,
funding amount, discount).

The factor's report lists one row per funded invoice. Each matched invoice gets the
three funding fields set and, optionally, two payments recorded against it: the funding
amount and the discount (so funding + discount clears the invoice). An upload is kept as
a FundingImport so it can be rolled back as a whole.
"""
import csv
import io
import json
from datetime import datetime, date
from typing import Any, Dict, List, Optional

from fastapi import HTTPException
from sqlalchemy import func
from sqlalchemy.orm import Session

from app.models import Invoice, InvoicePayment, FundingImport

# Same columns the portal requires; "Item Type" is optional (only rows typed "invoice" are used).
REQUIRED_COLUMNS = {
    "item number": "Item Number",
    "disbursement date": "Disbursement Date",
    "funding amount": "Funding Amount",
    "discount": "Discount",
}
PREFERRED_SHEET = "Associated Items"
FUNDING_METHOD = "factoring"
DISCOUNT_METHOD = "factoring discount"


def _cell_text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _to_number(value: Any) -> Optional[float]:
    if value is None or value == "":
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).replace("$", "").replace(",", "").strip()
    if text.startswith("(") and text.endswith(")"):  # accounting negative
        text = "-" + text[1:-1]
    try:
        return float(text) if text else 0.0
    except ValueError:
        return None


def _to_date(value: Any) -> Optional[datetime]:
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    text = _cell_text(value)
    if not text:
        return None
    for fmt in ("%m/%d/%Y", "%Y-%m-%d", "%m-%d-%Y", "%m/%d/%y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def _find_header(rows: List[List[Any]]) -> Optional[int]:
    for idx, row in enumerate(rows[:30]):
        names = {_cell_text(v).lower() for v in row if _cell_text(v)}
        if all(key in names for key in REQUIRED_COLUMNS):
            return idx
    return None


def read_rows(file_bytes: bytes, filename: str) -> List[Dict[str, Any]]:
    """Excel (.xlsx) or CSV -> [{invoice_code, disbursement_raw, funding_raw, discount_raw, row_number}]."""
    name = (filename or "").lower()
    if name.endswith(".csv"):
        text = file_bytes.decode("utf-8-sig", errors="replace")
        grid = [row for row in csv.reader(io.StringIO(text))]
    elif name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        try:
            workbook = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
        except Exception as exc:
            raise HTTPException(status_code=400, detail=f"Could not read the Excel file: {exc}")
        sheet = workbook[PREFERRED_SHEET] if PREFERRED_SHEET in workbook.sheetnames else workbook.worksheets[0]
        grid = [list(row) for row in sheet.iter_rows(values_only=True)]
    else:
        raise HTTPException(status_code=400, detail="Upload an .xlsx or .csv file (old .xls files: re-save as .xlsx)")

    header_idx = _find_header(grid)
    if header_idx is None:
        raise HTTPException(status_code=400, detail="Could not find the header row. Expected columns: "
                                                    + ", ".join(REQUIRED_COLUMNS.values()) + ".")
    header = [_cell_text(v).lower() for v in grid[header_idx]]
    col = {key: header.index(key) for key in REQUIRED_COLUMNS}
    type_col = header.index("item type") if "item type" in header else None

    rows = []
    for offset, raw in enumerate(grid[header_idx + 1:], start=header_idx + 2):
        cell = lambda i: raw[i] if i < len(raw) else None
        if type_col is not None and _cell_text(cell(type_col)).lower() != "invoice":
            continue
        code = _cell_text(cell(col["item number"]))
        if not code:
            continue
        rows.append({
            "row_number": offset,
            "invoice_code": code,
            "disbursement_raw": cell(col["disbursement date"]),
            "funding_raw": cell(col["funding amount"]),
            "discount_raw": cell(col["discount"]),
        })
    return rows


def _find_invoice(db: Session, code: str) -> Optional[Invoice]:
    """Match the report's invoice # to ours: INV-0012, Inv-0012, inv 12 and 12 all match INV-0012."""
    exact = db.query(Invoice).filter(func.lower(Invoice.code) == code.lower()).first()
    if exact:
        return exact
    digits = "".join(ch for ch in code if ch.isdigit())
    if not digits:
        return None
    candidates = [inv for inv in db.query(Invoice).all()
                  if "".join(ch for ch in inv.code if ch.isdigit()).lstrip("0") == digits.lstrip("0")]
    return candidates[0] if len(candidates) == 1 else None


def _same(a: Optional[float], b: Optional[float]) -> bool:
    return a is not None and b is not None and abs(a - b) < 0.005


def preview(db: Session, rows: List[Dict[str, Any]]) -> Dict[str, Any]:
    """Check every row against our invoices. action is one of:
    update | matches (already has these values) | discrepancy (has different values) | invalid."""
    out = []
    for row in rows:
        invoice = _find_invoice(db, row["invoice_code"])
        funding = _to_number(row["funding_raw"])
        discount = _to_number(row["discount_raw"])
        disbursed = _to_date(row["disbursement_raw"])
        errors = []
        if not invoice:
            errors.append("Invoice not found")
        elif invoice.status == "void":
            errors.append("Invoice is void")
        if disbursed is None:
            errors.append("Invalid Disbursement Date")
        if funding is None or funding <= 0:
            errors.append("Funding Amount must be greater than 0")
        if discount is None or discount < 0:
            errors.append("Discount cannot be negative")

        entry = {
            "row_number": row["row_number"],
            "invoice_code": row["invoice_code"],
            "invoice_id": invoice.id if invoice else None,
            "matched_code": invoice.code if invoice else None,
            "invoice_status": invoice.status if invoice else None,
            "invoice_total": round(invoice.total, 2) if invoice else None,
            "invoice_balance": round(invoice.balance, 2) if invoice else None,
            "disbursement_date": disbursed,
            "funding_amount": round(funding, 2) if funding is not None else None,
            "funding_discount": round(discount, 2) if discount is not None else None,
            "sum_check": round((funding or 0) + (discount or 0), 2),
            "errors": errors,
        }
        entry["matches_invoice_total"] = bool(invoice) and abs(entry["sum_check"] - round(invoice.total, 2)) < 0.01

        if errors:
            entry["action"] = "invalid"
        elif invoice.funding_amount is None and invoice.funding_discount is None and invoice.disbursement_date is None:
            entry["action"] = "update"
        elif (_same(invoice.funding_amount, funding) and _same(invoice.funding_discount or 0, discount)
              and invoice.disbursement_date and invoice.disbursement_date.date() == disbursed.date()):
            entry["action"] = "matches"
        else:
            entry["action"] = "discrepancy"
            entry["existing"] = {
                "disbursement_date": invoice.disbursement_date,
                "funding_amount": invoice.funding_amount,
                "funding_discount": invoice.funding_discount,
            }
        out.append(entry)

    counts = {a: sum(1 for r in out if r["action"] == a) for a in ("update", "matches", "discrepancy", "invalid")}
    return {"total_rows": len(out), "counts": counts, "rows": out}


def apply(db: Session, file_name: str, rows: List[Dict[str, Any]], record_payments: bool, by: str) -> Dict[str, Any]:
    """Set the funding fields on every "update" row (and record payments if asked), all in
    one transaction. Rows that already match or differ are reported, never overwritten."""
    checked = preview(db, rows)
    to_update = [r for r in checked["rows"] if r["action"] == "update"]
    if not to_update:
        raise HTTPException(status_code=400, detail="Nothing to update -- every row is invalid, already applied, or differs from what's on the invoice")

    batch = FundingImport(file_name=file_name, record_payments=record_payments, created_by=by)
    db.add(batch)
    db.flush()
    warnings = []
    try:
        for r in to_update:
            invoice = db.get(Invoice, r["invoice_id"])
            invoice.disbursement_date = r["disbursement_date"]
            invoice.funding_amount = r["funding_amount"]
            invoice.funding_discount = r["funding_discount"]
            invoice.funding_import_id = batch.id
            if not record_payments:
                continue
            if invoice.status == "draft":
                invoice.status = "sent"  # it was funded, so it went out
            room = round(invoice.balance, 2)
            for amount, method in ((r["funding_amount"], FUNDING_METHOD), (r["funding_discount"], DISCOUNT_METHOD)):
                take = round(min(amount, room), 2)
                if take <= 0:
                    continue
                db.add(InvoicePayment(invoice_id=invoice.id, amount=take, paid_date=r["disbursement_date"], method=method,
                                      reference=file_name, note=f"Bulk funding upload #{batch.id}",
                                      funding_import_id=batch.id, created_by=by))
                room = round(room - take, 2)
            if r["sum_check"] > round(invoice.balance, 2) + 0.005:
                warnings.append(f"{invoice.code}: funding + discount ({r['sum_check']:.2f}) is more than the open balance "
                                f"({invoice.balance:.2f}) -- only the balance was recorded as paid")
            db.flush()
            db.refresh(invoice)
            if invoice.balance <= 0.005:
                invoice.status = "paid"
            elif not r["matches_invoice_total"]:
                warnings.append(f"{invoice.code}: funding + discount ({r['sum_check']:.2f}) doesn't equal the invoice total "
                                f"({r['invoice_total']:.2f}) -- {invoice.balance:.2f} is still open")

        skipped = [r for r in checked["rows"] if r["action"] != "update"]
        batch.updated_count = len(to_update)
        batch.skipped_count = len(skipped)
        batch.summary = json.dumps({
            "updated": [r["matched_code"] for r in to_update],
            "skipped": [{"invoice_code": r["invoice_code"], "action": r["action"], "errors": r["errors"]} for r in skipped],
            "warnings": warnings,
        })
        db.commit()
    except Exception:
        db.rollback()
        raise
    return {"import_id": batch.id, "updated_count": len(to_update), "counts": checked["counts"],
            "rows": checked["rows"], "warnings": warnings}


def rollback(db: Session, import_id: int) -> FundingImport:
    """Undo one upload: remove the payments it recorded and clear the fields it set."""
    batch = db.get(FundingImport, import_id)
    if not batch:
        raise HTTPException(status_code=404, detail="Upload not found")
    if batch.rolled_back:
        raise HTTPException(status_code=400, detail="This upload was already rolled back")
    for payment in db.query(InvoicePayment).filter(InvoicePayment.funding_import_id == batch.id).all():
        db.delete(payment)
    db.flush()
    for invoice in db.query(Invoice).filter(Invoice.funding_import_id == batch.id).all():
        db.refresh(invoice)  # drop the deleted payments from invoice.payments before the fields change
        invoice.disbursement_date = invoice.funding_amount = invoice.funding_discount = None
        invoice.funding_import_id = None
        if invoice.status == "paid" and invoice.balance > 0.005:
            invoice.status = "sent"
    batch.rolled_back = True
    db.commit()
    db.refresh(batch)
    return batch


def history(db: Session, limit: int = 20) -> List[FundingImport]:
    return db.query(FundingImport).order_by(FundingImport.id.desc()).limit(limit).all()
