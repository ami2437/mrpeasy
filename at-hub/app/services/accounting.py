"""Accounting: the company's own books beside AT-HUB's operations data (routes/accounting.py, accounting.html).

Statement lines (bank, later cards) come in from CSV / Excel files, every line goes in an account -- by a person's
rule, by what AT-HUB learned from the lines already filed (acct_mapper.py), or by a person in the inbox -- and the
reports show where the money went: P&L by month / quarter / fiscal year (April-March by default), accounts, loans,
the partners' split, AP / AR, cash flow.

Room for typed-in entries everywhere: a line can be typed in ahead of the statement ("expected": matched, not
duplicated, when the statement brings it), "Book Entries" holds lines that are never on a statement (year-end
figures, adjustments), AP / AR take typed-in items and a person's word over AT-HUB's figure ("already paid").

Kept separate from AT-HUB's invoices and payments on purpose (user, 2026-10-10): nothing here changes them.
QuickBooks later: accounts carry a QuickBooks type, everything carries external_id, exports are QuickBooks-shaped."""
import csv
import hashlib
import io
import json
import re
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from typing import Dict, List, Optional, Tuple

from fastapi import HTTPException
from sqlalchemy import func
from sqlalchemy.orm import Session, selectinload

from app.models import (AcctAccount, AcctChange, AcctCheckOk, AcctIgnored, AcctImport, AcctMemory, AcctOpenItem, AcctRule, AcctSource, AcctSplit,
                        AcctTxn, AppSetting, Customer, CustomerOrder, Invoice, PurchaseOrder, Vendor)
from app.services import acct_mapper, acct_rules, acct_tails, clock
from app.services.money import cents, line_amount

KINDS = {  # kind -> (label, QuickBooks account type, in the P&L?)
    "income": ("Income", "Income", True),
    "cogs": ("Cost Of Goods", "Cost of Goods Sold", True),
    "expense": ("Expenses", "Expense", True),
    "loan": ("Loans", "Long Term Liabilities", False),
    "owner": ("Partner Draws", "Equity", False),
    "transfer": ("Transfers / Wash", "Bank", False),
    "pending": ("Set Aside: AR / AP Estimates", "Other Current Liabilities", False),  # in no total (user, 2026-10-10)
    "other": ("Other", "Other Expense", False),
}
# The user's workbook categories (FY2024-25_Input_File.xlsx) -> our kinds
WORKBOOK_KINDS = {"payments": "income", "income": "income", "cogs": "cogs", "expense": "expense", "expenses": "expense",
                  "loans": "loan", "loan": "loan", "shares": "owner", "share": "owner", "wash": "transfer",
                  "transfer": "transfer", "pending": "pending"}
# Their category names (the input / output workbooks) for each kind
CATEGORY = {"income": "Payments", "cogs": "COGS", "expense": "Expense", "loan": "Loans", "owner": "Shares", "transfer": "Wash",
            "pending": "Pending", "other": "Other"}
SETTINGS_KEY = "acct_settings"
DEFAULT_SETTINGS = {"fy_start_month": 4, "partners": [{"name": "Anuj", "share": 1 / 3}, {"name": "Niraj", "share": 1 / 3},
                                                      {"name": "Richard", "share": 1 / 3}],
                    "years": {}, "apar_groups": {"open_orders": False},
                    "auto_limit": 5000}  # AT-HUB / rules never file a line bigger than this by themselves (user, 2026-10-12)


# ---------------------------------------------------------------- settings, periods
def today() -> date:
    """The company's today as a date (clock.today() is a datetime at midnight)."""
    return clock.today().date()


def settings(db: Session) -> dict:
    row = db.get(AppSetting, SETTINGS_KEY)
    s = json.loads(row.value) if row and row.value else {}
    return {**DEFAULT_SETTINGS, **s}


def save_settings(db: Session, s: dict, by: str) -> dict:
    row = db.get(AppSetting, SETTINGS_KEY)
    if not row:
        row = AppSetting(key=SETTINGS_KEY)
        db.add(row)
    row.value, row.updated_by = json.dumps(s), by
    db.commit()
    return s


def fy_start(d: date, month: int) -> date:
    return date(d.year if d.month >= month else d.year - 1, month, 1)


def fy_label(start: date, month: int) -> str:
    if month == 1:
        return f"FY{start.year}"
    return f"FY{start.year}-{str(start.year + 1)[2:]}"


def fy_range(start_year: int, month: int) -> Tuple[date, date]:
    a = date(start_year, month, 1)
    b = date(start_year + 1, month, 1) - timedelta(days=1)
    return a, b


def years(db: Session) -> List[dict]:
    """The fiscal years that have lines (newest first), plus the current one."""
    m = settings(db)["fy_start_month"]
    starts = {fy_start(today(), m).year}
    for (d,) in db.query(AcctTxn.date).distinct():
        starts.add(fy_start(d.date(), m).year)
    out = []
    for y in sorted(starts, reverse=True):
        a, b = fy_range(y, m)
        out.append({"label": fy_label(a, m), "start": a.isoformat(), "end": b.isoformat(), "year": y})
    return out


def parse_day(v) -> Optional[date]:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for f in ("%m/%d/%Y", "%Y-%m-%d", "%m/%d/%y", "%m-%d-%Y", "%d-%b-%Y", "%b %d, %Y", "%Y-%m-%d %H:%M:%S", "%m/%d/%Y %H:%M"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def parse_money(v) -> Optional[float]:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace("$", "").replace(",", "")
    neg = s.startswith("(") and s.endswith(")")
    s = s.strip("()")
    try:
        x = float(s)
    except ValueError:
        return None
    return -x if neg else x


def _dt(d: date) -> datetime:
    return datetime(d.year, d.month, d.day)


# ---------------------------------------------------------------- accounts, sources
def book_source(db: Session) -> AcctSource:
    s = db.query(AcctSource).filter(AcctSource.kind == "book").first()
    if not s:
        s = AcctSource(name="Book Entries", kind="book")
        db.add(s)
        db.flush()
    return s


def account_by_name(db: Session, name: str, kind: str, partner: Optional[str] = None) -> AcctAccount:
    key = name.strip()
    a = next((x for x in db.query(AcctAccount).all() if x.name.strip().lower() == key.lower()), None)
    if not a:
        a = AcctAccount(name=key, kind=kind, partner=partner, qb_type=KINDS.get(kind, KINDS["other"])[1])
        db.add(a)
        db.flush()
    return a


SEP = " › "


def ancestors(a: AcctAccount, accounts: Dict[int, AcctAccount]) -> List[AcctAccount]:
    """Parent, grandparent ... (nearest first); stops on a loop."""
    out, seen, cur = [], {a.id}, a
    while cur.parent_id and cur.parent_id in accounts and cur.parent_id not in seen:
        cur = accounts[cur.parent_id]
        seen.add(cur.id)
        out.append(cur)
    return out


def path(a: AcctAccount, accounts: Dict[int, AcctAccount]) -> str:
    """'Labor › Faustino' -- how a sub-account is shown everywhere."""
    return SEP.join([x.name for x in reversed(ancestors(a, accounts))] + [a.name])


def descendants(aid: int, accounts: Dict[int, AcctAccount]) -> List[int]:
    kids = defaultdict(list)
    for x in accounts.values():
        if x.parent_id:
            kids[x.parent_id].append(x.id)
    out, todo = [], list(kids.get(aid, []))
    while todo:
        k = todo.pop()
        if k not in out and k != aid:
            out.append(k)
            todo += kids.get(k, [])
    return out


def tree_order(accs: List[AcctAccount], accounts: Dict[int, AcctAccount], key=None) -> List[Tuple[AcctAccount, int]]:
    """(account, depth) parents first, each followed by its sub-accounts; siblings by `key` (default name)."""
    key = key or (lambda x: x.name.lower())
    ids = {a.id for a in accs}
    kids = defaultdict(list)
    tops = []
    for a in accs:
        (kids[a.parent_id] if a.parent_id in ids else tops).append(a)
    out, seen = [], set()

    def walk(a, d):
        if a.id in seen:
            return
        seen.add(a.id)
        out.append((a, d))
        for k in sorted(kids.get(a.id, []), key=key):
            walk(k, d + 1)
    for a in sorted(tops, key=key):
        walk(a, 0)
    return out


def account_dict(a: AcctAccount, accounts: Optional[Dict[int, AcctAccount]] = None) -> dict:
    accounts = accounts or {}
    up = ancestors(a, accounts) if accounts else []
    return {"id": a.id, "name": a.name, "path": path(a, accounts) if accounts else a.name, "parent_id": a.parent_id, "depth": len(up),
            "kind": a.kind, "kind_label": KINDS.get(a.kind, KINDS["other"])[0], "partner": a.partner,
            "qb_type": a.qb_type or KINDS.get(a.kind, KINDS["other"])[1], "external_id": a.external_id, "note": a.note,
            "opening_balance": a.opening_balance or 0.0, "is_active": a.is_active}


def source_dict(db: Session, s: AcctSource) -> dict:
    rows = db.query(AcctTxn.amount, AcctTxn.date).filter(AcctTxn.source_id == s.id).all()
    last = max((d for _a, d in rows), default=None)
    return {"id": s.id, "name": s.name, "kind": s.kind, "last4": s.last4, "opening_balance": s.opening_balance or 0.0,
            "opening_date": s.opening_date.date().isoformat() if s.opening_date else None, "external_id": s.external_id,
            "is_active": s.is_active, "lines": len(rows), "last_date": last.date().isoformat() if last else None,
            "balance": cents((s.opening_balance or 0.0) + sum(a for a, d in rows if not s.opening_date or d >= s.opening_date)),
            "layout": json.loads(s.layout) if s.layout else None}


# ---------------------------------------------------------------- rules + what AT-HUB learned
def rule_specs(db: Session) -> List[dict]:
    """Every rule in its order (priority, then age) as an acct_rules spec."""
    rs = db.query(AcctRule).all()
    rs.sort(key=lambda r: (r.priority if r.priority is not None else 10 ** 6, r.id))
    return [acct_rules.spec(r) for r in rs]


def mapper(db: Session) -> acct_mapper.Mapper:
    """Learns from every line a person (or their workbook / a rule they applied) put in an account -- not from AT-HUB's
    own picks, not from lines someone told it to forget. A line a person corrected counts double."""
    ex = (db.query(AcctTxn.description, AcctTxn.bank_description, AcctTxn.amount, AcctTxn.account_id, AcctTxn.date, AcctTxn.corrected)
          .filter(AcctTxn.account_id.isnot(None), AcctTxn.how.in_(("person", "workbook", "rule")),
                  (AcctTxn.no_learn.is_(None)) | (AcctTxn.no_learn.is_(False))).all())
    return acct_mapper.Mapper([(b or d, a, acc, dt.day if dt else None, 2.0 if c else 1.0) for d, b, a, acc, dt, c in ex])


def rule_decision(specs: List[dict], t) -> Tuple[Optional[dict], List[dict]]:
    """(the rule that decides, the rules that disagree) for a line -- acct_rules.decide_rules."""
    return acct_rules.decide_rules(specs, acct_rules.line_of(t))


def _clash_info(clash: List[dict]) -> List[dict]:
    return [{"rule_id": r["id"], "name": r["name"] or acct_rules.describe(r), "account_id": r["account_id"],
             "splits": r["splits"], "action": r["action"]} for r in clash]


def decide(t: AcctTxn, m: acct_mapper.Mapper, specs: List[dict]) -> Tuple[Optional[int], Optional[str], str, Optional[dict]]:
    """(account, sure | guess | None, why, the rule that decided) -- rules first, then what was learned. Two rules that
    match and disagree decide nothing: the line waits in Review."""
    r, clash = rule_decision(specs, t)
    if clash:
        return None, None, "rules disagree: " + "; ".join(x["name"] or acct_rules.describe(x) for x in clash[:3]), None
    if r:
        label = r["name"] or acct_rules.describe(r)
        return r["account_id"] or (r["splits"][0]["account_id"] if r["splits"] else None), ("sure" if r["action"] != "suggest" else "guess"), f"rule: {label}", r
    acc, conf, why = m.suggest(t.bank_description or t.description, t.amount, day=t.date.day if t.date else None)
    return acc, conf, why, None


def file_by_rule(db: Session, t: AcctTxn, r: dict, by: Optional[str], reviewed: bool) -> None:
    """Put a line where a rule says: one account, or the rule's split; its payee / note when the rule sets them."""
    t.splits.clear()
    ps = acct_rules.parts(r, t.amount)
    if len(ps) == 1:
        t.account_id = ps[0][0]
    else:
        t.account_id = None
        for aid, amt in ps:
            t.splits.append(AcctSplit(account_id=aid, amount=amt, note=(r["name"] or None)))
    if r["set_payee"]:
        t.payee = r["set_payee"]
    if r["set_note"] and not t.note:
        t.note = r["set_note"]
    t.how, t.reviewed = "rule", reviewed
    t.rule_id = r.get("id")
    if r.get("review") and not reviewed:  # the rule files it, a person still looks
        t.flag = "check"
        t.flag_info = json.dumps({"reason": f"Filed by rule \"{r['name'] or acct_rules.describe(r)}\" -- check it", "rule_id": r.get("id")})
    elif t.flag in ("check", "conflict", "over_limit"):
        t.flag = t.flag_info = None
    if by:
        t.updated_by = by
    row = db.get(AcctRule, r["id"]) if r.get("id") else None
    if row:
        row.hits = (row.hits or 0) + 1


def autofill(db: Session, txns: Optional[List[AcctTxn]] = None, m=None) -> int:
    """Put the lines with no account in one when AT-HUB is sure (a rule that files by itself, or the same payee + note
    / amount as before). A rule set to "suggest" only suggests."""
    m = m or mapper(db)
    specs = rule_specs(db)
    limit = float(settings(db).get("auto_limit") or 0)
    if txns is None:
        txns = db.query(AcctTxn).options(selectinload(AcctTxn.splits)).filter(AcctTxn.account_id.is_(None)).all()
    n = 0

    def too_big(t, acc_id, why):
        """Over the limit: never filed by itself -- the account stays a suggestion and the line waits in Review."""
        if limit and abs(t.amount) > limit + 0.004:
            t.flag, t.flag_info = "over_limit", json.dumps({"reason": f"Over ${limit:,.0f} -- AT-HUB doesn't file it by itself ({why})", "account_id": acc_id})
            return True
        return False
    for t in txns:
        if t.account_id or t.splits:
            continue
        r, clash = rule_decision(specs, t)
        if clash:  # two rules say different things: nobody files it, it waits in Review with a red flag
            if t.flag != "conflict":
                t.flag, t.flag_info = "conflict", json.dumps({"reason": "Rules disagree", "rules": _clash_info(clash)})
            continue
        if t.flag == "conflict":
            t.flag = t.flag_info = None
        if r:
            if r["action"] != "suggest" and not too_big(t, r["account_id"], f"rule {r['name'] or acct_rules.describe(r)}"):
                file_by_rule(db, t, r, None, False)
                n += 1
            continue
        acc, conf, _why = m.suggest(t.bank_description or t.description, t.amount, day=t.date.day if t.date else None)
        if conf == "sure" and acc and not too_big(t, acc, "learned from earlier lines"):
            t.account_id, t.how, t.reviewed = acc, "auto", False
            n += 1
    return n


def already_in(db: Session, t: AcctTxn, limit: int = 40) -> List[dict]:
    """The same payee's lines already in the account this line just went into -- shown in the "like this" pop-up."""
    if not t.account_id:
        return []
    accounts, sources = maps(db)
    key = acct_mapper.payee(t.bank_description or t.description)
    out = [txn_dict(x, accounts, sources) for x in db.query(AcctTxn).options(selectinload(AcctTxn.splits))
           .filter(AcctTxn.account_id == t.account_id, AcctTxn.id != t.id).order_by(AcctTxn.date.desc()).all()
           if acct_mapper.payee(x.bank_description or x.description) == key]
    return out[:limit]


def last_splits(db: Session) -> Dict[Tuple[str, bool], dict]:
    """Per payee (and money in / out): the latest line a person split -- "split it like last time" for the next one."""
    out = {}
    for t in (db.query(AcctTxn).options(selectinload(AcctTxn.splits)).filter(AcctTxn.splits.any(), AcctTxn.how == "person")
              .order_by(AcctTxn.date).all()):
        if not t.amount:
            continue
        k = (acct_mapper.payee(t.bank_description or t.description), t.amount >= 0)
        out[k] = {"date": t.date.date().isoformat(), "txn_id": t.id,
                  "parts": [{"account_id": x.account_id, "pct": round(x.amount / t.amount * 100, 4), "note": x.note} for x in t.splits]}
    return out


def suggestions(db: Session, txns: List[AcctTxn]) -> Dict[int, dict]:
    m = mapper(db)
    specs = rule_specs(db)
    splits_by = None
    out = {}
    for t in txns:
        r, clash = rule_decision(specs, t)
        if clash:
            out[t.id] = {"account_id": None, "confidence": None, "why": "rules disagree", "conflict": _clash_info(clash)}
            continue
        acc, conf, why, r = decide(t, m, specs)
        if conf != "sure" and not r:
            if splits_by is None:
                splits_by = last_splits(db)
            ls = splits_by.get((acct_mapper.payee(t.bank_description or t.description), t.amount >= 0))
            if ls and len(ls["parts"]) > 1:
                out[t.id] = {"account_id": None, "confidence": "guess", "why": f"split like the {ls['date']} line", "split": ls["parts"]}
                continue
        if acc:
            out[t.id] = {"account_id": acc, "confidence": conf or "maybe", "why": why, "rule_id": r["id"] if r else None}
    return out


def rule_lines(db: Session, rule: dict, account_id: Optional[int] = None) -> dict:
    """Every line (all dates) a rule matches, grouped: not filed, filed by AT-HUB (not looked at), filed elsewhere by a
    person (or the workbook), already where the rule puts it, split (left alone unless the rule splits)."""
    s_ = acct_rules.spec(rule) if "conditions" not in rule or "match" not in rule else rule
    account_id = account_id or s_["account_id"]
    accounts, sources = maps(db)
    groups = {"unfiled": [], "auto": [], "elsewhere": [], "here": [], "split": []}
    for t in db.query(AcctTxn).options(selectinload(AcctTxn.splits)).order_by(AcctTxn.date.desc(), AcctTxn.id.desc()).all():
        if not acct_rules.matches(s_, acct_rules.line_of(t)):
            continue
        if t.splits and not s_["splits"]:
            g = "split"
        elif not s_["splits"] and account_id and t.account_id == account_id:
            g = "here"
        elif s_["splits"] and t.how == "rule" and t.reviewed and sorted((x.account_id, x.amount) for x in t.splits) == sorted(acct_rules.parts(s_, t.amount)):
            g = "here"
        elif not t.account_id and not t.splits:
            g = "unfiled"
        elif not t.reviewed:
            g = "auto"
        else:
            g = "elsewhere"
        groups[g].append(txn_dict(t, accounts, sources))
    return {"groups": groups, "count": sum(len(v) for v in groups.values()),
            "can_apply": len(groups["unfiled"]) + len(groups["auto"]) + len(groups["elsewhere"])}


def apply_rule(db: Session, r: AcctRule, ids: List[int], by: str) -> int:
    """Put the picked lines where the rule says (only lines it really matches)."""
    s_ = acct_rules.spec(r)
    n = 0
    for t in db.query(AcctTxn).options(selectinload(AcctTxn.splits)).filter(AcctTxn.id.in_(ids)).all():
        if (t.splits and not s_["splits"]) or not acct_rules.matches(s_, acct_rules.line_of(t)):
            continue
        file_by_rule(db, t, s_, by, True)  # counts a hit on the rule
        n += 1
    db.commit()
    return n


def look_alikes(db: Session, t: AcctTxn, limit: int = 60) -> List[dict]:
    """After a person files a line: the lines with no account that look like it (same payee, or AT-HUB now leans to
    the same account) but that it isn't sure enough about to file by itself -- for the "like this" pop-up."""
    if not t.account_id:
        return []
    m = mapper(db)
    specs = rule_specs(db)
    key = acct_mapper.payee(t.bank_description or t.description)
    accounts, sources = maps(db)
    out = []
    for x in (db.query(AcctTxn).options(selectinload(AcctTxn.splits)).filter(AcctTxn.account_id.is_(None), AcctTxn.id != t.id)
              .order_by(AcctTxn.date.desc()).all()):
        if x.splits:
            continue
        if (x.amount >= 0) != (t.amount >= 0):
            continue  # money in vs money out: not alike
        same_payee = acct_mapper.payee(x.bank_description or x.description) == key
        acc, conf, why, _r = decide(x, m, specs)
        if same_payee or acc == t.account_id:
            d = txn_dict(x, accounts, sources)
            d["lean"] = acc == t.account_id
            d["why"] = why
            out.append(d)
            if len(out) >= limit:
                break
    out.sort(key=lambda d: (not d["lean"], d["date"]), reverse=False)
    return out


def recheck_after_correction(db: Session, t: AcctTxn, old_account: Optional[int]) -> int:
    """A person moved a line out of the account AT-HUB had picked: AT-HUB's other unchecked picks of that payee into
    the same old account go back to To File (with fresh suggestions) unless it's still sure about them."""
    if not old_account:
        return 0
    db.flush()
    m = mapper(db)
    specs = rule_specs(db)
    key = acct_mapper.payee(t.bank_description or t.description)
    n = 0
    for x in db.query(AcctTxn).options(selectinload(AcctTxn.splits)).filter(AcctTxn.account_id == old_account, AcctTxn.how == "auto",
                                                                             AcctTxn.reviewed.is_(False), AcctTxn.id != t.id).all():
        if acct_mapper.payee(x.bank_description or x.description) != key:
            continue
        acc, conf, _w, _r = decide(x, m, specs)
        if not (acc == old_account and conf == "sure"):
            x.account_id, x.how = None, None
            n += 1
    return n


def learned(db: Session) -> List[dict]:
    """What AT-HUB learned: per payee, the accounts people put its lines in, how many times; mixed ones flagged."""
    accounts = {a.id: a for a in db.query(AcctAccount).all()}
    groups = defaultdict(lambda: {"accounts": Counter(), "forgotten": 0, "label": "", "last": None, "sample": ""})
    for t in (db.query(AcctTxn).filter(AcctTxn.account_id.isnot(None), AcctTxn.how.in_(("person", "workbook", "rule"))).all()):
        desc = t.bank_description or t.description
        k = acct_mapper.payee(desc)
        g = groups[k]
        g["label"] = g["label"] or acct_mapper.payee_label(desc)
        g["sample"] = g["sample"] or desc
        if t.no_learn:
            g["forgotten"] += 1
        else:
            g["accounts"][t.account_id] += 1
        d = t.date.date()
        g["last"] = max(g["last"], d) if g["last"] else d
    out = []
    for k, g in groups.items():
        total = sum(g["accounts"].values())
        accs = [{"account_id": a, "account": path(accounts[a], accounts) if a in accounts else "?", "count": c} for a, c in g["accounts"].most_common()]
        top = accs[0]["count"] / total if total else 0
        out.append({"key": k, "label": g["label"], "sample": g["sample"], "lines": total, "forgotten": g["forgotten"],
                    "accounts": accs, "how": "one account" if len(accs) == 1 else ("mostly one" if top >= 0.8 else "mixed -- the note words decide"),
                    "last": g["last"].isoformat() if g["last"] else None})
    out.sort(key=lambda x: (-x["lines"], x["label"]))
    return out


def forget(db: Session, key: str, on: bool = True) -> int:
    n = 0
    for t in db.query(AcctTxn).filter(AcctTxn.account_id.isnot(None)).all():
        if acct_mapper.payee(t.bank_description or t.description) == key:
            t.no_learn = True if on else None
            n += 1
    db.commit()
    return n


# ---- rules AT-HUB suggests from patterns in all the lines (filed and not)
GENERIC = set("""ZELLE TO FROM ON ONLINE TRANSFER DIRECT PAY PAYMENT PAYMENTS TRAN ID WF ACH BUSINESS WIRE WT MONEY AUTHORIZED CASH APP
DEPOSIT CHECK REF PMT WEB MOBILE EDI PYMNTS PYMT INVOICE AMERICAN TRADERS LLC INC BANK FARGO WELLS CHECKING EVERYDAY SERV
CARD CREDIT DEBIT FEE TRANS INDIVIDUAL MONTHLY BASE""".split())
_KIND_WORDS = {"zelle": "ZELLE", "wire": "", "ach": "", "transfer": "TRANSFER", "directpay": "DIRECT PAY", "cashapp": "CASH APP", "payee": "", "text": ""}
DISMISS_KEY = "acct_rule_dismissed"


def _pattern_words(key: str) -> str:
    kind, _, name = key.partition(":")
    return " ".join(x for x in (_KIND_WORDS.get(kind, ""), name) if x).strip()


def suggest_rules(db: Session, limit: int = 12) -> List[dict]:
    """Repeated lines (3+) sharing a payee -- or, for a payee whose lines go to different accounts, a payee + note word --
    turned into ready-made rules: the account people already used (or AT-HUB's guess when none is filed yet), how many
    lines it would file now. Patterns an existing rule covers, and ones dismissed, are left out."""
    row = db.get(AppSetting, DISMISS_KEY)
    dismissed = set(json.loads(row.value)) if row and row.value else set()
    specs = [x for x in rule_specs(db) if x["enabled"]]
    accounts = {a.id: a for a in db.query(AcctAccount).all()}
    txns = db.query(AcctTxn).options(selectinload(AcctTxn.splits)).all()
    info = []
    for t in txns:
        if t.splits:
            continue
        a = accounts.get(t.account_id)
        if a and a.kind == "pending":
            continue
        desc = t.bank_description or t.description
        info.append((t, acct_mapper.payee(desc), acct_mapper.words(desc), acct_rules.line_of(t)))
    by_key = defaultdict(list)
    for x in info:
        by_key[x[1]].append(x)
    m = mapper(db)
    cands = []

    def filed_by_person(t):
        return t.account_id and t.how in ("person", "workbook", "rule") and t.reviewed and not t.no_learn

    def add(words_, sign, members):
        if not words_ or len(members) < 3:
            return
        if not any(w not in GENERIC and len(w) >= 3 for w in words_.split()):
            return  # only bank wording ("DIRECT PAY", "ONLINE") -- says nothing about who
        spec_ = {"name": "", "match": "all", "conditions": [{"field": "text", "op": "contains", "value": words_}], "sign": sign,
                 "source_ids": [], "account_id": None, "splits": [], "set_payee": None, "set_note": None, "action": "file",
                 "enabled": True, "priority": 0, "id": None}
        hit = [x for x in info if acct_rules.matches(spec_, x[3])]
        if len(hit) < 3:
            return
        filed = Counter(x[0].account_id for x in hit if filed_by_person(x[0]))  # every line the rule would match, not just this payee's
        total_f = sum(filed.values())
        unfiled = [x[0] for x in hit if not x[0].account_id]
        if total_f:
            acc, k = filed.most_common(1)[0]
            if k < 2 or k / total_f < 0.9:
                return  # people put these in different places -- not one rule
            source = f"you put {k} of them in it"
        else:
            if not unfiled:
                return
            votes = Counter()
            for t in unfiled:
                a_, _c, _w = m.suggest(t.bank_description or t.description, t.amount)
                if a_:
                    votes[a_] += 1
            acc = votes.most_common(1)[0][0] if votes else None
            source = "AT-HUB's guess" if acc else "pick an account"
        if not unfiled and not total_f:
            return
        covered = sum(1 for x in hit if any(acct_rules.matches(r, x[3]) for r in specs))
        if covered >= 0.8 * len(hit):
            return
        dkey = f"{words_}|{sign or ''}|{acc or ''}"
        if dkey in dismissed:
            return
        n_unfiled = sum(1 for x in hit if not x[0].account_id)
        n_here = sum(1 for x in hit if acc and x[0].account_id == acc)
        cands.append({"key": dkey, "words": words_, "sign": sign, "account_id": acc,
                      "account": path(accounts[acc], accounts) if acc in accounts else None, "source": source,
                      "lines": len(hit), "unfiled": n_unfiled, "already": n_here, "elsewhere": len(hit) - n_unfiled - n_here,
                      "sample": (hit[0][0].bank_description or hit[0][0].description), "_ids": frozenset(x[0].id for x in hit)})

    for key, members in by_key.items():
        if len(members) < 3:
            continue
        signs = {1 if t.amount >= 0 else -1 for t, *_ in members}
        sign = None if len(signs) > 1 else ("in" if signs == {1} else "out")
        base = _pattern_words(key)
        filed = Counter(t.account_id for t, *_ in members if filed_by_person(t))
        mixed = len(filed) > 1 and filed.most_common(1)[0][1] / sum(filed.values()) < 0.9
        if not mixed:
            add(base, sign, members)
            continue
        # one payee, several accounts: a note word that keeps going to one account (ZELLE SHAH NIRAJ + MAALIKS -> Richard)
        frame = set(base.split()) | {"ZELLE", "TRANSFER", "DIRECT", "PAY"}
        wc = Counter(w for _t, _k, ws, _l in members for w in ws - frame)
        for w, c in wc.most_common(8):
            if c < 3:
                break
            sub = [x for x in members if w in x[2]]
            add(f"{base} {w}".strip(), sign, sub)
    seen, same_lines, out = set(), set(), []
    for c in sorted(cands, key=lambda c: (-c["unfiled"], -c["lines"], len(c["words"]))):
        ids = c.pop("_ids")
        if c["key"] in seen or ids in same_lines:
            continue  # two wordings that catch the very same lines: keep the shorter
        seen.add(c["key"])
        same_lines.add(ids)
        out.append(c)
    return out[:limit]


def dismiss_suggestion(db: Session, key: str, by: str) -> None:
    row = db.get(AppSetting, DISMISS_KEY)
    cur = set(json.loads(row.value)) if row and row.value else set()
    cur.add(key)
    if not row:
        row = AppSetting(key=DISMISS_KEY)
        db.add(row)
    row.value, row.updated_by = json.dumps(sorted(cur)), by
    db.commit()


# ---------------------------------------------------------------- statement files
def read_table(data: bytes, filename: str) -> List[List]:
    """Every non-empty row of a CSV or Excel file (first sheet; for a workbook, every sheet with a header)."""
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        rows = []
        for ws in wb.worksheets:
            sheet = [list(r) for r in ws.iter_rows(values_only=True) if any(c not in (None, "") for c in r)]
            if not rows:
                rows = sheet
            elif sheet and _header_score(sheet[0]) >= 3 and _header_score(sheet[0]) == _header_score(rows[0]):
                rows += sheet[1:]  # the same layout on another sheet (their workbook: "Sheet1" + "WF 0424-0325")
        return rows
    text = data.decode("utf-8-sig", errors="replace")
    return [r for r in csv.reader(io.StringIO(text)) if any(c.strip() for c in r)]


HEADER_WORDS = {"date": ("date", "posted", "transaction date", "trans date", "posting date"),
                "description": ("description", "memo", "details", "payee", "name", "transaction"),
                "amount": ("amount", "amt"), "debit": ("debit", "withdrawal", "withdrawals", "charge", "payment", "payments", "spent", "money out"),
                "credit": ("credit", "deposit", "deposits", "received", "money in"), "account": ("account",), "category": ("category",)}


def _header_score(row) -> int:
    cells = [str(c or "").strip().lower() for c in row]
    return sum(1 for c in cells if any(c == w or c.startswith(w) for ws in HEADER_WORDS.values() for w in ws))


def guess_layout(rows: List[List]) -> dict:
    """{has_header, date, description, amount | debit + credit, account?, category?} as column numbers."""
    if not rows:
        return {}
    head = rows[0]
    if _header_score(head) >= 2:
        cells = [str(c or "").strip().lower() for c in head]
        lay = {"has_header": True}
        used = set()
        for key, ws in HEADER_WORDS.items():
            for i, c in enumerate(cells):
                if i in used or key in lay:
                    continue
                if any(c == w for w in ws) or (key != "description" and any(c.startswith(w) for w in ws)):
                    lay[key] = i
                    used.add(i)
        if "date" not in lay:
            lay["date"] = next((i for i, c in enumerate(cells) if "date" in c), 0)
        if "description" not in lay:
            lay["description"] = next((i for i, c in enumerate(cells) if "desc" in c or "memo" in c), None)
        if "amount" in lay and "credit" in lay and "debit" not in lay:
            lay["debit"] = lay.pop("amount")  # QuickBooks register: "Amount" = what went out, "Deposit" = what came in
        if "account" in lay and "category" not in lay:
            lay["hint"] = lay.pop("account")  # QuickBooks' own account for the line -- a hint, not a workbook account
        lay["headers"] = [str(c or "").strip().lower() for c in head]  # this layout is only reused for files with these headers
        return lay
    # no header (Wells Fargo: date, amount, *, , description): find the columns by what's in them
    sample = rows[:30]
    width = max(len(r) for r in sample)
    def share(i, test):
        vals = [r[i] for r in sample if i < len(r) and r[i] not in (None, "")]
        return sum(1 for v in vals if test(v)) / max(1, len(vals)), len(vals)
    date_col = max(range(width), key=lambda i: share(i, lambda v: parse_day(v) is not None)[0])
    amt_cols = [i for i in range(width) if i != date_col and share(i, lambda v: parse_money(v) is not None)[0] > 0.9 and share(i, lambda v: True)[1] >= len(sample) * 0.8]
    def text_len(i):
        return sum(len(str(r[i] or "")) for r in sample if i < len(r))
    desc_col = max((i for i in range(width) if i != date_col and i not in amt_cols), key=text_len, default=None)
    return {"has_header": False, "date": date_col, "amount": amt_cols[0] if amt_cols else None, "description": desc_col}


def parse_rows(rows: List[List], lay: dict, kind: str = "bank") -> Tuple[List[dict], List[str]]:
    """Statement rows -> [{date, amount, description, account?, category?}] + the rows that couldn't be read."""
    out, bad = [], []
    body = rows[1:] if lay.get("has_header") else rows
    def cell(r, k):
        i = lay.get(k)
        return r[i] if i is not None and i < len(r) else None
    for n, r in enumerate(body, start=2 if lay.get("has_header") else 1):
        d = parse_day(cell(r, "date"))
        if lay.get("amount") is not None:
            a = parse_money(cell(r, "amount"))
        else:
            dr, cr = parse_money(cell(r, "debit")), parse_money(cell(r, "credit"))
            if dr is None and cr is None:
                a = None
            else:  # card statements: debit = a charge (money out), credit = a payment / refund
                a = (cr or 0.0) - abs(dr or 0.0)
        desc = str(cell(r, "description") or "").strip()
        if desc.lower().startswith("opening balance"):  # QuickBooks' "Opening Balance from Bank" -- not a transaction
            continue
        if d is None or a is None:
            if any(str(c or "").strip() for c in r):
                bad.append(f"row {n}: " + ", ".join(str(c) for c in r if c not in (None, ""))[:120])
            continue
        row = {"date": d, "amount": cents(a), "description": desc}
        if lay.get("account") is not None:
            row["account"] = str(cell(r, "account") or "").strip()
        if lay.get("category") is not None:
            row["category"] = str(cell(r, "category") or "").strip()
        if lay.get("hint") is not None:
            h = str(cell(r, "hint") or "").strip()
            if h and h.lower() not in ("-split-", "split", "-split"):
                row["hint"] = h
            else:  # a QuickBooks register line with no single account: QuickBooks has it split
                row["qb_split"] = True
        out.append(row)
    strip_tails(out)
    return out, bad


def strip_tails(rows: List[dict]) -> int:
    """QuickBooks' bank-rule text pasted after the bank's wording comes off (acct_tails.py) -- kept as qb_memo."""
    if not rows:
        return 0
    n = 0
    for r, (clean, memo, why) in zip(rows, acct_tails.detect(rows)):
        if memo:
            r["original"] = r["description"]
            r["description"], r["qb_memo"], r["memo_why"] = clean, memo, why
            n += 1
    return n


def _norm(desc: str) -> str:
    return re.sub(r"\s+", " ", (desc or "").upper()).strip()


def dedupe_keys(rows: List[dict]) -> List[str]:
    """date|amount|description|n: the n-th identical line of the day (two $25 wire fees the same day are two lines)."""
    seen = Counter()
    keys = []
    for r in rows:
        base = f"{r['date'].isoformat()}|{r['amount']:.2f}|{hashlib.sha1(_norm(r['description']).encode()).hexdigest()[:16]}"
        seen[base] += 1
        keys.append(f"{base}|{seen[base]}")
    return keys


def layout_for(rows: List[List], src: Optional[AcctSource], layout: Optional[dict]) -> dict:
    """The columns a person picked, else the ones this account's statements had last time (when this file looks the
    same: header row or not), else a guess."""
    if layout:
        return layout
    saved = json.loads(src.layout) if src and src.layout else None
    if saved and rows:
        has = _header_score(rows[0]) >= 2
        width = max(len(r) for r in rows[:10])
        cols = [v for k, v in saved.items() if k not in ("has_header", "headers") and isinstance(v, int)]
        same_head = not has or saved.get("headers") == [str(c or "").strip().lower() for c in rows[0]]
        if bool(saved.get("has_header")) == has and same_head and all(c < width for c in cols):
            return saved
    return guess_layout(rows)


def _sim_words(desc: str) -> set:
    return set(re.findall(r"[A-Z0-9]+", (desc or "").upper())) - {"REF", "ON", "TO", "FROM", "THE"}


def _refs(desc: str) -> set:
    """The bank's own transaction references (Cash App S-numbers, Zelle REF #, wire TRN#): long codes with digits."""
    return {w for w in re.findall(r"[A-Z0-9]+", (desc or "").upper()) if len(w) >= 8 and re.search(r"\d", w) and re.search(r"[A-Z]", w) or (len(w) >= 10 and w.isdigit())}


def same_line(a_desc: str, b_desc: str, days: int) -> bool:
    """Two lines of the same amount a few days apart: the same transaction? Different bank references = no."""
    ra, rb = _refs(a_desc), _refs(b_desc)
    if ra and rb and not (ra & rb):
        return False
    w1, w2 = _sim_words(a_desc), _sim_words(b_desc)
    sim = len(w1 & w2) / max(1, len(w1 | w2))
    return sim >= (0.7 if days <= 1 else 0.85)


def _toks(desc: str) -> List[str]:
    return re.findall(r"[A-Z0-9]+", (desc or "").upper())


def _tok_same(a: str, b: str) -> bool:
    """Two words of a description: equal, or a masked number on one side (XXXXXXXX7437, OWXXXXXXXX733661) whose
    visible part fits the other."""
    if a == b:
        return True
    for x, y in ((a, b), (b, a)):
        m = re.fullmatch(r"([A-WYZ]*)X{3,}(\d+)", x)
        if m and y.startswith(m.group(1)) and y.endswith(m.group(2)) and y[len(m.group(1)):].isdigit() and len(y) >= len(x) - 2:
            return True
    return False


def starts_same(a: str, b: str, k: int = 6) -> bool:
    """The first k words of the two descriptions match (the shorter one's start, either way round) -- QuickBooks keeps the
    bank's text but sometimes cuts it short and sometimes adds its own memo after it."""
    at, bt = _toks(a), _toks(b)
    n = min(k, len(at), len(bt))
    return n >= 2 and all(_tok_same(at[i], bt[i]) for i in range(n))


def _plain_refs(desc: str) -> set:
    return {w for w in _refs(desc) if "XXX" not in w}


def classify(db: Session, src: Optional[AcctSource], parsed: List[dict], keys: List[str], workbook: bool,
             pool_override: Optional[List[AcctTxn]] = None) -> List[dict]:
    """What happens to each row of a file (rules agreed with the user, tested on their bank + QuickBooks files):
      exact     -- already in: the same line key; or (1) the same start of the description / the same bank reference,
                   same size of amount, within 2 days; or (2a) the same day and the same amount
      possible  -- (2b) the same amount 1-3 days apart (banks post after weekends); (3) the same wording and amount on
                   another date (a mis-dated line) -- a person decides
      expected  -- a line typed in ahead (same amount within 7 days)
      estimate  -- a workbook "Pending" AR / AP estimate (not brought in)
      new       -- the rest
    Each line already in is matched once (two $25 fees the same day are two lines)."""
    sid = src.id if src else None
    pool = pool_override if pool_override is not None else (
        db.query(AcctTxn).options(selectinload(AcctTxn.splits)).filter(AcctTxn.source_id == sid, AcctTxn.expected.is_(False)).all() if sid else [])
    have = {}
    for t in pool:
        if t.dedupe:
            have.setdefault(t.dedupe, t.id)
    expected = db.query(AcctTxn).filter(AcctTxn.source_id == sid, AcctTxn.expected.is_(True)).all() if sid else []
    by_amt = defaultdict(list)
    for t in pool:
        by_amt[round(abs(t.amount), 2)].append(t)
    claimed = set()
    out: List[Optional[dict]] = [None] * len(parsed)

    def cands(r, days):
        return [t for t in by_amt.get(round(abs(r["amount"]), 2), []) if t.id not in claimed and abs((t.date.date() - r["date"]).days) <= days]

    def take(i, ts, status, rule, r):
        t = min(ts, key=lambda t: abs((t.date.date() - r["date"]).days))
        claimed.add(t.id)
        out[i] = {"status": status, "txn_id": t.id, "rule": rule, "sign_differs": (t.amount >= 0) != (r["amount"] >= 0)}

    for i, (r, k) in enumerate(zip(parsed, keys)):
        if workbook and (r.get("category") or "").strip().lower() == "pending":
            out[i] = {"status": "estimate"}
        elif k in have and have[k] not in claimed:
            claimed.add(have[k])
            out[i] = {"status": "exact", "txn_id": have[k], "rule": "same line"}
    for i, r in enumerate(parsed):
        if out[i]:
            continue
        hit = next((e for e in expected if e.id not in claimed and abs(e.amount - r["amount"]) < 0.005 and abs((e.date.date() - r["date"]).days) <= 7), None)
        if hit:
            claimed.add(hit.id)
            out[i] = {"status": "expected", "txn_id": hit.id}
    for i, r in enumerate(parsed):  # rule 1
        if out[i]:
            continue
        c = cands(r, 2)
        rr = _plain_refs(r["description"])
        h = [t for t in c if starts_same(t.bank_description or t.description, r["description"])] or \
            [t for t in c if rr and rr & _plain_refs(t.bank_description or t.description)]
        if h:
            take(i, h, "exact", "same start / reference", r)
    for i, r in enumerate(parsed):  # rule 2a: same day + same amount
        if not out[i]:
            c = [t for t in cands(r, 0) if t.date.date() == r["date"]]
            if c:
                take(i, c, "exact", "same day and amount", r)
    for i, r in enumerate(parsed):  # rule 2b: same amount, 1-3 days apart -> ask
        if not out[i]:
            c = cands(r, 3)
            if c:
                take(i, c, "possible", "same amount, a few days apart", r)
    for i, r in enumerate(parsed):  # rule 3: same wording + amount, another date -> ask
        if not out[i]:
            c = [t for t in by_amt.get(round(abs(r["amount"]), 2), []) if t.id not in claimed and starts_same(t.bank_description or t.description, r["description"], 8)]
            if c:
                take(i, c, "possible", "same wording and amount, another date", r)
    return [o or {"status": "new"} for o in out]


def _in_range(parsed: List[dict], date_from: Optional[date], date_to: Optional[date]) -> Tuple[List[dict], int]:
    keep = [r for r in parsed if (not date_from or r["date"] >= date_from) and (not date_to or r["date"] <= date_to)]
    return keep, len(parsed) - len(keep)


def _books_range(db: Session, source_id: Optional[int]) -> Tuple[Optional[str], Optional[str]]:
    if not source_id:
        return None, None
    lo, hi = db.query(func.min(AcctTxn.date), func.max(AcctTxn.date)).filter(AcctTxn.source_id == source_id, AcctTxn.expected.is_(False)).one()
    return (lo.date().isoformat() if lo else None), (hi.date().isoformat() if hi else None)


def _replace_scope(db: Session, src: AcctSource, parsed: List[dict], date_from: Optional[date] = None, date_to: Optional[date] = None) -> List[AcctTxn]:
    """Replace mode: the lines already in this account for the dates picked (else the file's first to last date) --
    they're what gets replaced."""
    if not parsed:
        return []
    lo, hi = date_from or min(r["date"] for r in parsed), date_to or max(r["date"] for r in parsed)
    return (db.query(AcctTxn).options(selectinload(AcctTxn.splits))
            .filter(AcctTxn.source_id == src.id, AcctTxn.expected.is_(False), AcctTxn.date >= _dt(lo), AcctTxn.date <= _dt(hi)).all())


def preview(db: Session, data: bytes, filename: str, source_id: Optional[int], layout: Optional[dict] = None,
            date_from: Optional[date] = None, date_to: Optional[date] = None, replace: bool = False) -> dict:
    rows = read_table(data, filename)
    src = db.get(AcctSource, source_id) if source_id else None
    lay = layout_for(rows, src, layout)
    allp, bad = parse_rows(rows, lay, src.kind if src else "bank")
    parsed, outside = _in_range(allp, date_from, date_to)
    keys = dedupe_keys(parsed)
    wb = lay.get("account") is not None and lay.get("category") is not None
    scope = _replace_scope(db, src, parsed, date_from, date_to) if replace and src and parsed else None
    cls = classify(db, src, parsed, keys, wb, pool_override=scope) if parsed else []
    st = Counter(c["status"] for c in cls)
    width = max((len(r) for r in rows[:30]), default=0)
    headers = [str(c or f"Column {i + 1}") for i, c in enumerate(rows[0])] if rows and lay.get("has_header") else [f"Column {chr(65 + i)}" for i in range(width)]
    ins = sum(1 for r in parsed if r["amount"] > 0)
    warn = None
    if len(parsed) >= 5 and (ins == 0 or ins == len(parsed)):
        warn = (f"Every line in this file reads as money {'in' if ins else 'out'}. If that's wrong, the amount column is the wrong one -- "
                "pick the money-out column as Debit and the money-in column as Credit below (a QuickBooks register: Amount = out, Deposit = in).")
    books_from, books_to = _books_range(db, source_id)
    matched_ids = {c["txn_id"] for c in cls if c.get("txn_id") and c["status"] in ("exact", "possible")}
    return {"layout": lay, "headers": headers, "sample": [[str(c if c is not None else "") for c in r] for r in rows[:8]],
            "rows": len(parsed), "outside": outside, "already": st["exact"], "possible": st["possible"], "expected": st["expected"],
            "estimate": st["estimate"], "new": st["new"], "bad": bad[:20], "bad_count": len(bad), "warning": warn,
            "books_from": books_from, "books_to": books_to, "replace": bool(replace),
            "replace_keep": len(matched_ids) if replace else 0,
            "replace_unmatched": (len(scope or []) - len(matched_ids)) if replace else 0,
            "workbook": wb, "hints": sum(1 for r in parsed if r.get("hint")),
            "file_first": min((r["date"] for r in allp), default=None) and min(r["date"] for r in allp).isoformat(),
            "file_last": max((r["date"] for r in allp), default=None) and max(r["date"] for r in allp).isoformat(),
            "first": min((r["date"] for r in parsed), default=None) and min(r["date"] for r in parsed).isoformat(),
            "last": max((r["date"] for r in parsed), default=None) and max(r["date"] for r in parsed).isoformat(),
            "money_in": cents(sum(r["amount"] for r in parsed if r["amount"] > 0)),
            "money_out": cents(sum(r["amount"] for r in parsed if r["amount"] < 0)),
            "parsed": [{"date": r["date"].isoformat(), "amount": r["amount"], "description": r["description"], "hint": r.get("hint"),
                        "qb_memo": r.get("qb_memo"), "qb_split": r.get("qb_split"),
                        "account": r.get("account"), "category": r.get("category"), "status": c["status"], "rule": c.get("rule")}
                       for r, c in list(zip(parsed, cls))[:12]],
            "memos": sum(1 for r in parsed if r.get("qb_memo")), "qb_splits": sum(1 for r in parsed if r.get("qb_split")),
            "memo_examples": [{"date": r["date"].isoformat(), "amount": r["amount"], "kept": r["description"], "removed": r["qb_memo"]}
                              for r in parsed if r.get("qb_memo")][:8],
            "split_lines": [{"date": r["date"].isoformat(), "amount": r["amount"], "description": r["description"]} for r in parsed if r.get("qb_split")][:20],
            "remembered": len(memory_hits(db, src.id if src else None, [r for r, c in zip(parsed, cls) if c["status"] == "new"])) if src else 0}


def _snapshot(t: AcctTxn) -> dict:
    """A line as it was (for Undo of a replace)."""
    cols = [c.name for c in AcctTxn.__table__.columns]
    d = {c: getattr(t, c) for c in cols}
    for k, v in list(d.items()):
        if isinstance(v, datetime):
            d[k] = v.isoformat()
    d["_splits"] = [{"account_id": x.account_id, "amount": x.amount, "note": x.note} for x in t.splits]
    return d


def _restore(db: Session, d: dict) -> None:
    splits = d.pop("_splits", [])
    for k in ("date", "created_at", "updated_at", "sent_back_at"):
        if d.get(k):
            d[k] = datetime.fromisoformat(d[k])
    t = AcctTxn(**d)
    db.add(t)
    db.flush()
    for x in splits:
        db.add(AcctSplit(txn_id=t.id, **x))


FILING = ("account_id", "how", "reviewed", "note", "rule_id", "corrected", "no_learn", "sent_back", "sent_back_reason", "sent_back_from",
          "sent_back_by", "sent_back_at", "external_id", "flag", "flag_info")


def _qb_split_flag():
    return "qb_split", json.dumps({"reason": "QuickBooks has this line split over several accounts (no single account) -- split it here too"})


def import_file(db: Session, data: bytes, filename: str, source_id: int, layout: Optional[dict], by: str,
                date_from: Optional[date] = None, date_to: Optional[date] = None, replace: bool = False) -> dict:
    """Bring a statement (or the user's categorized workbook) in. Lines already in are left out (Ignored); a line typed in
    ahead ("expected") with the same amount within 7 days becomes this line instead of a second copy. date_from / date_to:
    only the lines in those dates. replace: the file's lines REPLACE what's in the books for its dates -- each line already
    in that matches keeps how it was filed (account, split, note ...) on the file's line; the ones with no match on the
    file move to Book Entries (nothing is lost; Undo puts everything back)."""
    src = db.get(AcctSource, source_id)
    if not src:
        raise HTTPException(status_code=400, detail="Pick which account the statement is for")
    rows = read_table(data, filename)
    lay = layout_for(rows, src, layout)
    if lay.get("date") is None or lay.get("description") is None or (lay.get("amount") is None and lay.get("debit") is None):
        raise HTTPException(status_code=400, detail="Say which columns are the date, the description and the amount")
    allp, bad = parse_rows(rows, lay, src.kind)
    parsed, outside = _in_range(allp, date_from, date_to)
    if not parsed:
        raise HTTPException(status_code=400, detail="No lines with a date and an amount in this file" + (" for those dates" if outside else ""))
    workbook = lay.get("account") is not None and lay.get("category") is not None
    keys = dedupe_keys(parsed)
    scope = _replace_scope(db, src, parsed, date_from, date_to) if replace else None
    cls = classify(db, src, parsed, keys, workbook, pool_override=scope)
    imp = AcctImport(source_id=src.id, filename=filename, kind="workbook" if workbook else "statement", rows=len(parsed), created_by=by)
    db.add(imp)
    db.flush()
    new = []
    possible = estimates = replaced = 0
    snapshots, keep_ids = [], set()
    if replace:
        old = {t.id: t for t in scope}
        for r, k, c in zip(parsed, keys, cls):
            t_old = old.get(c.get("txn_id")) if c["status"] in ("exact", "possible") else None
            if t_old is None:
                continue
            snapshots.append(_snapshot(t_old))
            t = AcctTxn(source_id=src.id, date=_dt(r["date"]), amount=r["amount"], description=r["description"],
                        payee=t_old.payee if t_old.origin == "manual" else acct_mapper.payee_label(r["description"]),
                        origin="import", import_id=imp.id, dedupe=k, hint=r.get("hint"), qb_memo=r.get("qb_memo"), created_by=by, updated_by=by)
            for f in FILING:
                setattr(t, f, getattr(t_old, f))
            if r.get("qb_split") and not t_old.splits:
                t.flag, t.flag_info = _qb_split_flag()
            if t.how == "workbook":
                t.how = "person"  # a person's filing (from their workbook) -- still teaches AT-HUB
            for x in t_old.splits:
                t.splits.append(AcctSplit(account_id=x.account_id, amount=cents(x.amount * (r["amount"] / t_old.amount)) if t_old.amount else x.amount, note=x.note))
            db.add(t)
            db.flush()
            db.query(AcctCheckOk).filter(AcctCheckOk.txn_id == t_old.id).update({AcctCheckOk.txn_id: t.id}, synchronize_session=False)
            db.query(AcctIgnored).filter(AcctIgnored.dup_of_txn_id == t_old.id).update({AcctIgnored.dup_of_txn_id: t.id}, synchronize_session=False)
            db.delete(t_old)
            keep_ids.add(t_old.id)
            replaced += 1
            c["status"] = "replaced"
        book = book_source(db)
        moved = []
        for t_old in scope:  # in the books for these dates but not on the file: kept, as book entries
            if t_old.id in keep_ids:
                continue
            moved.append([t_old.id, t_old.source_id])
            t_old.source_id = book.id
            t_old.note = ((t_old.note + " · ") if t_old.note else "") + f"not on {filename}"
        imp.replaced = json.dumps(snapshots, default=str)
        imp.moved = json.dumps(moved)
        db.flush()
    for r, k, c in zip(parsed, keys, cls):
        st = c["status"]
        if st == "replaced":
            continue
        if st in ("exact", "possible", "estimate"):  # not brought in -- kept on the Ignored list (possible: a person decides)
            db.add(AcctIgnored(import_id=imp.id, source_id=src.id, date=_dt(r["date"]), amount=r["amount"], description=r["description"],
                               extra=json.dumps({x: r.get(x) for x in ("account", "category") if r.get(x)}) or None, dedupe=k,
                               reason={"exact": "exact", "possible": "near", "estimate": "estimate"}[st],
                               status="possible" if st == "possible" else "ignored", dup_of_txn_id=c.get("txn_id")))
            if st == "exact":
                imp.skipped += 1
            elif st == "possible":
                possible += 1
            else:
                estimates += 1
            continue
        if st == "expected":  # typed in earlier -- now it's on the statement
            hit = db.get(AcctTxn, c["txn_id"])
            hit.expected, hit.bank_description, hit.dedupe, hit.date = False, r["description"], k, _dt(r["date"])
            hit.updated_by = by
            imp.matched += 1
            continue
        t = AcctTxn(source_id=src.id, date=_dt(r["date"]), amount=r["amount"], description=r["description"],
                    payee=acct_mapper.payee_label(r["description"]), origin="workbook" if workbook else "import",
                    import_id=imp.id, dedupe=k, hint=r.get("hint"), qb_memo=r.get("qb_memo"), created_by=by)
        if r.get("qb_split"):
            t.flag, t.flag_info = _qb_split_flag()
        parts = workbook_parts(r.get("account") or "", r["amount"]) if workbook else None
        if parts:  # "Niraj 71472; Richard 69259" in their Account column: one line, split
            kind = WORKBOOK_KINDS.get((r.get("category") or "").strip().lower(), "other")
            for name, amt in parts:
                acc = account_by_name(db, name, kind, name if kind == "owner" else None)
                t.splits.append(AcctSplit(account_id=acc.id, amount=amt))
            t.how, t.reviewed = "workbook", True
        elif workbook and r.get("account"):
            kind = WORKBOOK_KINDS.get((r.get("category") or "").strip().lower(), "other")
            name = " ".join(w if w.isupper() and len(w) > 1 else w.capitalize() for w in r["account"].split())  # "factoring fee WF" -> "Factoring Fee WF"
            partner = name if kind == "owner" else None
            acc = account_by_name(db, name, kind, partner)
            t.account_id, t.how, t.reviewed = acc.id, "workbook", True
        db.add(t)
        new.append(t)
    db.flush()
    imp.added = len(new)
    remembered = 0
    if not workbook:
        remembered = apply_memory(db, src.id, new, by)
        imp.auto = autofill(db, new)
    if not workbook:  # a statement's columns are kept for next time (the workbook is a one-off)
        src.layout = json.dumps(lay)
    db.commit()
    return {"import_id": imp.id, "rows": imp.rows, "added": imp.added, "skipped": imp.skipped, "matched": imp.matched,
            "possible": possible, "estimates": estimates, "auto": imp.auto, "bad": bad[:20], "bad_count": len(bad),
            "outside": outside, "replaced": replaced, "moved_to_book": len(json.loads(imp.moved or "[]")),
            "left": sum(1 for t in new if not t.account_id and not t.splits), "remembered": remembered,
            "memos": sum(1 for r in parsed if r.get("qb_memo")), "qb_splits": sum(1 for r in parsed if r.get("qb_split")),
            "flagged": sum(1 for t in new if t.flag)}


def workbook_parts(cell: str, amount: float) -> Optional[List[Tuple[str, float]]]:
    """'Niraj 71472; Richard 69259' (or with $ / commas, ; or |) -> [(name, signed amount)] when they add up to the line."""
    if not re.search(r"[;|]", cell or ""):
        return None
    out = []
    for part in re.split(r"[;|]", cell):
        m = re.match(r"\s*(.+?)\s+\$?(-?[\d,]+(?:\.\d+)?)\s*$", part)
        if not m:
            return None
        v = abs(float(m.group(2).replace(",", ""))) * (-1 if amount < 0 else 1)
        out.append((" ".join(w if w.isupper() and len(w) > 1 else w.capitalize() for w in m.group(1).split()), cents(v)))
    if len(out) < 2 or abs(sum(a for _n, a in out) - amount) > 0.01:
        return None
    return out


def undo_import(db: Session, import_id: int) -> int:
    imp = db.get(AcctImport, import_id)
    if not imp:
        raise HTTPException(status_code=404, detail="Import not found")
    n = 0
    for t in db.query(AcctTxn).filter(AcctTxn.import_id == import_id).all():
        db.delete(t)
        n += 1
    db.query(AcctIgnored).filter(AcctIgnored.import_id == import_id).delete()
    db.flush()
    for d in json.loads(imp.replaced or "[]"):  # a replace: the lines it replaced come back as they were
        _restore(db, d)
    for tid, sid in json.loads(imp.moved or "[]"):
        t = db.get(AcctTxn, tid)
        if t:
            t.source_id = sid
            t.note = re.sub(r"( · )?not on [^·]+$", "", t.note or "").strip() or None
    db.delete(imp)
    db.commit()
    return n


# ---------------------------------------------------------------- lines
def txn_dict(t: AcctTxn, accounts: Dict[int, AcctAccount], sources: Dict[int, AcctSource]) -> dict:
    a = accounts.get(t.account_id)
    s = sources.get(t.source_id)
    return {"id": t.id, "date": t.date.date().isoformat(), "amount": t.amount, "description": t.description,
            "bank_description": t.bank_description, "payee": t.payee, "account_id": t.account_id,
            "account": path(a, accounts) if a else None, "kind": a.kind if a else None, "how": t.how, "reviewed": t.reviewed,
            "note": t.note, "origin": t.origin, "expected": t.expected, "source_id": t.source_id, "hint": t.hint,
            "source": s.name if s else "", "import_id": t.import_id,
            "splits": [{"id": x.id, "account_id": x.account_id, "account": path(accounts[x.account_id], accounts) if x.account_id in accounts else "",
                        "kind": accounts[x.account_id].kind if x.account_id in accounts else None, "amount": x.amount, "note": x.note} for x in t.splits],
            "created_by": t.created_by, "updated_by": t.updated_by, "rule_id": t.rule_id,
            "sent_back": bool(t.sent_back), "sent_back_reason": t.sent_back_reason, "sent_back_by": t.sent_back_by,
            "sent_back_at": t.sent_back_at, "flag": t.flag, "flag_info": json.loads(t.flag_info) if t.flag_info else None,
            "qb_memo": t.qb_memo}


def maps(db: Session):
    return ({a.id: a for a in db.query(AcctAccount).all()}, {s.id: s for s in db.query(AcctSource).all()})


def set_account(db: Session, t: AcctTxn, account_id: Optional[int], splits: Optional[List[dict]], by: str) -> None:
    if t.how in ("auto", "rule") and not t.reviewed and t.account_id and (splits or account_id != t.account_id):
        t.corrected = True  # a person moved AT-HUB's pick: counts double when learning
    t.splits.clear()
    db.flush()
    if splits:
        total = cents(sum(float(x["amount"]) for x in splits))
        if abs(total - t.amount) > 0.005:
            raise HTTPException(status_code=400, detail=f"The parts add up to {total:,.2f}, the line is {t.amount:,.2f}")
        for x in splits:
            if not db.get(AcctAccount, int(x["account_id"])):
                raise HTTPException(status_code=400, detail="Pick an account for every part")
            t.splits.append(AcctSplit(account_id=int(x["account_id"]), amount=cents(float(x["amount"])), note=(x.get("note") or None)))
        t.account_id = None
    else:
        if account_id and not db.get(AcctAccount, account_id):
            raise HTTPException(status_code=400, detail="Account not found")
        t.account_id = account_id or None
    t.how = "person" if (account_id or splits) else None
    t.reviewed = True
    t.rule_id = None
    if account_id or splits:
        t.sent_back = None
        t.flag = t.flag_info = None
    t.updated_by = by
    if splits:
        remember(db, t, "split", by)
    else:
        forget_memory(db, t, "split")


# ---------------------------------------------------------------- reports
def lines(db: Session, start: Optional[date] = None, end: Optional[date] = None, source_id: Optional[int] = None,
          set_aside: bool = False) -> List[dict]:
    """Every line, a split line counted once per part: {date, amount, account_id, kind, source_id, txn_id}. The
    year-end AR / AP estimates (kind "pending") are set aside -- in no report total -- unless set_aside=True."""
    q = db.query(AcctTxn).options(selectinload(AcctTxn.splits))
    if start:
        q = q.filter(AcctTxn.date >= _dt(start))
    if end:
        q = q.filter(AcctTxn.date <= _dt(end))
    if source_id:
        q = q.filter(AcctTxn.source_id == source_id)
    accounts = {a.id: a for a in db.query(AcctAccount).all()}
    out = []
    for t in q.all():
        parts = [(x.account_id, x.amount) for x in t.splits] or [(t.account_id, t.amount)]
        for acc, amt in parts:
            a = accounts.get(acc)
            if a and a.kind == "pending" and not set_aside:
                continue
            out.append({"date": t.date.date(), "amount": amt, "account_id": acc, "kind": a.kind if a else None,
                        "source_id": t.source_id, "txn_id": t.id, "expected": t.expected})
    return out


def _bucket(d: date, by: str, fy_month: int) -> str:
    if by == "month":
        return f"{d.year}-{d.month:02d}"
    if by == "quarter":
        s = fy_start(d, fy_month)
        q = ((d.year - s.year) * 12 + d.month - s.month) // 3 + 1
        return f"{fy_label(s, fy_month)} Q{q}"
    return fy_label(fy_start(d, fy_month), fy_month)


def _buckets(start: date, end: date, by: str, fy_month: int) -> List[str]:
    out, d = [], date(start.year, start.month, 1)
    while d <= end:
        b = _bucket(d, by, fy_month)
        if b not in out:
            out.append(b)
        d = date(d.year + (d.month == 12), d.month % 12 + 1, 1)
    return out


def pnl(db: Session, start: date, end: date, by: str = "month") -> dict:
    """Income, cost of goods, expenses by account and period; profit; what's below the line (loans, draws, wash)."""
    m = settings(db)["fy_start_month"]
    cols = _buckets(start, end, by, m)
    accounts = {a.id: a for a in db.query(AcctAccount).all()}
    cells = defaultdict(lambda: defaultdict(float))
    unfiled = defaultdict(float)
    for l in lines(db, start, end):
        b = _bucket(l["date"], by, m)
        if l["account_id"] is None:
            unfiled[b] += l["amount"]
        else:
            cells[l["account_id"]][b] += l["amount"]
    sections = []
    for kind in ("income", "cogs", "expense", "loan", "owner", "transfer", "pending", "other"):
        rows = tree_rows([a for a in accounts.values() if a.kind == kind], accounts, cells, cols, kind)
        tot = [cents(sum(cells[aid].get(c, 0.0) for aid in cells if aid in accounts and accounts[aid].kind == kind)) for c in cols]
        sections.append({"kind": kind, "label": KINDS[kind][0], "in_pnl": KINDS[kind][2], "rows": rows, "totals": tot,
                         "total": cents(sum(tot))})
    by_kind = {s["kind"]: s for s in sections}
    profit = [cents(by_kind["income"]["totals"][i] + by_kind["cogs"]["totals"][i] + by_kind["expense"]["totals"][i]) for i in range(len(cols))]
    gross = [cents(by_kind["income"]["totals"][i] + by_kind["cogs"]["totals"][i]) for i in range(len(cols))]
    return {"columns": cols, "sections": sections, "gross": gross, "gross_total": cents(sum(gross)), "profit": profit,
            "profit_total": cents(sum(profit)), "unfiled": [cents(unfiled.get(c, 0.0)) for c in cols],
            "unfiled_total": cents(sum(unfiled.values())), "set_aside": set_aside(db, start, end)}


def tree_rows(accs: List[AcctAccount], accounts: Dict[int, AcctAccount], cells, cols, kind: str) -> List[dict]:
    """P&L rows for one section: accounts with money in the period (or a sub-account that has), a parent's figures
    including its sub-accounts; when a parent also has lines of its own, they get their own indented row."""
    own = {a.id: [cents(cells.get(a.id, {}).get(c, 0.0)) for c in cols] for a in accs}
    roll = {}
    for a in accs:
        ids = [a.id] + [d for d in descendants(a.id, accounts) if d in own]
        roll[a.id] = [cents(sum(own[i][k] for i in ids)) for k in range(len(cols))]
    used = {a.id for a in accs if any(roll[a.id])}
    sign = -1 if kind == "income" else 1
    out = []
    for a, d in tree_order([a for a in accs if a.id in used], accounts, key=lambda x: (sign * sum(roll[x.id]), x.name.lower())):
        kids = [k for k in descendants(a.id, accounts) if k in used]
        out.append({"account_id": a.id, "name": a.name, "path": path(a, accounts), "depth": d, "has_children": bool(kids),
                    "cells": roll[a.id], "total": cents(sum(roll[a.id]))})
        if kids and any(own[a.id]):
            out.append({"account_id": a.id, "name": f"{a.name} (Other)", "path": path(a, accounts), "depth": d + 1,
                        "has_children": False, "own_only": True, "cells": own[a.id], "total": cents(sum(own[a.id]))})
    return out


def set_aside(db: Session, start: date, end: date) -> dict:
    """The year-end AR / AP estimates in the period: kept, shown on their own, counted nowhere."""
    ls = [l for l in lines(db, start, end, set_aside=True) if l["kind"] == "pending"]
    return {"count": len(ls), "total": cents(sum(l["amount"] for l in ls))}


def overview(db: Session, start: date, end: date) -> dict:
    ls = lines(db, start, end)
    accounts = {a.id: a for a in db.query(AcctAccount).all()}
    months = _buckets(start, end, "month", 1)
    inflow, outflow = defaultdict(float), defaultdict(float)
    kinds = defaultdict(float)
    for l in ls:
        b = _bucket(l["date"], "month", 1)
        (inflow if l["amount"] > 0 else outflow)[b] += l["amount"]
        kinds[l["kind"] or "unfiled"] += l["amount"]
    payees = defaultdict(float)
    for l in ls:
        if l["amount"] < 0 and l["account_id"] and l["kind"] in ("cogs", "expense"):
            payees[l["account_id"]] += l["amount"]
    top = sorted(payees.items(), key=lambda x: x[1])[:8]
    income = kinds.get("income", 0.0)
    cogs = kinds.get("cogs", 0.0)
    expense = kinds.get("expense", 0.0)
    unfiled = db.query(AcctTxn).filter(AcctTxn.account_id.is_(None), ~AcctTxn.splits.any()).count()
    unreviewed = db.query(AcctTxn).filter(AcctTxn.reviewed.is_(False)).count()
    sources = [source_dict(db, s) for s in db.query(AcctSource).filter(AcctSource.is_active.is_(True)).order_by(AcctSource.id).all()]
    return {"months": [{"month": m_, "in": cents(inflow.get(m_, 0.0)), "out": cents(outflow.get(m_, 0.0))} for m_ in months],
            "kinds": {k: cents(v) for k, v in kinds.items()}, "income": cents(income), "cogs": cents(cogs), "expense": cents(expense),
            "profit": cents(income + cogs + expense), "money_in": cents(sum(inflow.values())), "money_out": cents(sum(outflow.values())),
            "top_spend": [{"account_id": a, "name": accounts[a].name if a in accounts else "?", "kind": accounts[a].kind if a in accounts else None,
                           "amount": cents(v)} for a, v in top],
            "unfiled": unfiled, "unreviewed": unreviewed, "sources": sources, "lines": len(ls)}


def account_summary(db: Session, start: date, end: date) -> List[dict]:
    acc = defaultdict(lambda: {"count": 0, "in": 0.0, "out": 0.0, "last": None})
    for l in lines(db, start, end, set_aside=True):
        x = acc[l["account_id"]]
        x["count"] += 1
        x["in" if l["amount"] > 0 else "out"] += l["amount"]
        x["last"] = max(x["last"], l["date"]) if x["last"] else l["date"]
    accounts = {a.id: a for a in db.query(AcctAccount).all()}
    empty = {"count": 0, "in": 0.0, "out": 0.0, "last": None}
    out = []
    for a, d in tree_order(list(accounts.values()), accounts):
        ids = [a.id] + descendants(a.id, accounts)
        xs = [acc.get(i, empty) for i in ids]
        last = max((x["last"] for x in xs if x["last"]), default=None)
        x = acc.get(a.id, empty)
        out.append({**account_dict(a, accounts), "count": sum(y["count"] for y in xs), "in": cents(sum(y["in"] for y in xs)),
                    "out": cents(sum(y["out"] for y in xs)), "net": cents(sum(y["in"] + y["out"] for y in xs)),
                    "own_net": cents(x["in"] + x["out"]), "subs": len(ids) - 1, "last": last.isoformat() if last else None})
    return out


def loans(db: Session, end: Optional[date] = None) -> List[dict]:
    """Per lender: borrowed (money in), repaid (out), what's still owed (opening + borrowed - repaid), all time to `end`."""
    out = []
    for a in db.query(AcctAccount).filter(AcctAccount.kind == "loan").order_by(AcctAccount.name).all():
        ls = [l for l in lines(db, None, end) if l["account_id"] == a.id]
        got = sum(l["amount"] for l in ls if l["amount"] > 0)
        paid = sum(l["amount"] for l in ls if l["amount"] < 0)
        out.append({"account_id": a.id, "name": a.name, "opening": a.opening_balance or 0.0, "borrowed": cents(got),
                    "repaid": cents(paid), "owed": cents((a.opening_balance or 0.0) + got + paid), "count": len(ls),
                    "last": max((l["date"] for l in ls), default=None)})
    return out


# ---------------------------------------------------------------- AP / AR
def _names(db: Session, model, ids) -> Dict[int, str]:
    ids = {i for i in ids if i}
    return {r.id: r.name for r in db.query(model).filter(model.id.in_(ids)).all()} if ids else {}


def open_items(db: Session) -> dict:
    """AP / AR now: AT-HUB's own figures (unpaid invoices, invoices not sent, POs not paid, open customer orders) with
    a person's word on top ("already paid", a different amount, a note), plus items typed in."""
    over = {(o.source, o.source_id): o for o in db.query(AcctOpenItem).filter(AcctOpenItem.source != "manual").all()}
    now = today()
    items = []

    def add(side, group, source, sid, party, ref, amount, d, due, link=None):
        o = over.get((source, sid))
        row = {"key": f"{source}:{sid}", "side": side, "group": group, "source": source, "source_id": sid, "party": party,
               "ref": ref, "athub_amount": cents(amount), "amount": cents(o.amount if o and o.amount is not None else amount),
               "date": d.date().isoformat() if d else None, "due_date": due.date().isoformat() if due else None,
               "overdue": bool(due and due.date() < now), "status": o.status if o else "open",
               "paid_date": o.paid_date.date().isoformat() if o and o.paid_date else None, "note": o.note if o else None,
               "override_id": o.id if o else None, "changed_by": o.updated_by or o.created_by if o else None, "link": link}
        items.append(row)

    invs = db.query(Invoice).options(selectinload(Invoice.lines), selectinload(Invoice.payments)).filter(Invoice.status.in_(("draft", "sent", "paid"))).all()
    cust = _names(db, Customer, [i.customer_id for i in invs])
    for i in invs:
        if i.status in ("sent", "paid") and i.balance > 0.005:
            add("ar", "Unpaid Invoices", "invoice", i.id, cust.get(i.customer_id, ""), i.code, i.balance, i.invoice_date, i.due_date, f"invoices.html?id={i.id}")
        elif i.status == "draft" and i.total > 0.005:
            add("ar", "Invoices Not Sent", "invoice", i.id, cust.get(i.customer_id, ""), i.code, i.total, i.invoice_date, None, f"invoices.html?id={i.id}")
    orders = db.query(CustomerOrder).options(selectinload(CustomerOrder.lines)).filter(CustomerOrder.status == "confirmed").all()
    cust2 = _names(db, Customer, [o.customer_id for o in orders])
    for o in orders:
        left = sum(line_amount(max(0, l.quantity - min(l.shipped_quantity or 0, l.quantity)), l.unit_price) for l in o.lines)
        if left > 0.005:
            add("ar", "Open Customer Orders", "order", o.id, cust2.get(o.customer_id, ""), o.code + (f" · PO {o.po_number}" if o.po_number else ""),
                left, o.created_at, None, f"customer-orders.html?id={o.id}")
    from app.services.crud import po_left_to_pay
    pos = (db.query(PurchaseOrder).options(selectinload(PurchaseOrder.lines), selectinload(PurchaseOrder.payments), selectinload(PurchaseOrder.bills),
                                           selectinload(PurchaseOrder.charges))
           .filter(PurchaseOrder.status.notin_(("draft", "validation", "cancelled"))).all())
    vend = _names(db, Vendor, [p.vendor_id for p in pos])
    for p in pos:
        left = po_left_to_pay(p)
        if left > 0.005:
            due = min((b.due_date for b in p.bills if b.balance > 0.005 and b.due_date), default=None)
            add("ap", "Purchase Orders Not Paid", "po", p.id, vend.get(p.vendor_id, ""), p.code + (f" · SO {p.vendor_so_number}" if p.vendor_so_number else ""),
                left, p.order_date, due, f"purchase-orders.html?id={p.id}")
    for o in db.query(AcctOpenItem).filter(AcctOpenItem.source == "manual").order_by(AcctOpenItem.id).all():
        items.append({"key": f"manual:{o.id}", "side": o.side, "group": "Typed In", "source": "manual", "source_id": None, "id": o.id,
                      "party": o.party, "ref": o.ref, "athub_amount": None, "amount": cents(o.amount or 0.0),
                      "date": o.item_date.date().isoformat() if o.item_date else None,
                      "due_date": o.due_date.date().isoformat() if o.due_date else None,
                      "overdue": bool(o.due_date and o.due_date.date() < now and o.status == "open"), "status": o.status,
                      "paid_date": o.paid_date.date().isoformat() if o.paid_date else None, "note": o.note, "override_id": o.id,
                      "changed_by": o.updated_by or o.created_by, "link": None})
    counted = settings(db)["apar_groups"]
    def counts(r):
        return r["status"] == "open" and (r["group"] != "Open Customer Orders" or counted.get("open_orders"))
    tot = {side: cents(sum(r["amount"] for r in items if r["side"] == side and counts(r))) for side in ("ap", "ar")}
    groups = defaultdict(lambda: {"count": 0, "amount": 0.0})
    for r in items:
        if r["status"] == "open":
            g = groups[(r["side"], r["group"])]
            g["count"] += 1
            g["amount"] = cents(g["amount"] + r["amount"])
    return {"items": items, "ap": tot["ap"], "ar": tot["ar"], "open_orders_counted": bool(counted.get("open_orders")),
            "groups": [{"side": s, "group": g, **v} for (s, g), v in sorted(groups.items())]}


def set_open_item(db: Session, data: dict, by: str) -> AcctOpenItem:
    """A typed-in AP / AR item (new or changed), or a person's word on an AT-HUB figure (source + source_id)."""
    src = data.get("source") or "manual"
    o = None
    if data.get("id"):
        o = db.get(AcctOpenItem, int(data["id"]))
    elif src != "manual":
        o = db.query(AcctOpenItem).filter(AcctOpenItem.source == src, AcctOpenItem.source_id == data.get("source_id")).first()
    if not o:
        if src == "manual" and not data.get("side") in ("ap", "ar"):
            raise HTTPException(status_code=400, detail="Payable (we owe) or receivable (owed to us)?")
        o = AcctOpenItem(source=src, source_id=data.get("source_id"), side=data.get("side") or ("ar" if src in ("invoice", "order") else "ap"),
                         created_by=by)
        db.add(o)
    for k in ("party", "ref", "note"):
        if k in data:
            setattr(o, k, (data[k] or None))
    if "side" in data and data["side"] in ("ap", "ar"):
        o.side = data["side"]
    if "amount" in data:
        o.amount = cents(float(data["amount"])) if data["amount"] not in (None, "") else None
    for k in ("item_date", "due_date", "paid_date"):
        if k in data:
            d = parse_day(data[k])
            setattr(o, k, _dt(d) if d else None)
    if "status" in data:
        if data["status"] not in ("open", "paid", "left_out"):
            raise HTTPException(status_code=400, detail="Status is open, paid or left out")
        o.status = data["status"]
        if o.status == "paid" and not o.paid_date:
            o.paid_date = _dt(today())
        if o.status == "open":
            o.paid_date = None
    if src == "manual" and (o.amount is None or not (o.party or o.ref)):
        raise HTTPException(status_code=400, detail="Who and how much")
    o.updated_by = by
    db.commit()
    return o


# ---------------------------------------------------------------- partners
def partners(db: Session, start: date, end: date, fy_key: str) -> dict:
    """The user's year-end split (their Summary sheet): to share = money in - cost of goods - expenses - what was
    borrowed (net of repayments: borrowed money in the bank isn't profit) + the bank balance counted; each partner's
    share, what they already took, what's left. AR / AP estimates are set aside, not counted (user, 2026-10-10)."""
    s = settings(db)
    yr = s.get("years", {}).get(fy_key, {})
    p = pnl(db, start, end, "fy")
    k = {x["kind"]: x["total"] for x in p["sections"]}
    bank = float(yr.get("bank_balance") or 0.0)
    steps = [("Money In (Payments)", k.get("income", 0.0)), ("Cost Of Goods", k.get("cogs", 0.0)), ("Expenses", k.get("expense", 0.0)),
             ("Less Loans (Borrowed − Repaid)", -k.get("loan", 0.0)), ("Bank Balance Counted", bank)]
    total = cents(sum(v for _l, v in steps))
    accounts = {a.id: a for a in db.query(AcctAccount).all()}
    taken = defaultdict(float)
    for l in lines(db, start, end):
        if l["kind"] == "owner":
            a = accounts.get(l["account_id"])
            who = next((x.partner for x in ([a] + ancestors(a, accounts)) if x.partner), None) if a else None
            top = ([a] + ancestors(a, accounts))[-1] if a else None
            taken[(who or (top.name if top else "?")).strip().lower()] += l["amount"]
    people = []
    adj = yr.get("adjust", {})
    for pt in s["partners"]:
        share = cents(total * float(pt["share"]))
        t = cents(taken.get(pt["name"].strip().lower(), 0.0))
        a = float(adj.get(pt["name"], 0.0) or 0.0)
        people.append({"name": pt["name"], "share_pct": float(pt["share"]), "share": share, "taken": t, "adjust": a,
                       "remaining": cents(share + t + a)})
    known = {pt["name"].strip().lower() for pt in s["partners"]}
    others = [{"name": n, "taken": cents(v)} for n, v in taken.items() if n not in known]
    return {"steps": [{"label": l, "amount": cents(v)} for l, v in steps], "total": total, "partners": people, "others": others,
            "bank_balance": bank, "notes": yr.get("notes", ""), "operating_profit": p["profit_total"],
            "set_aside": p["set_aside"], "unfiled": p["unfiled_total"]}


def cashflow(db: Session, start: date, end: date, source_id: Optional[int] = None) -> dict:
    """By month: in / out by kind and the running balance (from the bank account's opening balance when it's set)."""
    months = _buckets(start, end, "month", 1)
    by = defaultdict(lambda: defaultdict(float))
    for l in lines(db, start, end, source_id):
        if l["expected"]:
            continue
        by[_bucket(l["date"], "month", 1)][l["kind"] or "unfiled"] += l["amount"]
    srcs = [db.get(AcctSource, source_id)] if source_id else db.query(AcctSource).filter(AcctSource.kind != "book").all()
    opening = 0.0
    for s in srcs:
        if not s:
            continue
        opening += s.opening_balance or 0.0
        q = db.query(AcctTxn.amount).filter(AcctTxn.source_id == s.id, AcctTxn.date < _dt(start), AcctTxn.expected.is_(False))
        if s.opening_date:
            q = q.filter(AcctTxn.date >= s.opening_date)
        opening += sum(a for (a,) in q.all())
    run = opening
    rows = []
    for m_ in months:
        k = by.get(m_, {})
        net = sum(k.values())
        run += net
        rows.append({"month": m_, "by_kind": {x: cents(v) for x, v in k.items()}, "in": cents(sum(v for v in k.values() if v > 0)),
                     "net": cents(net), "balance": cents(run)})
    return {"opening": cents(opening), "rows": rows}


# ---------------------------------------------------------------- Excel: the analysis, laid out like the user's own output
def export_xlsx(db: Session, start: date, end: date, fy_key: str) -> bytes:
    """Like their FY2024-25_Auto_Analyzed.xlsx (made by their "ATind Analysis Script.py"): All Transactions, a tab per
    category (its lines, then the total per account; Loans also money out / in per lender), Summary (count and total
    per category, then the split). The AR / AP estimates go on "Not Included" and are counted nowhere."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    accounts, _sources = maps(db)
    bold = Font(bold=True)
    money_fmt = '#,##0.00;[Red]-#,##0.00'
    head = ["Date ", "Amount", "Description", "Account", "Category", "Note"]
    txns = (db.query(AcctTxn).options(selectinload(AcctTxn.splits)).filter(AcctTxn.date >= _dt(start), AcctTxn.date <= _dt(end))
            .order_by(AcctTxn.date, AcctTxn.id).all())
    rows, aside = [], []
    for t in txns:
        parts = [(x.account_id, x.amount, x.note) for x in t.splits] or [(t.account_id, t.amount, None)]
        for aid, amt, note in parts:
            a = accounts.get(aid)
            r = [t.date, amt, t.bank_description or t.description, path(a, accounts) if a else "",
                 CATEGORY.get(a.kind, "Other") if a else "Not Filed", " · ".join(x for x in (note, t.note) if x)]
            (aside if a and a.kind == "pending" else rows).append(r)

    def sheet(ws, data):
        ws.append(head)
        for c in ws[ws._current_row]:
            c.font = bold
        for r in data:
            ws.append(r)
            ws.cell(ws._current_row, 1).number_format = "yyyy-mm-dd"
            ws.cell(ws._current_row, 2).number_format = money_fmt
        for col, w in zip("ABCDEF", (13, 14, 80, 22, 12, 30)):
            ws.column_dimensions[col].width = w

    def totals(ws, data, loans=False):
        tot = defaultdict(float)
        for r in data:
            tot[r[3]] += r[1]
        ws.append([])
        ws.append(["Account", "Total Amount"])
        for c in ws[ws._current_row][:2]:
            c.font = bold
        for n in sorted(tot, key=str.lower):
            ws.append([n, cents(tot[n])])
            ws.cell(ws._current_row, 2).number_format = money_fmt
        parents = sorted({n.split(SEP)[0] for n in tot if SEP in n}, key=str.lower)
        for top in parents:  # Labor (All) = Labor + Labor › Faustino + ...
            ws.append([f"{top} (All)", cents(sum(v for n, v in tot.items() if n == top or n.startswith(top + SEP)))])
            ws.cell(ws._current_row, 1).font = bold
            ws.cell(ws._current_row, 2).number_format = money_fmt
        if loans:
            ws.append([])
            ws.append(["Account", "Total_Outgoing", "Total_Incoming"])
            for c in ws[ws._current_row][:3]:
                c.font = bold
            for n in sorted(tot, key=str.lower):
                ws.append([n, cents(sum(r[1] for r in data if r[3] == n and r[1] < 0)), cents(sum(r[1] for r in data if r[3] == n and r[1] > 0))])
                ws.cell(ws._current_row, 2).number_format = ws.cell(ws._current_row, 3).number_format = money_fmt

    wb = Workbook()
    ws = wb.active
    ws.title = "All Transactions"
    sheet(ws, rows)
    cats = sorted({r[4] for r in rows}, key=lambda c: (c == "Not Filed", c))
    for cat in cats:
        data = [r for r in rows if r[4] == cat]
        sh = wb.create_sheet(cat[:31])
        if cat == "Not Filed":
            sh.append(["Not in an account yet -- file them in AT-HUB (Accounting > To File) and make this file again."])
            sh["A1"].font = Font(bold=True, color="C00000")
            sh.append([])
        sheet(sh, data)
        totals(sh, data, loans=cat == "Loans")
    if aside:
        sh = wb.create_sheet("Not Included")
        sh.append(["Year-end AR / AP estimates -- set aside: not included in any total, in the Summary or in the split."])
        sh["A1"].font = Font(bold=True, color="C00000")
        sh.append([])
        sheet(sh, aside)
        sh.append([])
        sh.append(["Total (not included)", cents(sum(r[1] for r in aside))])
        sh.cell(sh._current_row, 1).font = bold
        sh.cell(sh._current_row, 2).number_format = money_fmt

    # Summary: like theirs, with live formulas
    pt = partners(db, start, end, fy_key)
    sm = wb.create_sheet("Summary")
    sm.append(["Category", "Transaction Count", "Total Amount"])
    for c in sm[1]:
        c.font = bold
    at = {}
    for cat in cats:
        data = [r for r in rows if r[4] == cat]
        sm.append([cat, len(data), cents(sum(r[1] for r in data))])
        sm.cell(sm._current_row, 3).number_format = money_fmt
        at[cat] = sm._current_row
    sm.append(["Account Balance", None, pt["bank_balance"]])
    sm.cell(sm._current_row, 3).number_format = money_fmt
    bal = sm._current_row
    sm.append([])

    def ref(cat):
        return f"C{at[cat]}" if cat in at else "0"
    first = sm._current_row + 1
    sm.append(["Money Out [Expenses + COGS]", f"={ref('Expense')}+{ref('COGS')}"])
    sm.append(["Money In [Payments]", f"={ref('Payments')}"])
    sm.append(["Debt / Loans Payable [borrowed - repaid, taken off]", f"=-{ref('Loans')}"])
    sm.append(["Account Balance", f"=C{bal}"])
    last = sm._current_row
    sm.append([])
    sm.append(["Profit", f"=SUM(B{first}:B{last})"])
    prow = sm._current_row
    sm.cell(prow, 1).font = sm.cell(prow, 2).font = bold
    for r in range(first, prow + 1):
        sm.cell(r, 2).number_format = money_fmt
    sm.append([])
    sm.append(["Partner", "Already Taken", "Share", "Adjust", "Remaining"])
    for c in sm[sm._current_row]:
        c.font = bold
    for x in pt["partners"]:
        n = sm._current_row + 1
        sm.append([x["name"], x["taken"], f"=B${prow}*{x['share_pct']}", x["adjust"], f"=C{n}+B{n}+D{n}"])
        for col in range(2, 6):
            sm.cell(n, col).number_format = money_fmt
    sm.append([])
    sm.append(["Notes:"])
    sm.cell(sm._current_row, 1).font = bold
    notes = [l for l in (pt["notes"] or "").splitlines() if l.strip()]
    if pt["set_aside"]["count"]:
        notes.append(f"Not included: {pt['set_aside']['count']} year-end AR / AP estimate line(s), net {pt['set_aside']['total']:,.2f} "
                     "-- set aside (they were estimates), see the Not Included tab.")
    if "Not Filed" in at:
        notes.append(f"{sum(1 for r in rows if r[4] == 'Not Filed')} line(s) not in an account yet ({pt['unfiled']:,.2f}) -- see the Not Filed tab.")
    notes.append(f"Period {start.isoformat()} to {end.isoformat()} ({fy_key}). Made by AT-HUB Accounting.")
    for l in notes:
        sm.append([l])
    sm.column_dimensions["A"].width = 48
    for c in "BCDE":
        sm.column_dimensions[c].width = 18

    # one extra the old script didn't have: the P&L by month
    p = pnl(db, start, end, "month")
    sh = wb.create_sheet("P&L By Month")
    sh.append(["Account"] + [colname for colname in p["columns"]] + ["Total"])
    for c in sh[1]:
        c.font = bold
    for sec in p["sections"]:
        if not sec["rows"]:
            continue
        name = CATEGORY.get(sec["kind"], sec["label"])
        sh.append([name])
        sh.cell(sh._current_row, 1).font = bold
        for r in sec["rows"]:
            sh.append([r["name"]] + r["cells"] + [r["total"]])
        sh.append([f"Total {name}"] + sec["totals"] + [sec["total"]])
        sh.cell(sh._current_row, 1).font = bold
        if sec["kind"] == "expense":
            sh.append(["Profit (Payments - COGS - Expense)"] + p["profit"] + [p["profit_total"]])
            sh.cell(sh._current_row, 1).font = bold
    sh.column_dimensions["A"].width = 34
    for row in sh.iter_rows(min_row=2, min_col=2):
        for c in row:
            c.number_format = money_fmt
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_quickbooks_csv(db: Session, source_id: int, start: date, end: date) -> str:
    """QuickBooks Online bank upload (3 columns: Date, Description, Amount) for one account -- the first step of a
    QuickBooks hand-over; the account each line is in goes along as a 4th column for whoever books it there."""
    accounts, _ = maps(db)
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["Date", "Description", "Amount", "Account"])
    for t in (db.query(AcctTxn).filter(AcctTxn.source_id == source_id, AcctTxn.date >= _dt(start), AcctTxn.date <= _dt(end),
                                       AcctTxn.expected.is_(False)).order_by(AcctTxn.date, AcctTxn.id).all()):
        a = accounts.get(t.account_id)
        w.writerow([t.date.strftime("%m/%d/%Y"), t.bank_description or t.description, f"{t.amount:.2f}", path(a, accounts) if a else ""])
    return buf.getvalue()


# ---------------------------------------------------------------- remembered filings (survive a fresh start / re-import)
def _mem_same(m: AcctMemory, d: date, amount: float, desc: str) -> bool:
    if abs(m.amount - amount) > 0.005:
        return False
    days = abs((m.date.date() - d).days)
    if days > 3:
        return False
    if days == 0 and starts_same(m.description or "", desc, 3):
        return True
    rr = _plain_refs(desc)
    return starts_same(m.description or "", desc) or bool(rr and rr & _plain_refs(m.description or ""))


def remember(db: Session, t: AcctTxn, kind: str, by: Optional[str]) -> AcctMemory:
    """Keep how a person filed this line (kind split: always, from set_account; filing: on a fresh start)."""
    desc = t.bank_description or t.description
    m = next((x for x in db.query(AcctMemory).filter(AcctMemory.amount == t.amount).all() if _mem_same(x, t.date.date(), t.amount, desc)), None)
    if not m:
        m = AcctMemory(date=t.date, amount=t.amount, description=desc, created_by=by)
        db.add(m)
    m.source_id, m.kind, m.payee, m.note = t.source_id, kind, t.payee, t.note
    if t.splits or kind == "split":
        m.kind = "split"
        m.account_id = None
        m.splits = json.dumps([{"account_id": x.account_id, "amount": x.amount, "note": x.note} for x in t.splits])
    else:
        m.account_id, m.splits = t.account_id, None
    return m


def forget_memory(db: Session, t: AcctTxn, kind: Optional[str] = None) -> int:
    n = 0
    for m in db.query(AcctMemory).filter(AcctMemory.amount == t.amount).all():
        if (kind is None or m.kind == kind) and _mem_same(m, t.date.date(), t.amount, t.bank_description or t.description):
            db.delete(m)
            n += 1
    return n


def memory_hits(db: Session, source_id: Optional[int], items) -> Dict[int, AcctMemory]:
    """index in items -> the remembered filing for it (items: AcctTxn or parsed rows)."""
    mems = db.query(AcctMemory).all()
    if not mems:
        return {}
    by_amt = defaultdict(list)
    for m in mems:
        by_amt[round(m.amount, 2)].append(m)
    out, used = {}, set()
    for i, x in enumerate(items):
        g = (lambda k: x.get(k)) if isinstance(x, dict) else (lambda k: getattr(x, k))
        d = g("date")
        d = d.date() if isinstance(d, datetime) else d
        desc = (g("bank_description") if not isinstance(x, dict) else None) or g("description")
        for m in by_amt.get(round(g("amount"), 2), []):
            if m.id in used or (m.source_id and source_id and m.source_id != source_id):
                continue
            if _mem_same(m, d, g("amount"), desc):
                out[i] = m
                used.add(m.id)
                break
    return out


def apply_memory(db: Session, source_id: int, txns: List[AcctTxn], by: str) -> int:
    """New statement lines a person had filed before (a fresh start, a re-import): filed the same way again."""
    accounts = {x[0] for x in db.query(AcctAccount.id).all()}
    n = 0
    for i, m in memory_hits(db, source_id, txns).items():
        t = txns[i]
        if t.account_id or t.splits:
            continue
        parts = json.loads(m.splits or "[]")
        if parts and all(p["account_id"] in accounts for p in parts):
            total = sum(p["amount"] for p in parts) or t.amount
            for p in parts:
                t.splits.append(AcctSplit(account_id=p["account_id"], amount=cents(p["amount"] * t.amount / total), note=p.get("note")))
            diff = cents(t.amount - sum(x.amount for x in t.splits))
            if abs(diff) > 0.004:
                t.splits[-1].amount = cents(t.splits[-1].amount + diff)
        elif m.account_id in accounts:
            t.account_id = m.account_id
        else:
            continue
        t.how, t.reviewed = "person", True
        t.note = t.note or m.note
        if t.flag == "qb_split" and t.splits:
            t.flag = t.flag_info = None
        m.used, m.last_used = (m.used or 0) + 1, datetime.utcnow()
        n += 1
    return n


def memory_list(db: Session) -> List[dict]:
    accounts, _ = maps(db)
    out = []
    for m in db.query(AcctMemory).order_by(AcctMemory.date.desc()).all():
        parts = json.loads(m.splits or "[]")
        out.append({"id": m.id, "date": m.date.date().isoformat(), "amount": m.amount, "description": m.description, "kind": m.kind,
                    "account": path(accounts[m.account_id], accounts) if m.account_id in accounts else None,
                    "splits": [{**p, "account": path(accounts[p["account_id"]], accounts) if p["account_id"] in accounts else "?"} for p in parts],
                    "note": m.note, "used": m.used, "last_used": m.last_used, "created_by": m.created_by, "created_at": m.created_at})
    return out


SEEDED_KEY = "acct_memory_seeded"


def seed_memory_once(db: Session) -> int:
    """Once: every line a person has split so far is remembered (later splits are remembered as they're made)."""
    if db.get(AppSetting, SEEDED_KEY):
        return 0
    n = 0
    for t in db.query(AcctTxn).options(selectinload(AcctTxn.splits)).filter(AcctTxn.splits.any()).all():
        remember(db, t, "split", t.updated_by or t.created_by)
        n += 1
    db.add(AppSetting(key=SEEDED_KEY, value=json.dumps({"lines": n, "at": datetime.utcnow().isoformat()})))
    db.commit()
    return n


def start_fresh(db: Session, keep_filings: bool, by: str) -> dict:
    """Clear the books' lines (every statement line, typed-in line, import, ignored line) to bring the statements in again --
    accounts, sub-accounts, rules, bank accounts, comments, projections, AP / AR notes and settings stay. Every split is
    remembered first (and, keep_filings, how a person filed each line), so the same lines come back filed the same way."""
    splits = filings = 0
    for t in db.query(AcctTxn).options(selectinload(AcctTxn.splits)).all():
        if t.expected or t.origin == "manual":
            continue
        if t.splits and t.how in ("person", "workbook"):
            remember(db, t, "split", by)
            splits += 1
        elif keep_filings and t.account_id and t.how in ("person", "workbook") and t.reviewed:
            remember(db, t, "filing", by)
            filings += 1
    db.flush()
    lines = db.query(AcctTxn).count()
    db.query(AcctSplit).delete(synchronize_session=False)
    db.query(AcctTxn).delete(synchronize_session=False)
    db.query(AcctIgnored).delete(synchronize_session=False)
    db.query(AcctCheckOk).delete(synchronize_session=False)
    imports = db.query(AcctImport).count()
    db.query(AcctImport).delete(synchronize_session=False)
    db.query(AcctChange).delete(synchronize_session=False)
    db.query(AcctRule).update({AcctRule.hits: 0}, synchronize_session=False)
    for k in ("acct_rec_decisions",):
        row = db.get(AppSetting, k)
        if row:
            db.delete(row)
    db.commit()
    return {"lines": lines, "imports": imports, "remembered_splits": splits, "remembered_filings": filings,
            "memory": db.query(AcctMemory).count()}


# ---------------------------------------------------------------- QuickBooks' pasted memo on lines already in
def tails_existing(db: Session) -> List[dict]:
    """Statement lines already in the books that carry QuickBooks' pasted text: what comes off, and what that does to a
    line a rule / AT-HUB filed (the rule no longer matches -> back to Review)."""
    accounts, _ = maps(db)
    txns = (db.query(AcctTxn).options(selectinload(AcctTxn.splits)).filter(AcctTxn.origin != "manual", AcctTxn.qb_memo.is_(None))
            .order_by(AcctTxn.date).all())
    rows = [{"date": t.date.date(), "description": t.description} for t in txns]
    specs = rule_specs(db)
    out = []
    for t, (clean, memo, why) in zip(txns, acct_tails.detect(rows)):
        if not memo:
            continue
        after = {"description": clean, "bank_description": None, "payee": t.payee, "note": t.note, "amount": t.amount,
                 "source_id": t.source_id, "date": t.date}
        effect = None
        if t.how in ("rule", "auto") and (t.account_id or t.splits):
            r, clash = acct_rules.decide_rules(specs, acct_rules.line_of(after))
            if clash:
                effect = "rules disagree now -- back to Review"
            elif t.how == "rule" and (not r or r["id"] != t.rule_id):
                effect = "its rule no longer matches -- back to Review"
        a = accounts.get(t.account_id)
        out.append({"id": t.id, "date": t.date.date().isoformat(), "amount": t.amount, "before": t.description, "after": clean, "removed": memo,
                    "why": why, "account": path(a, accounts) if a else ("split" if t.splits else None), "how": t.how, "effect": effect})
    return out


def clean_tails(db: Session, ids: List[int], by: str) -> dict:
    found = {x["id"]: x for x in tails_existing(db)}
    n = back = 0
    for t in db.query(AcctTxn).options(selectinload(AcctTxn.splits)).filter(AcctTxn.id.in_([i for i in ids if i in found])).all():
        x = found[t.id]
        t.description, t.qb_memo, t.updated_by = x["after"], x["removed"], by
        t.payee = acct_mapper.payee_label(x["after"]) if t.origin != "manual" else t.payee
        n += 1
        if x["effect"]:
            was = {"account_id": t.account_id} if t.account_id else {"splits": [{"account_id": s_.account_id, "amount": s_.amount, "note": s_.note} for s_ in t.splits]}
            t.sent_back_from = json.dumps(was)
            t.splits.clear()
            t.account_id, t.how, t.rule_id, t.reviewed = None, None, None, True
            t.sent_back, t.sent_back_reason, t.sent_back_by, t.sent_back_at = True, "QuickBooks' pasted text removed -- " + x["effect"].split(" -- ")[0], by, datetime.utcnow()
            back += 1
    db.flush()
    return {"cleaned": n, "back_to_review": back}
