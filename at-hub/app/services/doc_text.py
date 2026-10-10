"""Text out of any document we accept -- so the AI readers work on Excel / CSV / Word / text / email files too,
not only PDFs and photos -- and a quick preview (a table for spreadsheets, text for the rest) for the file viewer.

    family(name)        -> "pdf" | "image" | "sheet" | "csv" | "word" | "text" | "email" | "other"
    text_of(data, name) -> the document's words ("" when there are none to read: a scan, a photo, .xls, .msg)
    preview(data, name) -> {"type": "table", "sheets": [{"name", "rows"}]} | {"type": "text", "text"} | None
"""
import csv
import email
import io
import re
import zipfile
from email import policy

IMAGES = (".png", ".jpg", ".jpeg", ".gif", ".webp", ".heic", ".heif")
MAX_ROWS, MAX_COLS = 300, 40


def family(name: str) -> str:
    n = (name or "").lower()
    if n.endswith(".pdf"):
        return "pdf"
    if n.endswith(IMAGES):
        return "image"
    if n.endswith((".xlsx", ".xlsm")):
        return "sheet"
    if n.endswith(".csv"):
        return "csv"
    if n.endswith(".docx"):
        return "word"
    if n.endswith(".txt"):
        return "text"
    if n.endswith(".eml"):
        return "email"
    return "other"  # .xls, .doc, .msg: kept and attached, not read


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", "replace")


def _cell(v) -> str:
    if v is None:
        return ""
    if hasattr(v, "strftime"):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _sheets(data: bytes):
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        rows = []
        for r in ws.iter_rows(values_only=True):
            cells = [_cell(v) for v in r[:MAX_COLS]]
            if any(cells):
                rows.append(cells)
            if len(rows) >= MAX_ROWS:
                break
        if rows:
            width = max(len(r) for r in rows)
            while width and not any(len(r) >= width and r[width - 1] for r in rows):
                width -= 1  # drop empty columns on the right
            out.append({"name": ws.title, "rows": [r[:width] + [""] * (width - len(r[:width])) for r in rows]})
    wb.close()
    return out


def _csv_rows(data: bytes):
    text = _decode(data)
    try:
        dialect = csv.Sniffer().sniff(text[:4000], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = []
    for r in csv.reader(io.StringIO(text), dialect):
        if any(c.strip() for c in r):
            rows.append([c.strip() for c in r[:MAX_COLS]])
        if len(rows) >= MAX_ROWS:
            break
    return rows


def _docx_text(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        xml = z.read("word/document.xml").decode("utf-8", "replace")
    xml = re.sub(r"</w:p>", "\n", xml)
    xml = re.sub(r"<w:tab/>|</w:tc>", "\t", xml)
    return re.sub(r"<[^>]+>", "", xml).replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">").strip()


def _email(data: bytes):
    """(header + body text, [(attachment name, bytes)])"""
    msg = email.message_from_bytes(data, policy=policy.default)
    head = "\n".join(f"{k}: {msg[k]}" for k in ("From", "To", "Subject", "Date") if msg[k])
    body, files = "", []
    for part in msg.walk():
        if part.is_multipart():
            continue
        name = part.get_filename()
        if name:
            files.append((name, part.get_payload(decode=True) or b""))
        elif part.get_content_type() == "text/plain" and not body:
            body = part.get_content()
        elif part.get_content_type() == "text/html" and not body:
            body = re.sub(r"<[^>]+>", " ", part.get_content())
    return f"{head}\n\n{body}".strip(), files


def text_of(data: bytes, name: str) -> str:
    f = family(name)
    try:
        if f == "pdf":
            from app.services.ai_orders import pdf_text
            return pdf_text(data)
        if f == "sheet":
            return "\n\n".join(f"Sheet {s['name']}:\n" + "\n".join(" | ".join(c for c in r) for r in s["rows"]) for s in _sheets(data))
        if f == "csv":
            return "\n".join(" | ".join(r) for r in _csv_rows(data))
        if f == "word":
            return _docx_text(data)
        if f == "text":
            return _decode(data)
        if f == "email":
            text, files = _email(data)
            # an emailed PO is usually the attachment: read the attached documents too
            for n, b in files:
                inner = text_of(b, n) if family(n) != "email" else ""
                if inner:
                    text += f"\n\n--- attached {n} ---\n{inner}"
            return text
    except Exception:
        return ""
    return ""


def preview(data: bytes, name: str):
    f = family(name)
    try:
        if f == "sheet":
            return {"type": "table", "sheets": _sheets(data)}
        if f == "csv":
            return {"type": "table", "sheets": [{"name": name, "rows": _csv_rows(data)}]}
        if f in ("word", "text", "email"):
            return {"type": "text", "text": text_of(data, name)[:60000]}
    except Exception:
        return None
    return None
