"""Customer open-lines check: a customer's own report of the POs they still have open with us (Hudson's "Open Purchase
Orders by Request Date") read line by line against our orders.

    read_file(data, name)                 -> (headers, rows)  -- Excel or CSV; the header row is found, not assumed
    guess_mapping(headers)                -> {our field: their column}
    check(db, customer_id, rows, mapping, rules) -> {"summary", "rows", "ours_only", "missing_orders"}
    create_orders(db, report, pos, by)    -> the orders missing from AT-HUB, made from the report (in Validation)
    apply_theirs(db, report, idx, field, by)  -> take their value for one difference (qty / price / date / job / add line)

Matching: their PO # -> our order for that customer (else any customer's, said so), their item # -> our line (the same
item twice on an order: same ordered qty first, then the same line #). Every matched line is compared: ordered qty,
open qty (ours = ordered - shipped, with the shipments that went), price, wanted date, job #. Rules leave out what
isn't a real line on one side (their kit header lines, our $0 companion nut lines)."""
import io
import json
import re
from collections import defaultdict
from datetime import datetime, date

from fastapi import HTTPException
from sqlalchemy.orm import Session

from app.models import Customer, CustomerOrder, CustomerOrderLine, CustomerReport, CustomerReportProfile, Shipment, ShipmentLine, StockItem

FIELDS = {  # our field -> words that name it in their header (first match wins)
    "po": ["order number", "po number", "po #", "po no", "purchase order", "order #", "order no", "po"],
    "line": ["line number", "line #", "line no", "line", "ln"],
    "line_type": ["ln ty", "line type"],
    "item": ["item number", "item #", "item no", "part number", "part #", "our part", "item", "part"],
    "desc": ["description"],
    "desc2": ["description line 2", "description 2"],
    "qty_ordered": ["order quantity", "qty ordered", "ordered qty", "quantity ordered", "order qty"],
    "qty_open": ["quantity open", "open qty", "qty open", "open quantity", "balance", "remaining", "qty due", "backorder"],
    "price": ["unit cost", "unit price", "price"],
    "request_date": ["request date", "required date", "need date", "due date"],
    "promised_date": ["promised delivery date", "promised date", "promise date"],
    "job": ["reference 2", "job", "job number", "project"],
    "ship_to": ["ship to name", "ship to"],
}
LABELS = {"po": "Their PO #", "line": "Line #", "line_type": "Line Type", "item": "Item #", "desc": "Description", "desc2": "Description 2",
          "qty_ordered": "Ordered Qty", "qty_open": "Open Qty", "price": "Unit Price", "request_date": "Request Date",
          "promised_date": "Promised Date", "job": "Job #", "ship_to": "Ship To"}
DEFAULT_RULES = {"skip_their_line_types": ["S"], "skip_their_max_price": 0.01,  # kit header lines: type S at a penny or less
                 "skip_our_zero_price": True}  # our $0 lines (companion nuts) aren't on their PO
EPS = 1e-6


def _norm(v) -> str:
    return re.sub(r"[^A-Z0-9]", "", str(v if v is not None else "").upper())


def _num(v):
    if v in (None, ""):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", "").replace("$", "").strip())
    except ValueError:
        return None


def _day(v):
    if v in (None, ""):
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(s[:19], fmt).date().isoformat()
        except ValueError:
            continue
    return None


def _code(v) -> str:
    """Item / PO numbers as text: 4156932.0 -> "4156932", "  7621" -> "7621"."""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v if v is not None else "").strip()


# ---------------------------------------------------------------------------------------------------- reading
def read_file(data: bytes, name: str):
    n = (name or "").lower()
    if n.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        raw = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
        wb.close()
    elif n.endswith((".csv", ".txt")):
        import csv
        from app.services.doc_text import _decode
        raw = [r for r in csv.reader(io.StringIO(_decode(data)))]
    else:
        raise HTTPException(status_code=400, detail="Upload the report as Excel (.xlsx) or CSV")
    raw = [r for r in raw if any(c not in (None, "") for c in r)]
    # the header row: the first with several words we know ("order number", "item", "quantity"...)
    known = [w for words in FIELDS.values() for w in words]
    hi = next((i for i, r in enumerate(raw[:15]) if sum(1 for c in r if isinstance(c, str) and any(k in c.lower() for k in known)) >= 3), None)
    if hi is None:
        raise HTTPException(status_code=400, detail="Couldn't find the header row (PO #, item #, quantities) in this file")
    headers = [str(c).strip() if c is not None else "" for c in raw[hi]]
    rows = []
    for r in raw[hi + 1:]:
        row = {headers[i]: r[i] for i in range(min(len(headers), len(r))) if headers[i]}
        if any(v not in (None, "") for v in row.values()):
            rows.append(row)
    return headers, rows


def guess_mapping(headers) -> dict:
    out, used = {}, set()
    low = {h: h.lower().strip() for h in headers if h}
    # longer, more specific names first ("description line 2" before "description")
    for field in sorted(FIELDS, key=lambda f: -max(len(w) for w in FIELDS[f])):
        for word in FIELDS[field]:
            hit = next((h for h, l in low.items() if h not in used and (l == word or (len(word) > 3 and word in l))), None)
            if hit:
                out[field], _ = hit, used.add(hit)
                break
    return out


# ---------------------------------------------------------------------------------------------------- checking
def _their(row: dict, mapping: dict) -> dict:
    g = lambda f: row.get(mapping.get(f)) if mapping.get(f) else None
    desc = " ".join(str(x).strip() for x in (g("desc"), g("desc2")) if x not in (None, ""))
    return {"po": _code(g("po")), "line": _code(g("line")), "line_type": _code(g("line_type")).upper(), "item": _code(g("item")), "desc": desc,
            "qty_ordered": _num(g("qty_ordered")), "qty_open": _num(g("qty_open")), "price": _num(g("price")),
            "request_date": _day(g("request_date")), "promised_date": _day(g("promised_date")), "job": _code(g("job")),
            "ship_to": _code(g("ship_to"))}


def _shipments_by_line(db: Session, line_ids):
    out = defaultdict(list)
    if not line_ids:
        return out
    from app.models import Attachment
    rows = (db.query(ShipmentLine.order_line_id, ShipmentLine.quantity, Shipment)
            .join(Shipment, Shipment.id == ShipmentLine.shipment_id)
            .filter(ShipmentLine.order_line_id.in_(list(line_ids)), Shipment.status != "cancelled").all())
    pods = {sid for (sid,) in db.query(Attachment.entity_id).filter(Attachment.entity_type == "shipment", Attachment.category == "pod").all()}
    for lid, qty, s in rows:
        out[lid].append({"id": s.id, "code": s.code, "status": s.status, "qty": qty,
                         "ship_date": s.ship_date.isoformat() if s.ship_date else None,
                         "delivered_at": s.delivered_at.isoformat() if s.delivered_at else None, "pod": s.id in pods})
    return out


def check(db: Session, customer_id: int, rows: list, mapping: dict, rules: dict = None) -> dict:
    rules = {**DEFAULT_RULES, **(rules or {})}
    cust = db.get(Customer, customer_id)
    if not cust:
        raise HTTPException(status_code=400, detail="Pick the customer")
    for f in ("po", "item"):
        if not mapping.get(f):
            raise HTTPException(status_code=400, detail=f"Say which column is {LABELS[f]}")
    theirs = [_their(r, mapping) for r in rows]
    theirs = [t for t in theirs if t["po"] or t["item"]]
    # our orders by PO #: this customer first, else anyone's (a second customer record for the same company)
    mine, anyone = defaultdict(list), defaultdict(list)
    for o in db.query(CustomerOrder).filter(CustomerOrder.status != "cancelled").all():
        k = _norm(o.po_number)
        if k:
            (mine if o.customer_id == customer_id else anyone)[k].append(o)
    pos = {_norm(t["po"]) for t in theirs}
    orders = {}
    for k in pos:
        os_ = mine.get(k) or anyone.get(k) or []
        if os_:
            orders[k] = sorted(os_, key=lambda o: o.id)
    order_ids = [o.id for os_ in orders.values() for o in os_]
    lines = defaultdict(list)
    for l in db.query(CustomerOrderLine).filter(CustomerOrderLine.order_id.in_(order_ids or [0])).all():
        lines[l.order_id].append(l)
    codes = {i.id: i for i in db.query(StockItem).filter(StockItem.id.in_({l.item_id for ls in lines.values() for l in ls} or {0})).all()}
    ships = _shipments_by_line(db, [l.id for ls in lines.values() for l in ls])
    used = set()
    out_rows, summary = [], defaultdict(int)

    def our_line_out(l, o):
        it = codes.get(l.item_id)
        return {"order_id": o.id, "order_code": o.code, "order_status": o.status, "line_id": l.id, "line_no": l.line_no,
                "item": it.code if it else "", "title": it.title if it else "", "qty": l.quantity, "shipped": l.shipped_quantity or 0,
                "open": max(0.0, l.quantity - (l.shipped_quantity or 0)), "price": l.unit_price,
                "date": (l.delivery_date or o.delivery_date).date().isoformat() if (l.delivery_date or o.delivery_date) else None,
                "job": o.job_number or "", "shipments": ships.get(l.id, [])}

    for idx, t in enumerate(theirs):
        row = {"idx": idx, "theirs": t, "ours": None, "status": None, "flags": [], "notes": []}
        if t["line_type"] in [x.upper() for x in rules["skip_their_line_types"]] and (t["price"] or 0) <= rules["skip_their_max_price"] + EPS:
            row["status"] = "ignored"
            row["notes"].append(f"line type {t['line_type']} at {t['price'] or 0:g} -- their kit / header line")
        else:
            os_ = orders.get(_norm(t["po"]))
            if not os_:
                row["status"] = "missing_order"
            else:
                o = os_[0]
                if o.customer_id != customer_id:
                    row["notes"].append(f"PO is on {o.code} for another customer ({db.get(Customer, o.customer_id).name})")
                cand = [l for x in os_ for l in lines[x.id] if l.id not in used
                        and _norm(codes[l.item_id].code if l.item_id in codes else "") == _norm(t["item"])]
                if not cand:
                    row["status"] = "missing_line"
                    row["ours"] = {"order_id": o.id, "order_code": o.code, "order_status": o.status}
                else:
                    l = next((c for c in cand if t["qty_ordered"] is not None and abs(c.quantity - t["qty_ordered"]) < EPS), None) \
                        or next((c for c in cand if t["line"] and str(c.line_no) == t["line"]), None) or cand[0]
                    used.add(l.id)
                    ox = next(x for x in os_ if x.id == l.order_id)
                    ours = our_line_out(l, ox)
                    row["ours"] = ours
                    if t["qty_ordered"] is not None and abs(ours["qty"] - t["qty_ordered"]) > EPS:
                        row["flags"].append("qty")
                    if t["qty_open"] is not None and abs(ours["open"] - t["qty_open"]) > EPS:
                        row["flags"].append("shipped_open" if ours["open"] < t["qty_open"] else "open_more")
                    if t["price"] is not None and ours["price"] is not None and abs(ours["price"] - t["price"]) > 0.00005:
                        row["flags"].append("price")
                    wanted = {d for d in (t["request_date"], t["promised_date"]) if d}  # ours may follow either of their dates
                    if wanted and ours["date"] and ours["date"] not in wanted:
                        row["flags"].append("date")
                    if t["job"] and _norm(t["job"]) != _norm(ours["job"]):
                        row["flags"].append("job")
                    row["status"] = "match" if not row["flags"] else "diff"
        summary[row["status"]] += 1
        for f in row["flags"]:
            summary["flag_" + f] += 1
        out_rows.append(row)
    # ours that their report doesn't have (on the POs they listed)
    ours_only = []
    for os_ in orders.values():
        for o in os_:
            for l in lines[o.id]:
                if l.id in used:
                    continue
                if rules.get("skip_our_zero_price") and not (l.unit_price or 0):
                    continue
                if l.quantity - (l.shipped_quantity or 0) > EPS:
                    ours_only.append(our_line_out(l, o))
    summary["ours_only"] = len(ours_only)
    missing = defaultdict(list)
    for r in out_rows:
        if r["status"] == "missing_order":
            missing[r["theirs"]["po"]].append(r["idx"])
    item_codes = {_norm(i.code): i for i in db.query(StockItem).all()}
    missing_orders = [{"po": po, "rows": idxs, "lines": len(idxs), "job": out_rows[idxs[0]]["theirs"]["job"],
                       "known_items": sum(1 for i in idxs if _norm(out_rows[i]["theirs"]["item"]) in item_codes)} for po, idxs in missing.items()]
    summary["missing_orders"] = len(missing_orders)
    summary["lines"] = len(out_rows)
    summary["pos"] = len({_norm(t["po"]) for t in theirs if t["po"]})
    return {"customer_id": customer_id, "customer": cust.name, "summary": dict(summary), "rows": out_rows, "ours_only": ours_only,
            "missing_orders": missing_orders, "rules": rules, "checked_at": datetime.utcnow().isoformat() + "Z"}


def changes_since(prev: dict, cur: dict) -> dict:
    """What moved since the last upload for this customer: by (PO, line)."""
    if not prev:
        return {}
    key = lambda r: (r["theirs"]["po"], r["theirs"]["line"], _norm(r["theirs"]["item"]))
    before = {key(r): r["status"] for r in prev.get("rows", [])}
    now = {key(r): r["status"] for r in cur.get("rows", [])}
    return {"gone": len([k for k in before if k not in now]),  # off their report: received / closed by them
            "new": len([k for k in now if k not in before]),
            "fixed": len([k for k in now if k in before and before[k] != "match" and now[k] == "match"]),
            "worse": len([k for k in now if k in before and before[k] == "match" and now[k] != "match"])}


def mark_orders(db: Session, result: dict, report_id: int) -> None:
    """Each of our orders on the report remembers the check: all lines matched, or how many differences."""
    per = defaultdict(lambda: {"matched": 0, "diffs": 0, "missing": 0, "shipped_open": 0})
    for r in result["rows"]:
        o = (r.get("ours") or {}).get("order_id")
        if not o:
            continue
        real = [f for f in r.get("flags") or [] if f != "shipped_open"]  # shipped but open on theirs = their receiving, not our error
        if "shipped_open" in (r.get("flags") or []):
            per[o]["shipped_open"] += 1
        if r["status"] == "missing_line":
            per[o]["missing"] += 1
        elif real:
            per[o]["diffs"] += 1
        else:
            per[o]["matched"] += 1
    for oid, c in per.items():
        o = db.get(CustomerOrder, oid)
        if o:
            o.report_check = json.dumps({"report_id": report_id, "at": result["checked_at"], "customer": result["customer"], **c})


# ---------------------------------------------------------------------------------------------------- acting
def create_orders(db: Session, report: CustomerReport, pos, by: str) -> list:
    """Orders their report has and we don't: made from the report, in Validation (lines whose item # we don't have
    wait as "pick the item"), the report file attached."""
    from app.services import ai_pending
    from app.routes.attachments import store_file, upload_root
    res = json.loads(report.result or "{}")
    rows = res.get("rows") or []
    wanted = {_norm(p) for p in pos}
    groups = defaultdict(list)
    for r in rows:
        if r["status"] == "missing_order" and _norm(r["theirs"]["po"]) in wanted:
            groups[r["theirs"]["po"]].append(r)
    if not groups:
        raise HTTPException(status_code=400, detail="Those POs aren't missing any more -- check the report again")
    items = {_norm(i.code): i for i in db.query(StockItem).filter(StockItem.is_active == True).all()}  # noqa: E712
    data = (upload_root() / report.stored_name).read_bytes()
    made = []
    have = {_norm(o.po_number) for o in db.query(CustomerOrder).filter(CustomerOrder.customer_id == report.customer_id,
                                                                       CustomerOrder.status != "cancelled").all() if o.po_number}
    for po, rs in groups.items():
        if _norm(po) in have:  # made since the check: not twice
            continue
        t0 = rs[0]["theirs"]
        lines = []
        for r in sorted(rs, key=lambda r: (_num(r["theirs"]["line"]) or 0)):
            t = r["theirs"]
            it = items.get(_norm(t["item"]))
            lines.append({"item_id": it.id if it else None, "item_code": t["item"], "description": t["desc"],
                          "quantity": t["qty_open"] if t["qty_open"] is not None else (t["qty_ordered"] or 1),
                          "unit_price": t["price"] or 0, "delivery_date": t["promised_date"] or t["request_date"],
                          "line_note": None, "candidates": []})
        draft = {"customer": {"customer_id": report.customer_id}, "po_number": po, "job_number": t0["job"] or None,
                 "delivery_date": t0["promised_date"] or t0["request_date"], "lines": lines,
                 "notes": f"Made from {report.filename} (customer open-lines report) -- check it against their PO"}
        rec = ai_pending.create_for_validation(db, "customer", draft, report.filename, by)
        rec.status = "validation"
        store_file(db, "customer_order", rec.id, "other", report.filename, None, data, "Customer open-lines report", by)
        made.append({"po": po, "order_id": rec.id, "code": rec.code, "lines": len(rec.lines), "waiting": len(rec.ai_pending_lines)})
    db.commit()
    return made


def apply_theirs(db: Session, report: CustomerReport, idx: int, field: str, by: str) -> dict:
    """Take their value for one flagged line: qty / price / date / job, or add their line to our order."""
    from app import schemas
    from app.services.crud import CustomerOrderService
    res = json.loads(report.result or "{}")
    rows = res.get("rows") or []
    if not 0 <= idx < len(rows):
        raise HTTPException(status_code=400, detail="That line isn't on the report -- check it again")
    r = rows[idx]
    t, ours = r["theirs"], r.get("ours") or {}
    if field == "add_line":
        if r["status"] != "missing_line" or not ours.get("order_id"):
            raise HTTPException(status_code=400, detail="That line isn't missing from our order")
        it = db.query(StockItem).filter(StockItem.code == t["item"]).first()
        if not it:
            raise HTTPException(status_code=400, detail=f"We have no item {t['item']} -- add the item first")
        CustomerOrderService.add_line(db, ours["order_id"], schemas.CustomerOrderLineAdd(
            item_id=it.id, quantity=t["qty_ordered"] or t["qty_open"] or 1, unit_price=t["price"] or 0,
            delivery_date=t["promised_date"] or t["request_date"]))
        return {"done": f"added {t['item']} to {ours['order_code']}"}
    if not ours.get("line_id"):
        raise HTTPException(status_code=400, detail="That line isn't matched to one of ours")
    if field == "job":
        CustomerOrderService.update(db, ours["order_id"], schemas.CustomerOrderUpdate(job_number=t["job"]))
        return {"done": f"job # on {ours['order_code']} is now {t['job']}"}
    changes = {}
    if field == "qty":
        changes["quantity"] = t["qty_ordered"]
    elif field == "price":
        changes["unit_price"] = t["price"]
    elif field == "date":
        changes["delivery_date"] = t["promised_date"] or t["request_date"]
    else:
        raise HTTPException(status_code=400, detail="Unknown change")
    CustomerOrderService.update_line(db, ours["order_id"], ours["line_id"], schemas.CustomerOrderLineUpdate(**changes))
    return {"done": f"{ours['order_code']} line #{ours['line_no']}: {field} set to theirs"}
