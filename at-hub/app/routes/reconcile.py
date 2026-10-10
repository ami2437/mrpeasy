"""Customer open-lines check (services/open_lines.py): upload a customer's report of their open POs with us, see what
matches our orders and what doesn't, create the missing orders, take their value where ours is wrong. Money permission
"reconcile" (prices are compared) -- Admin and Super admin until a role is given it."""
import io
import json
import uuid
from pathlib import Path
from typing import List, Optional

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import Response
from pydantic import BaseModel
from sqlalchemy.orm import Session

from app.config.database import get_db
from app.dependencies import get_current_active_user, require_perm
from app.models import Customer, CustomerReport, CustomerReportProfile, User
from app.routes.attachments import MAX_BYTES, upload_root
from app.services import open_lines

router = APIRouter(prefix="/api/reconcile", tags=["reconcile"], dependencies=[Depends(require_perm("reconcile"))])


def _profile(db: Session, customer_id: int):
    return db.query(CustomerReportProfile).filter(CustomerReportProfile.customer_id == customer_id).first()


def _rows(report: CustomerReport):
    return open_lines.read_file((upload_root() / report.stored_name).read_bytes(), report.filename)


def _brief(r: CustomerReport, db: Session) -> dict:
    res = json.loads(r.result or "{}")
    c = db.get(Customer, r.customer_id)
    return {"id": r.id, "customer_id": r.customer_id, "customer": c.name if c else "", "filename": r.filename, "uploaded_by": r.uploaded_by,
            "created_at": r.created_at, "summary": res.get("summary") or {}, "checked": bool(res.get("rows") is not None)}


def _full(r: CustomerReport, db: Session) -> dict:
    headers, _rows_ = _rows(r)
    res = json.loads(r.result or "{}")
    prev = (db.query(CustomerReport).filter(CustomerReport.customer_id == r.customer_id, CustomerReport.id < r.id)
            .order_by(CustomerReport.id.desc()).first())
    return {**_brief(r, db), "headers": headers, "mapping": json.loads(r.mapping or "{}"), "labels": open_lines.LABELS,
            "result": res, "since_last": open_lines.changes_since(json.loads(prev.result or "{}") if prev else None, res),
            "previous": {"id": prev.id, "created_at": prev.created_at, "filename": prev.filename} if prev else None}


def _run(db: Session, r: CustomerReport, mapping: dict, rules: Optional[dict], by: str, save_profile: bool = True) -> None:
    _h, rows = _rows(r)
    prof = _profile(db, r.customer_id)
    rules = rules if rules is not None else (json.loads(prof.rules) if prof and prof.rules else None)
    res = open_lines.check(db, r.customer_id, rows, mapping, rules)
    r.mapping, r.result = json.dumps(mapping), json.dumps(res, default=str)
    if save_profile:
        if not prof:
            prof = CustomerReportProfile(customer_id=r.customer_id)
            db.add(prof)
        prof.mapping, prof.rules, prof.updated_by = json.dumps(mapping), json.dumps(res["rules"]), by
    db.flush()
    open_lines.mark_orders(db, res, r.id)
    db.commit()


@router.get("/reports")
def reports(customer_id: Optional[int] = None, db: Session = Depends(get_db)):
    q = db.query(CustomerReport)
    if customer_id:
        q = q.filter(CustomerReport.customer_id == customer_id)
    return [_brief(r, db) for r in q.order_by(CustomerReport.id.desc()).limit(200).all()]


@router.post("/reports")
def upload(customer_id: int = Form(...), file: UploadFile = File(...), db: Session = Depends(get_db),
           user: User = Depends(get_current_active_user)):
    """Keep the file, read it with the customer's saved column layout (or a guess), and check it straight away when
    the PO # and item # columns are known."""
    if not db.get(Customer, customer_id):
        raise HTTPException(status_code=400, detail="Pick the customer")
    name = Path(file.filename or "report.xlsx").name
    data = file.file.read()
    if not data or len(data) > MAX_BYTES:
        raise HTTPException(status_code=400, detail="Empty, or larger than 25 MB")
    headers, rows = open_lines.read_file(data, name)
    folder = upload_root() / "reports"
    folder.mkdir(parents=True, exist_ok=True)
    stored = f"reports/{uuid.uuid4().hex[:12]}{Path(name).suffix.lower()}"
    (upload_root() / stored).write_bytes(data)
    prof = _profile(db, customer_id)
    saved = json.loads(prof.mapping) if prof and prof.mapping else {}
    mapping = {k: v for k, v in saved.items() if v in headers} if saved else {}
    if not mapping.get("po") or not mapping.get("item"):
        mapping = {**open_lines.guess_mapping(headers), **mapping}
    r = CustomerReport(customer_id=customer_id, filename=name, stored_name=stored, mapping=json.dumps(mapping), uploaded_by=user.username)
    db.add(r)
    db.flush()
    if mapping.get("po") and mapping.get("item"):
        _run(db, r, mapping, None, user.username)
    else:
        db.commit()
    return _full(r, db)


@router.get("/reports/{rid}")
def get_report(rid: int, db: Session = Depends(get_db)):
    r = db.get(CustomerReport, rid)
    if not r:
        raise HTTPException(status_code=404, detail="Report not found")
    return _full(r, db)


class RecheckIn(BaseModel):
    mapping: Optional[dict] = None
    rules: Optional[dict] = None


@router.post("/reports/{rid}/check")
def recheck(rid: int, data: RecheckIn, db: Session = Depends(get_db), user: User = Depends(get_current_active_user)):
    """Check again (after fixing orders, or with a different column layout / rules -- saved for this customer)."""
    r = db.get(CustomerReport, rid)
    if not r:
        raise HTTPException(status_code=404, detail="Report not found")
    _run(db, r, data.mapping or json.loads(r.mapping or "{}"), data.rules, user.username)
    return _full(r, db)


class CreateIn(BaseModel):
    pos: List[str]


@router.post("/reports/{rid}/create-orders")
def create_orders(rid: int, data: CreateIn, db: Session = Depends(get_db), user: User = Depends(get_current_active_user)):
    from app.services.permissions import has
    if not has(user, "orders.edit"):
        raise HTTPException(status_code=403, detail="Creating orders needs the orders permission")
    r = db.get(CustomerReport, rid)
    if not r:
        raise HTTPException(status_code=404, detail="Report not found")
    made = open_lines.create_orders(db, r, data.pos, user.username)
    _run(db, r, json.loads(r.mapping or "{}"), None, user.username, save_profile=False)
    return {"made": made, **_full(r, db)}


class ApplyIn(BaseModel):
    idx: int
    field: str  # qty | price | date | job | add_line


@router.post("/reports/{rid}/apply")
def apply(rid: int, data: ApplyIn, db: Session = Depends(get_db), user: User = Depends(get_current_active_user)):
    from app.services.permissions import has
    if not has(user, "orders.edit"):
        raise HTTPException(status_code=403, detail="Changing orders needs the orders permission")
    r = db.get(CustomerReport, rid)
    if not r:
        raise HTTPException(status_code=404, detail="Report not found")
    done = open_lines.apply_theirs(db, r, data.idx, data.field, user.username)
    _run(db, r, json.loads(r.mapping or "{}"), None, user.username, save_profile=False)
    return {**done, **_full(r, db)}


@router.delete("/reports/{rid}", status_code=204)
def delete(rid: int, db: Session = Depends(get_db)):
    r = db.get(CustomerReport, rid)
    if r:
        (upload_root() / r.stored_name).unlink(missing_ok=True)
        db.delete(r)
        db.commit()
    return Response(status_code=204)


@router.get("/reports/{rid}/export")
def export(rid: int, db: Session = Depends(get_db)):
    """The check as Excel: every line with its verdict (e.g. to send the "shipped, still open on your side" list)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    r = db.get(CustomerReport, rid)
    if not r:
        raise HTTPException(status_code=404, detail="Report not found")
    res = json.loads(r.result or "{}")
    words = {"match": "Matches", "diff": "Differs", "missing_order": "PO not in our system", "missing_line": "Line not on our order",
             "ignored": "Left out (rule)"}
    flag_words = {"qty": "ordered qty", "price": "price", "date": "date", "job": "job #", "shipped_open": "shipped, still open on yours",
                  "open_more": "open qty differs"}
    wb = Workbook()
    ws = wb.active
    ws.title = "Check"
    head = ["Their PO #", "Line", "Item #", "Description", "Ordered (theirs)", "Open (theirs)", "Price (theirs)", "Result", "Details",
            "Our Order", "Ordered (ours)", "Shipped (ours)", "Open (ours)", "Price (ours)", "Shipments (date, POD)"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
    fills = {"match": "DCFCE7", "diff": "FEF3C7", "missing_order": "FEE2E2", "missing_line": "FEE2E2", "ignored": "F1F5F9"}
    for row in res.get("rows", []):
        t, o = row["theirs"], row.get("ours") or {}
        ships = "; ".join(f"{s['code']} {str(s.get('ship_date') or '')[:10]}{' POD' if s.get('pod') else ''}" for s in o.get("shipments", []))
        ws.append([t["po"], t["line"], t["item"], t["desc"], t["qty_ordered"], t["qty_open"], t["price"], words.get(row["status"], row["status"]),
                   ", ".join(flag_words.get(f, f) for f in row["flags"]) + ("; " if row["flags"] and row["notes"] else "") + "; ".join(row["notes"]),
                   o.get("order_code", ""), o.get("qty"), o.get("shipped"), o.get("open"), o.get("price"), ships])
        ws.cell(row=ws.max_row, column=8).fill = PatternFill("solid", fgColor=fills.get(row["status"], "FFFFFF"))
    if res.get("ours_only"):
        ws2 = wb.create_sheet("Ours, not on theirs")
        ws2.append(["Our Order", "Line", "Item #", "Ordered", "Shipped", "Open", "Price"])
        for o in res["ours_only"]:
            ws2.append([o["order_code"], o["line_no"], o["item"], o["qty"], o["shipped"], o["open"], o["price"]])
    for sheet in wb.worksheets:
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = min(48, max(10, max(len(str(c.value or "")) for c in col) + 2))
    buf = io.BytesIO()
    wb.save(buf)
    from app.services.filenames import disposition
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": disposition(f"Open lines check - {res.get('customer', '')} - {str(r.created_at)[:10]}.xlsx", inline=False)})
