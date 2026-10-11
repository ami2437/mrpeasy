"""Accounting (services/accounting.py): statements in, every line in an account, reports. Money permission
"accounting" -- Admin and Super admin until a role is given it on Users & Roles."""
import json
from collections import defaultdict
from datetime import date, datetime
from typing import List, Optional

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import Response
from pydantic import BaseModel
from sqlalchemy import or_
from sqlalchemy.orm import Session, selectinload

from app.config.database import get_db
from app.dependencies import get_current_active_user, require_perm
from app.models import AcctAccount, AcctCheckOk, AcctComment, AcctImport, AcctOpenItem, AcctReconcile, AcctRule, AcctSheet, AcctSource, AcctSplit, AcctTxn, User
from app.routes.attachments import MAX_BYTES
from app.services import accounting as A
from app.services import acct_mapper, acct_more as M, acct_rules
from app.services.money import cents

router = APIRouter(prefix="/api/accounting", tags=["accounting"], dependencies=[Depends(require_perm("accounting"))])


def _who(u: User) -> str:
    return u.full_name or u.username


def _period(db: Session, start: Optional[str], end: Optional[str]):
    a, b = A.parse_day(start), A.parse_day(end)
    if not a or not b:
        m = A.settings(db)["fy_start_month"]
        a, b = A.fy_range(A.fy_start(A.today(), m).year, m)
    if b < a:
        raise HTTPException(status_code=400, detail="The period ends before it starts")
    return a, b


def _accounts(db: Session) -> list:
    accounts = {a.id: a for a in db.query(AcctAccount).all()}
    return [A.account_dict(a, accounts) for a, _d in A.tree_order(list(accounts.values()), accounts)]


def _fy_key(db: Session, a: date) -> str:
    m = A.settings(db)["fy_start_month"]
    return A.fy_label(A.fy_start(a, m), m)


# ---------------------------------------------------------------- setup
@router.get("/meta")
def meta(db: Session = Depends(get_db)):
    A.book_source(db)
    db.commit()
    return {"settings": A.settings(db), "years": A.years(db), "kinds": [{"key": k, "label": v[0], "qb_type": v[1], "in_pnl": v[2]} for k, v in A.KINDS.items()],
            "accounts": _accounts(db),
            "sources": [A.source_dict(db, s) for s in db.query(AcctSource).order_by(AcctSource.kind, AcctSource.id).all()],
            "today": A.today().isoformat()}


class SettingsIn(BaseModel):
    fy_start_month: Optional[int] = None
    partners: Optional[List[dict]] = None
    year: Optional[str] = None  # FY label the next three are for
    bank_balance: Optional[float] = None
    notes: Optional[str] = None
    adjust: Optional[dict] = None
    open_orders_counted: Optional[bool] = None


@router.put("/settings")
def put_settings(body: SettingsIn, db: Session = Depends(get_db), user: User = Depends(get_current_active_user)):
    s = A.settings(db)
    if body.fy_start_month is not None:
        if not 1 <= body.fy_start_month <= 12:
            raise HTTPException(status_code=400, detail="Month 1-12")
        s["fy_start_month"] = body.fy_start_month
    if body.partners is not None:
        ps = [{"name": str(p.get("name") or "").strip(), "share": float(p.get("share") or 0)} for p in body.partners if str(p.get("name") or "").strip()]
        if ps and abs(sum(p["share"] for p in ps) - 1) > 0.001:
            raise HTTPException(status_code=400, detail=f"The shares add up to {sum(p['share'] for p in ps) * 100:.1f}%, not 100%")
        s["partners"] = ps
    if body.year:
        y = dict(s.get("years", {}).get(body.year, {}))
        if body.bank_balance is not None:
            y["bank_balance"] = cents(body.bank_balance)
        if body.notes is not None:
            y["notes"] = body.notes
        if body.adjust is not None:
            y["adjust"] = {k: cents(float(v or 0)) for k, v in body.adjust.items()}
        s.setdefault("years", {})[body.year] = y
    if body.open_orders_counted is not None:
        s.setdefault("apar_groups", {})["open_orders"] = body.open_orders_counted
    return A.save_settings(db, s, _who(user))


class AccountIn(BaseModel):
    name: Optional[str] = None
    kind: Optional[str] = None
    partner: Optional[str] = None
    qb_type: Optional[str] = None
    external_id: Optional[str] = None
    note: Optional[str] = None
    opening_balance: Optional[float] = None
    is_active: Optional[bool] = None
    parent_id: Optional[int] = None  # 0 = no parent (a main account)


def _apply_account(a: AcctAccount, body: AccountIn, db: Session):
    if body.kind is not None and body.kind not in A.KINDS:
        raise HTTPException(status_code=400, detail="Unknown account type")
    if body.name is not None:
        name = body.name.strip()
        if not name:
            raise HTTPException(status_code=400, detail="Give the account a name")
        clash = next((x for x in db.query(AcctAccount).all() if x.id != a.id and x.name.strip().lower() == name.lower()), None)
        if clash:
            raise HTTPException(status_code=400, detail=f"There's already an account called {clash.name}")
        a.name = name
    for k in ("kind", "partner", "qb_type", "external_id", "note", "opening_balance", "is_active"):
        v = getattr(body, k)
        if v is not None:
            setattr(a, k, (v.strip() or None) if isinstance(v, str) else v)
    accounts = {x.id: x for x in db.query(AcctAccount).all()}
    if body.parent_id is not None:
        if not body.parent_id:
            a.parent_id = None
        else:
            parent = accounts.get(body.parent_id)
            if not parent or parent.id == a.id or (a.id and parent.id in A.descendants(a.id, accounts)):
                raise HTTPException(status_code=400, detail="An account can't be under itself or one of its own sub-accounts")
            a.parent_id = parent.id
    if a.parent_id and a.parent_id in accounts:  # a sub-account is the same type as its parent (and the same partner's)
        parent = accounts[a.parent_id]
        a.kind = parent.kind
        if parent.kind == "owner" and not a.partner:
            a.partner = next((x.partner for x in [parent] + A.ancestors(parent, accounts) if x.partner), None)
    if a.id:
        for d in A.descendants(a.id, accounts):  # a parent's type change carries down
            accounts[d].kind = a.kind
            if a.kind != "owner":
                accounts[d].partner = None
    if a.kind not in ("owner", "loan"):
        a.partner = None


@router.post("/accounts")
def add_account(body: AccountIn, db: Session = Depends(get_db)):
    a = AcctAccount(name="", kind=body.kind or "expense")
    _apply_account(a, body, db)
    if not a.name:
        raise HTTPException(status_code=400, detail="Give the account a name")
    a.qb_type = a.qb_type or A.KINDS[a.kind][1]
    db.add(a)
    db.commit()
    return A.account_dict(a, {x.id: x for x in db.query(AcctAccount).all()})


@router.put("/accounts/{account_id}")
def edit_account(account_id: int, body: AccountIn, db: Session = Depends(get_db)):
    a = db.get(AcctAccount, account_id)
    if not a:
        raise HTTPException(status_code=404, detail="Account not found")
    _apply_account(a, body, db)
    db.commit()
    return A.account_dict(a, {x.id: x for x in db.query(AcctAccount).all()})


@router.post("/accounts/{account_id}/merge")
def merge_account(account_id: int, into_id: int, db: Session = Depends(get_db)):
    """Move every line, part and rule of this account into another, then remove it (two names for one thing)."""
    a, b = db.get(AcctAccount, account_id), db.get(AcctAccount, into_id)
    if not a or not b or a.id == b.id:
        raise HTTPException(status_code=400, detail="Pick another account to merge into")
    n = db.query(AcctTxn).filter(AcctTxn.account_id == a.id).update({AcctTxn.account_id: b.id})
    n += db.query(AcctSplit).filter(AcctSplit.account_id == a.id).update({AcctSplit.account_id: b.id})
    db.query(AcctRule).filter(AcctRule.account_id == a.id).update({AcctRule.account_id: b.id})
    accounts = {x.id: x for x in db.query(AcctAccount).all()}
    if b.id in A.descendants(a.id, accounts):
        raise HTTPException(status_code=400, detail=f"{b.name} is one of {a.name}'s own sub-accounts -- move it out first")
    db.query(AcctAccount).filter(AcctAccount.parent_id == a.id).update({AcctAccount.parent_id: b.id})  # its sub-accounts go along
    b.opening_balance = (b.opening_balance or 0) + (a.opening_balance or 0)
    db.delete(a)
    db.commit()
    return {"moved": n}


@router.delete("/accounts/{account_id}")
def delete_account(account_id: int, db: Session = Depends(get_db)):
    a = db.get(AcctAccount, account_id)
    if not a:
        raise HTTPException(status_code=404, detail="Account not found")
    used = db.query(AcctTxn).filter(AcctTxn.account_id == a.id).count() + db.query(AcctSplit).filter(AcctSplit.account_id == a.id).count()
    if used:
        raise HTTPException(status_code=400, detail=f"{used} line(s) are in {a.name} -- merge it into another account, or switch it off")
    subs = db.query(AcctAccount).filter(AcctAccount.parent_id == a.id).count()
    if subs:
        raise HTTPException(status_code=400, detail=f"{a.name} has {subs} sub-account(s) -- move or delete them first")
    db.query(AcctRule).filter(AcctRule.account_id == a.id).delete()
    db.delete(a)
    db.commit()
    return {"ok": True}


class SourceIn(BaseModel):
    name: Optional[str] = None
    kind: Optional[str] = None
    last4: Optional[str] = None
    opening_balance: Optional[float] = None
    opening_date: Optional[str] = None
    external_id: Optional[str] = None
    is_active: Optional[bool] = None


def _apply_source(s: AcctSource, body: SourceIn):
    if body.kind is not None:
        if body.kind not in ("bank", "card"):
            raise HTTPException(status_code=400, detail="A bank account or a credit card")
        s.kind = body.kind
    for k in ("name", "last4", "external_id"):
        v = getattr(body, k)
        if v is not None:
            setattr(s, k, v.strip() or None)
    if body.opening_balance is not None:
        s.opening_balance = cents(body.opening_balance)
    if body.opening_date is not None:
        d = A.parse_day(body.opening_date)
        s.opening_date = A._dt(d) if d else None
    if body.is_active is not None:
        s.is_active = body.is_active
    if not s.name:
        raise HTTPException(status_code=400, detail="Give it a name (e.g. Wells Fargo Checking)")


@router.post("/sources")
def add_source(body: SourceIn, db: Session = Depends(get_db)):
    s = AcctSource(name="", kind=body.kind or "bank")
    _apply_source(s, body)
    db.add(s)
    db.commit()
    return A.source_dict(db, s)


@router.put("/sources/{source_id}")
def edit_source(source_id: int, body: SourceIn, db: Session = Depends(get_db)):
    s = db.get(AcctSource, source_id)
    if not s:
        raise HTTPException(status_code=404, detail="Not found")
    if s.kind == "book":
        body.kind = None
    _apply_source(s, body)
    db.commit()
    return A.source_dict(db, s)


@router.delete("/sources/{source_id}")
def delete_source(source_id: int, db: Session = Depends(get_db)):
    s = db.get(AcctSource, source_id)
    if not s or s.kind == "book":
        raise HTTPException(status_code=400, detail="Book Entries stays")
    n = db.query(AcctTxn).filter(AcctTxn.source_id == s.id).count()
    if n:
        raise HTTPException(status_code=400, detail=f"{s.name} has {n} line(s) -- switch it off instead")
    db.delete(s)
    db.commit()
    return {"ok": True}


# ---------------------------------------------------------------- statement files
async def _read(file: UploadFile) -> bytes:
    data = await file.read()
    if len(data) > MAX_BYTES:
        raise HTTPException(status_code=400, detail="File too big")
    if not (file.filename or "").lower().endswith((".csv", ".xlsx", ".xlsm", ".txt")):
        raise HTTPException(status_code=400, detail="A CSV or Excel (.xlsx) statement")
    return data


@router.post("/import/preview")
async def import_preview(source_id: int = Form(...), layout: Optional[str] = Form(None), file: UploadFile = File(...),
                         db: Session = Depends(get_db)):
    data = await _read(file)
    return A.preview(db, data, file.filename, source_id, json.loads(layout) if layout else None)


@router.post("/import")
async def import_statement(source_id: int = Form(...), layout: Optional[str] = Form(None), file: UploadFile = File(...),
                           db: Session = Depends(get_db), user: User = Depends(get_current_active_user)):
    data = await _read(file)
    return A.import_file(db, data, file.filename, source_id, json.loads(layout) if layout else None, _who(user))


@router.get("/imports")
def imports(db: Session = Depends(get_db)):
    src = {s.id: s.name for s in db.query(AcctSource).all()}
    return [{"id": i.id, "source": src.get(i.source_id, ""), "filename": i.filename, "kind": i.kind, "rows": i.rows, "added": i.added,
             "skipped": i.skipped, "matched": i.matched, "auto": i.auto, "created_by": i.created_by, "created_at": i.created_at,
             "still_there": db.query(AcctTxn).filter(AcctTxn.import_id == i.id).count()}
            for i in db.query(AcctImport).order_by(AcctImport.id.desc()).limit(100).all()]


@router.delete("/imports/{import_id}")
def undo_import(import_id: int, db: Session = Depends(get_db)):
    return {"removed": A.undo_import(db, import_id)}


# ---------------------------------------------------------------- lines
def _search_params(start, end, **kw) -> dict:
    return {"start": start, "end": end, **{k: v for k, v in kw.items() if v not in (None, "")}}


@router.get("/txns")
def txns(start: Optional[str] = None, end: Optional[str] = None, source_id: Optional[int] = None, account_id: Optional[int] = None,
         kind: Optional[str] = None, status: Optional[str] = None, q: Optional[str] = None, all_dates: bool = False,
         date_from: Optional[str] = None, date_to: Optional[str] = None, amin: Optional[str] = None, amax: Optional[str] = None,
         amt: Optional[str] = None, sign: Optional[str] = None, has_note: bool = False, how: Optional[str] = None,
         sort: Optional[str] = None, group: Optional[str] = None, limit: int = 300, offset: int = 0, db: Session = Depends(get_db)):
    """Bank-style search (services/acct_more.py): status unfiled|review / auto / expected / typed / split / sent_back;
    q understands amounts ("1072.50", ">500", "500-1000", "-300"), words and "phrases". Totals, stats and groups are for
    EVERYTHING that matches, the rows for the page asked for."""
    if not all_dates and not (date_from or date_to):
        a, b = _period(db, start, end)
        start, end = a.isoformat(), b.isoformat()
    r = M.search(db, _search_params(start, end, source_id=source_id, account_id=account_id, kind=kind, status=status, q=q,
                                    all_dates=all_dates, date_from=date_from, date_to=date_to, amin=amin, amax=amax, amt=amt,
                                    sign=sign, has_note=has_note, how=how, sort=sort, group=group, limit=limit, offset=offset))
    rows = r.pop("rows_obj")
    r.pop("all_rows")
    accounts, sources = A.maps(db)
    unfiled = [t for t in rows if not t.account_id and not t.splits]
    sugg = A.suggestions(db, unfiled) if unfiled else {}
    out = []
    for t in rows:
        d = A.txn_dict(t, accounts, sources)
        was = json.loads(t.sent_back_from) if t.sent_back_from and not t.account_id and not t.splits else None
        if was and was.get("account_id") in accounts:  # sent back: where it was is the top suggestion
            d["suggest"] = {"account_id": was["account_id"], "account": A.path(accounts[was["account_id"]], accounts), "confidence": "guess",
                            "why": "where it was before it was sent back"}
        elif t.id in sugg:
            sg = sugg[t.id]
            d["suggest"] = {**sg, "account": A.path(accounts[sg["account_id"]], accounts) if sg["account_id"] in accounts else ""}
        if t.id in r["balance"]:
            d["balance"] = r["balance"][t.id]
        out.append(d)
    r.pop("balance")
    return {**r, "rows": out}


@router.get("/txns/export.xlsx")
def txns_export(start: Optional[str] = None, end: Optional[str] = None, source_id: Optional[int] = None, account_id: Optional[int] = None,
                kind: Optional[str] = None, status: Optional[str] = None, q: Optional[str] = None, all_dates: bool = False,
                date_from: Optional[str] = None, date_to: Optional[str] = None, amin: Optional[str] = None, amax: Optional[str] = None,
                amt: Optional[str] = None, sign: Optional[str] = None, has_note: bool = False, how: Optional[str] = None,
                sort: Optional[str] = None, db: Session = Depends(get_db)):
    if not all_dates and not (date_from or date_to):
        a, b = _period(db, start, end)
        start, end = a.isoformat(), b.isoformat()
    r = M.search(db, _search_params(start, end, source_id=source_id, account_id=account_id, kind=kind, status=status, q=q,
                                    all_dates=all_dates, date_from=date_from, date_to=date_to, amin=amin, amax=amax, amt=amt,
                                    sign=sign, has_note=has_note, how=how, sort=sort, limit=1, offset=0))
    return Response(M.found_xlsx(db, r["all_rows"]), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": 'attachment; filename="Transactions.xlsx"'})


class SendBackIn(BaseModel):
    ids: List[int]
    reason: Optional[str] = None


@router.post("/txns/send-back")
def send_back(body: SendBackIn, db: Session = Depends(get_db), user: User = Depends(get_current_active_user)):
    """Back to Review: out of its account (where it was is kept as the top suggestion), with an optional reason."""
    return {"done": M.send_back(db, body.ids, body.reason, _who(user))}


@router.post("/txns/undo-send-back")
def undo_send_back(body: SendBackIn, db: Session = Depends(get_db), user: User = Depends(get_current_active_user)):
    return {"done": M.undo_send_back(db, body.ids, _who(user))}


class TxnIn(BaseModel):
    source_id: Optional[int] = None
    date: Optional[str] = None
    amount: Optional[float] = None
    description: Optional[str] = None
    payee: Optional[str] = None
    account_id: Optional[int] = None
    splits: Optional[List[dict]] = None
    note: Optional[str] = None
    expected: Optional[bool] = None
    learn: bool = True  # after a person files a line, file the look-alikes AT-HUB is now sure about


@router.post("/txns")
def add_txn(body: TxnIn, db: Session = Depends(get_db), user: User = Depends(get_current_active_user)):
    """A typed-in line: on a bank account and not on a statement yet (expected -- matched when the statement brings it),
    or a book entry that never will be (year-end figures, adjustments, interest to add)."""
    d = A.parse_day(body.date)
    if not d or body.amount in (None, 0):
        raise HTTPException(status_code=400, detail="A date and an amount (+ money in, - money out)")
    src = db.get(AcctSource, body.source_id) if body.source_id else A.book_source(db)
    if not src:
        raise HTTPException(status_code=400, detail="Account not found")
    desc = (body.description or "").strip() or (body.payee or "").strip()
    if not desc:
        raise HTTPException(status_code=400, detail="Say what it is")
    t = AcctTxn(source_id=src.id, date=A._dt(d), amount=cents(body.amount), description=desc, payee=(body.payee or "").strip() or acct_mapper.payee_label(desc),
                note=(body.note or "").strip() or None, origin="manual", expected=bool(body.expected) and src.kind != "book",
                created_by=_who(user))
    db.add(t)
    db.flush()
    if body.account_id or body.splits:
        A.set_account(db, t, body.account_id, body.splits, _who(user))
    db.commit()
    accounts, sources = A.maps(db)
    return A.txn_dict(t, accounts, sources)


@router.put("/txns/{txn_id}")
def edit_txn(txn_id: int, body: TxnIn, db: Session = Depends(get_db), user: User = Depends(get_current_active_user)):
    t = db.get(AcctTxn, txn_id)
    if not t:
        raise HTTPException(status_code=404, detail="Line not found")
    fields = body.model_dump(exclude_unset=True)
    src = db.get(AcctSource, t.source_id)
    if t.origin == "manual" or src.kind == "book":  # typed in / book entries can be changed; a statement line's date / amount / wording can't
        if "date" in fields:
            d = A.parse_day(body.date)
            if not d:
                raise HTTPException(status_code=400, detail="Date?")
            t.date = A._dt(d)
        if "amount" in fields and body.amount:
            t.amount = cents(body.amount)
        if "description" in fields and (body.description or "").strip():
            t.description = body.description.strip()
        if "expected" in fields:
            t.expected = bool(body.expected) and src.kind != "book"
        if "source_id" in fields and body.source_id and db.get(AcctSource, body.source_id):
            t.source_id = body.source_id
    elif {"date", "amount", "description"} & {k for k, v in fields.items() if v is not None}:
        raise HTTPException(status_code=400, detail="This line came from a statement -- its date, amount and wording stay as the bank has them")
    if "payee" in fields:
        t.payee = (body.payee or "").strip() or None
    if "note" in fields:
        t.note = (body.note or "").strip() or None
    filed_more = taken_back = 0
    alike = []
    if "account_id" in fields or "splits" in fields:
        old, was_auto = t.account_id, (t.how in ("auto", "rule") and not t.reviewed)
        A.set_account(db, t, body.account_id, body.splits, _who(user))
        db.flush()
        if body.learn and (body.account_id or body.splits):
            if was_auto and old and old != t.account_id:
                taken_back = A.recheck_after_correction(db, t, old)
            filed_more = A.autofill(db)
            db.flush()
            alike = A.look_alikes(db, t)
    t.updated_by = _who(user)
    db.commit()
    accounts, sources = A.maps(db)
    return {**A.txn_dict(t, accounts, sources), "filed_more": filed_more, "taken_back": taken_back, "alike": alike}


class BulkIn(BaseModel):
    ids: List[int]
    action: str  # file | review | unfile | delete
    account_id: Optional[int] = None


@router.post("/txns/bulk")
def bulk(body: BulkIn, db: Session = Depends(get_db), user: User = Depends(get_current_active_user)):
    ts = db.query(AcctTxn).filter(AcctTxn.id.in_(body.ids)).all()
    n = skipped = 0
    for t in ts:
        if body.action == "file":
            if not body.account_id:
                raise HTTPException(status_code=400, detail="Pick the account")
            A.set_account(db, t, body.account_id, None, _who(user))
        elif body.action == "review":
            t.reviewed = True
            t.how = "person" if t.how in ("auto", "rule") and t.account_id else t.how
        elif body.action == "unfile":
            A.set_account(db, t, None, None, _who(user))
        elif body.action == "delete":
            if t.origin != "manual" and db.get(AcctSource, t.source_id).kind != "book":
                skipped += 1
                continue
            db.delete(t)
        else:
            raise HTTPException(status_code=400, detail="Unknown action")
        n += 1
    db.flush()
    more = A.autofill(db) if body.action == "file" else 0
    db.commit()
    return {"done": n, "skipped": skipped, "filed_more": more}


@router.delete("/txns/{txn_id}")
def delete_txn(txn_id: int, db: Session = Depends(get_db)):
    """Typed-in lines only -- a statement line is undone with its whole import (Statements tab)."""
    t = db.get(AcctTxn, txn_id)
    if not t:
        raise HTTPException(status_code=404, detail="Line not found")
    if t.origin != "manual" and db.get(AcctSource, t.source_id).kind != "book":
        raise HTTPException(status_code=400, detail="This line came from a statement -- undo the import it came with instead")
    db.delete(t)
    db.commit()
    return {"ok": True}


# ---------------------------------------------------------------- rules (services/acct_rules.py)
class RuleIn(BaseModel):
    name: Optional[str] = None
    match: Optional[str] = "all"  # all | any
    conditions: Optional[List[dict]] = None  # [{field, op, value}]
    contains: Optional[str] = None  # the old one-box form: words in the line (made into a condition)
    sign: Optional[str] = None  # in | out | None
    min_amount: Optional[float] = None
    max_amount: Optional[float] = None
    source_ids: Optional[List[int]] = None
    account_id: Optional[int] = None
    splits: Optional[List[dict]] = None  # [{account_id, pct | amount}]
    set_payee: Optional[str] = None
    set_note: Optional[str] = None
    action: Optional[str] = "file"  # file | suggest
    enabled: Optional[bool] = True


def _spec_in(body: RuleIn) -> dict:
    d = body.model_dump()
    if not d.get("conditions"):
        d["conditions"] = None  # spec() builds them from contains / min / max
    d["match_mode"] = d.pop("match") or "all"
    s_ = acct_rules.spec(d)
    s_["splits"] = [{"account_id": int(x["account_id"]), **({"pct": float(x["pct"])} if x.get("pct") not in (None, "") else {"amount": abs(float(x.get("amount") or 0))})}
                    for x in (body.splits or []) if x.get("account_id")]
    if s_["sign"] not in ("in", "out"):
        s_["sign"] = None
    if s_["action"] not in ("file", "suggest"):
        s_["action"] = "file"
    return s_


def _rule_row(db: Session, r: AcctRule, s_: dict) -> None:
    err = acct_rules.check(s_)
    if err:
        raise HTTPException(status_code=400, detail=err)
    for aid in [s_["account_id"]] + [x["account_id"] for x in s_["splits"]]:
        if aid and not db.get(AcctAccount, aid):
            raise HTTPException(status_code=400, detail="Account not found")
    r.name = (s_["name"] or "").strip() or None
    r.match_mode = s_["match"]
    r.conditions = json.dumps(s_["conditions"])
    r.source_ids = json.dumps(s_["source_ids"])
    r.splits = json.dumps(s_["splits"])
    r.set_payee, r.set_note = s_["set_payee"], s_["set_note"]
    r.action, r.enabled, r.sign = s_["action"], s_["enabled"], s_["sign"]
    r.account_id = s_["account_id"] or s_["splits"][0]["account_id"]
    r.contains = " / ".join(str(c["value"]) for c in s_["conditions"] if c.get("field") in acct_rules.TEXT_FIELDS)[:200]
    r.min_amount = r.max_amount = None  # the conditions say it now


def _rule_out(r: AcctRule, accounts, sources) -> dict:
    s_ = acct_rules.spec(r)
    return {**s_, "id": r.id, "hits": r.hits, "created_by": r.created_by, "contains": r.contains,
            "account": A.path(accounts[r.account_id], accounts) if r.account_id in accounts else "",
            "splits": [{**x, "account": A.path(accounts[x["account_id"]], accounts) if x["account_id"] in accounts else "?"} for x in s_["splits"]],
            "summary": acct_rules.describe(s_, {k: v.name for k, v in sources.items()})}


@router.get("/rules")
def rules(db: Session = Depends(get_db)):
    accounts, sources = A.maps(db)
    rows = db.query(AcctRule).all()
    rows.sort(key=lambda r: (r.priority if r.priority is not None else 10 ** 6, r.id))
    filed = defaultdict(lambda: [0, 0.0])
    for rid, amt in db.query(AcctTxn.rule_id, AcctTxn.amount).filter(AcctTxn.rule_id.isnot(None)).all():
        filed[rid][0] += 1
        filed[rid][1] += amt
    return [{**_rule_out(r, accounts, sources), "filed_lines": filed[r.id][0], "filed_total": cents(filed[r.id][1])} for r in rows]


@router.get("/rules/meta")
def rules_meta():
    return {"text_fields": acct_rules.TEXT_FIELDS, "num_fields": acct_rules.NUM_FIELDS, "text_ops": acct_rules.TEXT_OPS, "num_ops": acct_rules.NUM_OPS}


@router.post("/rules/preview")
def preview_rule(body: RuleIn, db: Session = Depends(get_db)):
    """The lines a rule would cover (before or after saving it), grouped -- for the live count and the pop-up."""
    s_ = _spec_in(body)
    if not [c for c in s_["conditions"] if c.get("value") not in (None, "", [])]:
        raise HTTPException(status_code=400, detail="Add a condition")
    return A.rule_lines(db, s_, s_["account_id"])


def _next_priority(db: Session) -> int:
    return max([r.priority or 0 for r in db.query(AcctRule).all()] + [0]) + 10


@router.post("/rules")
def add_rule(body: RuleIn, db: Session = Depends(get_db), user: User = Depends(get_current_active_user)):
    """Saved only -- nothing is filed until someone picks the lines in the pop-up (POST /rules/{id}/apply). Lines
    brought in later that match are filed by it by themselves (or suggested, if it's a suggest-only rule)."""
    r = AcctRule(created_by=_who(user), priority=_next_priority(db), contains="")
    _rule_row(db, r, _spec_in(body))
    db.add(r)
    db.commit()
    return {"id": r.id}


@router.put("/rules/{rule_id}")
def edit_rule(rule_id: int, body: RuleIn, db: Session = Depends(get_db)):
    r = db.get(AcctRule, rule_id)
    if not r:
        raise HTTPException(status_code=404, detail="Rule not found")
    _rule_row(db, r, _spec_in(body))
    db.commit()
    return {"id": r.id}


class OnOff(BaseModel):
    enabled: bool


@router.put("/rules/{rule_id}/enabled")
def rule_on_off(rule_id: int, body: OnOff, db: Session = Depends(get_db)):
    r = db.get(AcctRule, rule_id)
    if not r:
        raise HTTPException(status_code=404, detail="Rule not found")
    r.enabled = body.enabled
    db.commit()
    return {"ok": True}


@router.post("/rules/{rule_id}/copy")
def copy_rule(rule_id: int, db: Session = Depends(get_db), user: User = Depends(get_current_active_user)):
    r = db.get(AcctRule, rule_id)
    if not r:
        raise HTTPException(status_code=404, detail="Rule not found")
    c = AcctRule(created_by=_who(user), priority=(r.priority or 0) + 1, contains=r.contains or "")
    s_ = acct_rules.spec(r)
    s_["name"] = ((r.name or r.contains or "Rule") + " (copy)")[:120]
    s_["enabled"] = False  # a copy starts switched off, so it can't fight the original
    _rule_row(db, c, s_)
    db.add(c)
    db.commit()
    return {"id": c.id}


class OrderIn(BaseModel):
    ids: List[int]


@router.put("/rules-order")
def rules_order(body: OrderIn, db: Session = Depends(get_db)):
    """The order rules run in (first match wins): the ids top to bottom."""
    for n, rid in enumerate(body.ids):
        r = db.get(AcctRule, rid)
        if r:
            r.priority = (n + 1) * 10
    db.commit()
    return {"ok": True}


@router.get("/rules/{rule_id}/lines")
def rule_lines(rule_id: int, db: Session = Depends(get_db)):
    r = db.get(AcctRule, rule_id)
    if not r:
        raise HTTPException(status_code=404, detail="Rule not found")
    return A.rule_lines(db, acct_rules.spec(r))


class ApplyIn(BaseModel):
    ids: List[int]


@router.post("/rules/{rule_id}/apply")
def apply_rule(rule_id: int, body: ApplyIn, db: Session = Depends(get_db), user: User = Depends(get_current_active_user)):
    r = db.get(AcctRule, rule_id)
    if not r:
        raise HTTPException(status_code=404, detail="Rule not found")
    return {"filed": A.apply_rule(db, r, body.ids, _who(user))}


@router.delete("/rules/{rule_id}")
def delete_rule(rule_id: int, db: Session = Depends(get_db)):
    r = db.get(AcctRule, rule_id)
    if r:
        db.delete(r)
        db.commit()
    return {"ok": True}


@router.get("/rules-export.xlsx")
def export_rules(db: Session = Depends(get_db)):
    """Every rule in order, one row each (like QuickBooks' rule export) -- to keep, or for a QuickBooks hand-over."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    accounts, sources = A.maps(db)
    wb = Workbook()
    ws = wb.active
    ws.title = "Rules"
    ws.append(["Order", "Name", "On", "Money", "Bank Accounts", "Match", "Conditions", "Account", "Split", "Set Payee", "Set Note", "Files Or Suggests", "Used"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for n, r in enumerate(rules(db), start=1):
        ws.append([n, r["name"], "yes" if r["enabled"] else "no", {"in": "In", "out": "Out"}.get(r["sign"] or "", "Both"),
                   ", ".join(sources[i].name for i in r["source_ids"] if i in sources) or "All", "All" if r["match"] != "any" else "Any",
                   r["summary"], r["account"], ", ".join(f"{x['account']} {x['pct']}%" if "pct" in x else f"{x['account']} ${x['amount']}" for x in r["splits"]),
                   r["set_payee"] or "", r["set_note"] or "", "Suggests" if r["action"] == "suggest" else "Files", r["hits"]])
    for col, w in zip("ABCDEFGHIJKLM", (7, 24, 5, 8, 22, 7, 60, 28, 30, 18, 22, 14, 7)):
        ws.column_dimensions[col].width = w
    import io as _io
    buf = _io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": 'attachment; filename="Accounting Rules.xlsx"'})


@router.get("/rule-suggestions")
def rule_suggestions(db: Session = Depends(get_db)):
    return A.suggest_rules(db)


class DismissIn(BaseModel):
    key: str


@router.post("/rule-suggestions/dismiss")
def dismiss_rule_suggestion(body: DismissIn, db: Session = Depends(get_db), user: User = Depends(get_current_active_user)):
    A.dismiss_suggestion(db, body.key, _who(user))
    return {"ok": True}


# ---------------------------------------------------------------- what AT-HUB learned
@router.get("/learned")
def learned(db: Session = Depends(get_db)):
    return A.learned(db)


class ForgetIn(BaseModel):
    key: str
    forget: bool = True


@router.post("/learned/forget")
def forget(body: ForgetIn, db: Session = Depends(get_db)):
    """Stop (or start again) learning from a payee's lines -- the lines keep their accounts."""
    return {"lines": A.forget(db, body.key, body.forget)}


# ---------------------------------------------------------------- reports
@router.get("/overview")
def overview(start: Optional[str] = None, end: Optional[str] = None, db: Session = Depends(get_db)):
    a, b = _period(db, start, end)
    pa, pb = M.prev_range(a, b)
    prev = A.overview(db, pa, pb)
    return {**A.overview(db, a, b), "prev": {k: prev[k] for k in ("money_in", "money_out", "income", "cogs", "expense", "profit")},
            "prev_period": {"start": pa.isoformat(), "end": pb.isoformat()}}


@router.get("/pnl")
def pnl(start: Optional[str] = None, end: Optional[str] = None, by: str = "month", db: Session = Depends(get_db)):
    a, b = _period(db, start, end)
    if by not in ("month", "quarter", "fy"):
        raise HTTPException(status_code=400, detail="month, quarter or fy")
    cur = A.pnl(db, a, b, by)
    pa, pb = M.prev_range(a, b)
    prev = A.pnl(db, pa, pb, "fy")
    prev_acc = {r["account_id"]: r["total"] for sec in prev["sections"] for r in sec["rows"] if not r.get("own_only")}
    for sec in cur["sections"]:
        for r in sec["rows"]:
            if not r.get("own_only"):
                r["prev"] = prev_acc.get(r["account_id"], 0.0)
    cur["prev"] = {"sections": {sec["kind"]: sec["total"] for sec in prev["sections"]}, "profit": prev["profit_total"], "gross": prev["gross_total"],
                   "start": pa.isoformat(), "end": pb.isoformat()}
    return cur


@router.get("/accounts-summary")
def accounts_summary(start: Optional[str] = None, end: Optional[str] = None, db: Session = Depends(get_db)):
    a, b = _period(db, start, end)
    return A.account_summary(db, a, b)


@router.get("/loans")
def loans(end: Optional[str] = None, db: Session = Depends(get_db)):
    return A.loans(db, A.parse_day(end))


@router.get("/partners")
def partners(start: Optional[str] = None, end: Optional[str] = None, db: Session = Depends(get_db)):
    a, b = _period(db, start, end)
    return {**A.partners(db, a, b, _fy_key(db, a)), "year": _fy_key(db, a)}


@router.get("/cashflow")
def cashflow(start: Optional[str] = None, end: Optional[str] = None, source_id: Optional[int] = None, db: Session = Depends(get_db)):
    a, b = _period(db, start, end)
    return A.cashflow(db, a, b, source_id)


@router.get("/open-items")
def open_items(db: Session = Depends(get_db)):
    return A.open_items(db)


class OpenItemIn(BaseModel):
    id: Optional[int] = None
    source: Optional[str] = None
    source_id: Optional[int] = None
    side: Optional[str] = None
    party: Optional[str] = None
    ref: Optional[str] = None
    amount: Optional[float] = None
    item_date: Optional[str] = None
    due_date: Optional[str] = None
    status: Optional[str] = None
    paid_date: Optional[str] = None
    note: Optional[str] = None


@router.post("/open-items")
def set_open_item(body: OpenItemIn, db: Session = Depends(get_db), user: User = Depends(get_current_active_user)):
    o = A.set_open_item(db, body.model_dump(exclude_unset=True), _who(user))
    return {"id": o.id}


@router.delete("/open-items/{item_id}")
def delete_open_item(item_id: int, db: Session = Depends(get_db)):
    """A typed-in item goes; a change to an AT-HUB figure is dropped (AT-HUB's figure again)."""
    o = db.get(AcctOpenItem, item_id)
    if o:
        db.delete(o)
        db.commit()
    return {"ok": True}


@router.get("/export.xlsx")
def export_xlsx(start: Optional[str] = None, end: Optional[str] = None, db: Session = Depends(get_db)):
    a, b = _period(db, start, end)
    data = A.export_xlsx(db, a, b, _fy_key(db, a))
    name = f"Accounting {a.isoformat()} to {b.isoformat()}.xlsx"
    return Response(data, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{name}"'})


@router.get("/export/quickbooks.csv")
def export_qb(source_id: int, start: Optional[str] = None, end: Optional[str] = None, db: Session = Depends(get_db)):
    a, b = _period(db, start, end)
    s = db.get(AcctSource, source_id)
    if not s:
        raise HTTPException(status_code=404, detail="Not found")
    return Response(A.export_quickbooks_csv(db, source_id, a, b), media_type="text/csv",
                    headers={"Content-Disposition": f'attachment; filename="{s.name} {a.isoformat()} to {b.isoformat()} QuickBooks.csv"'})


# ---------------------------------------------------------------- comments (every section, per period)
class CommentIn(BaseModel):
    section: str
    text: str
    period_key: Optional[str] = None
    ref_type: Optional[str] = None
    ref_id: Optional[int] = None


@router.get("/comments")
def list_comments(section: Optional[str] = None, period_key: Optional[str] = None, all_periods: bool = False, db: Session = Depends(get_db)):
    return M.comments(db, section, period_key, all_periods)


@router.post("/comments")
def add_comment(body: CommentIn, db: Session = Depends(get_db), user: User = Depends(get_current_active_user)):
    if not body.text.strip():
        raise HTTPException(status_code=400, detail="Write the comment")
    c = AcctComment(section=body.section, period_key=body.period_key, ref_type=body.ref_type, ref_id=body.ref_id, text=body.text.strip(), created_by=_who(user))
    db.add(c)
    db.commit()
    return M.comment_dict(c)


@router.put("/comments/{comment_id}")
def edit_comment(comment_id: int, body: CommentIn, db: Session = Depends(get_db), user: User = Depends(get_current_active_user)):
    c = db.get(AcctComment, comment_id)
    if not c:
        raise HTTPException(status_code=404, detail="Comment not found")
    c.text, c.updated_by, c.updated_at = body.text.strip(), _who(user), datetime.utcnow()
    db.commit()
    return M.comment_dict(c)


@router.delete("/comments/{comment_id}")
def delete_comment(comment_id: int, db: Session = Depends(get_db)):
    c = db.get(AcctComment, comment_id)
    if c:
        db.delete(c)
        db.commit()
    return {"ok": True}


# ---------------------------------------------------------------- Ignored (duplicates / estimates)
@router.get("/ignored")
def list_ignored(import_id: Optional[int] = None, status: Optional[str] = None, db: Session = Depends(get_db)):
    return M.ignored(db, import_id, status)


class DecideIn(BaseModel):
    ids: List[int]
    action: str  # ignore | keep | restore


@router.post("/ignored/decide")
def decide_ignored(body: DecideIn, db: Session = Depends(get_db), user: User = Depends(get_current_active_user)):
    return {"done": M.decide_ignored(db, body.ids, body.action, _who(user))}


# ---------------------------------------------------------------- Check My Books
@router.get("/check")
def check_books(start: Optional[str] = None, end: Optional[str] = None, all_dates: bool = False, db: Session = Depends(get_db)):
    a, b = (None, None) if all_dates else _period(db, start, end)
    return M.check_books(db, a, b)


class CheckOkIn(BaseModel):
    txn_id: int
    check: str


@router.post("/check/ok")
def check_ok(body: CheckOkIn, db: Session = Depends(get_db), user: User = Depends(get_current_active_user)):
    """"It's right": this finding isn't shown again for this line."""
    if not db.query(AcctCheckOk).filter(AcctCheckOk.txn_id == body.txn_id, AcctCheckOk.check == body.check).first():
        db.add(AcctCheckOk(txn_id=body.txn_id, check=body.check, created_by=_who(user)))
        db.commit()
    return {"ok": True}


# ---------------------------------------------------------------- Partner Accounts, Payees, 1099, Recurring
@router.get("/partner-accounts")
def partner_accounts(start: Optional[str] = None, end: Optional[str] = None, db: Session = Depends(get_db)):
    a, b = _period(db, start, end)
    return M.partner_accounts(db, a, b, _fy_key(db, a))


@router.get("/payees")
def payees(start: Optional[str] = None, end: Optional[str] = None, db: Session = Depends(get_db)):
    a, b = _period(db, start, end)
    return M.payees(db, a, b)


@router.get("/tax1099")
def tax1099(year: int, threshold: float = 600.0, db: Session = Depends(get_db)):
    return M.tax1099(db, year, threshold)


class MarkIn(BaseModel):
    year: int
    payee: str
    mark: Optional[str] = None  # yes | no | None


@router.post("/tax1099/mark")
def mark1099(body: MarkIn, db: Session = Depends(get_db), user: User = Depends(get_current_active_user)):
    M.mark1099(db, body.year, body.payee, body.mark, _who(user))
    return {"ok": True}


@router.get("/tax1099/export.xlsx")
def export1099(year: int, threshold: float = 600.0, db: Session = Depends(get_db)):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    d = M.tax1099(db, year, threshold)
    wb = Workbook()
    ws = wb.active
    ws.title = f"1099 {year}"
    ws.append(["Payee", "Paid In " + str(year), "Payments", "First", "Last", "Accounts", "Needs A 1099"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in d["rows"]:
        ws.append([r["payee"], r["total"], r["count"], r["first"], r["last"], ", ".join(r["accounts"]), {"yes": "Yes", "no": "No"}.get(r["mark"] or "", "?")])
    for col, w in zip("ABCDEFG", (34, 14, 10, 12, 12, 40, 12)):
        ws.column_dimensions[col].width = w
    import io as _io
    buf = _io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="1099 list {year}.xlsx"'})


@router.get("/recurring")
def recurring(db: Session = Depends(get_db)):
    return M.recurring(db)


# ---------------------------------------------------------------- scratch sheets: Projections + Loans quick calculations (NOT the books)
class SheetIn(BaseModel):
    kind: Optional[str] = "projection"  # projection | loancalc
    name: Optional[str] = None
    doc: Optional[dict] = None


@router.get("/sheets")
def list_sheets(kind: str = "projection", db: Session = Depends(get_db), user: User = Depends(get_current_active_user)):
    q = db.query(AcctSheet).filter(AcctSheet.kind == kind)
    if kind == "loancalc":  # each person's own scratch work
        q = q.filter(AcctSheet.owner == user.username)
    return [M.sheet_dict(x) for x in q.order_by(AcctSheet.id).all()]


@router.post("/sheets")
def add_sheet(body: SheetIn, db: Session = Depends(get_db), user: User = Depends(get_current_active_user)):
    kind = body.kind if body.kind in ("projection", "loancalc") else "projection"
    x = AcctSheet(kind=kind, name=(body.name or "").strip() or ("Quick Calculations" if kind == "loancalc" else "New Projection"),
                  owner=user.username if kind == "loancalc" else None, doc=json.dumps(body.doc or {}), created_by=_who(user))
    db.add(x)
    db.commit()
    return M.sheet_dict(x)


@router.put("/sheets/{sheet_id}")
def save_sheet(sheet_id: int, body: SheetIn, db: Session = Depends(get_db), user: User = Depends(get_current_active_user)):
    x = db.get(AcctSheet, sheet_id)
    if not x:
        raise HTTPException(status_code=404, detail="Not found")
    if body.name is not None and body.name.strip():
        x.name = body.name.strip()
    if body.doc is not None:
        x.doc = json.dumps(body.doc)
    x.updated_by, x.updated_at = _who(user), datetime.utcnow()
    db.commit()
    return M.sheet_dict(x)


@router.post("/sheets/{sheet_id}/copy")
def copy_sheet(sheet_id: int, db: Session = Depends(get_db), user: User = Depends(get_current_active_user)):
    x = db.get(AcctSheet, sheet_id)
    if not x:
        raise HTTPException(status_code=404, detail="Not found")
    c = AcctSheet(kind=x.kind, name=f"{x.name} (copy)", owner=x.owner, doc=x.doc, created_by=_who(user))
    db.add(c)
    db.commit()
    return M.sheet_dict(c)


@router.delete("/sheets/{sheet_id}")
def delete_sheet(sheet_id: int, db: Session = Depends(get_db)):
    x = db.get(AcctSheet, sheet_id)
    if x:
        db.delete(x)
        db.commit()
    return {"ok": True}


@router.get("/sheets-books")
def sheets_books(months: int = 6, db: Session = Depends(get_db)):
    """Today's figures from the books for "Fill from the books"."""
    return M.books_snapshot(db, max(1, min(months, 24)))


@router.get("/sheets/{sheet_id}/export.xlsx")
def sheet_export(sheet_id: int, db: Session = Depends(get_db)):
    x = db.get(AcctSheet, sheet_id)
    if not x:
        raise HTTPException(status_code=404, detail="Not found")
    return Response(M.sheet_xlsx(x.name, json.loads(x.doc or "{}")), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{x.name}.xlsx"'})


# ---------------------------------------------------------------- reconcile against QuickBooks
@router.post("/reconcile/parse")
async def reconcile_parse(text: Optional[str] = Form(None), file: Optional[UploadFile] = File(None)):
    data = await file.read() if file else None
    if data is not None and len(data) > MAX_BYTES:
        raise HTTPException(status_code=400, detail="File too big")
    return M.parse_qb(data, file.filename if file else None, text)


class ReconcileIn(BaseModel):
    account_ids: List[int]
    start: Optional[str] = None
    end: Optional[str] = None
    rows: List[dict]
    mapping: Optional[dict] = None  # QuickBooks account name -> AT-HUB account id (saved for next time)
    tol_days: int = 5
    name: Optional[str] = None
    save: bool = True


@router.get("/reconcile/map")
def reconcile_map(db: Session = Depends(get_db)):
    return M.qb_map(db)


@router.post("/reconcile/run")
def reconcile_run(body: ReconcileIn, db: Session = Depends(get_db), user: User = Depends(get_current_active_user)):
    if not body.account_ids:
        raise HTTPException(status_code=400, detail="Pick the AT-HUB accounts to check")
    a, b = _period(db, body.start, body.end)
    mapping = M.qb_map(db)
    for k, v in (body.mapping or {}).items():
        if v:
            mapping[M._norm(k)] = int(v)
        else:
            mapping.pop(M._norm(k), None)
    M._save_setting(db, M.QB_MAP_KEY, mapping, _who(user))
    res = M.reconcile(db, body.account_ids, a, b, body.rows, mapping, max(0, min(body.tol_days, 30)))
    rid = None
    if body.save:
        r = AcctReconcile(name=(body.name or "").strip() or None, account_ids=json.dumps(body.account_ids), start=A._dt(a), end=A._dt(b),
                          qb_rows=json.dumps(body.rows), result=json.dumps({"counts": res["counts"], "totals": res["totals"]}), created_by=_who(user))
        db.add(r)
    db.commit()
    if body.save:
        rid = r.id
    return {**res, "id": rid, "start": a.isoformat(), "end": b.isoformat()}


@router.get("/reconcile/history")
def reconcile_history(db: Session = Depends(get_db)):
    accounts, _ = A.maps(db)
    out = []
    for r in db.query(AcctReconcile).order_by(AcctReconcile.id.desc()).limit(50).all():
        res = json.loads(r.result or "{}")
        out.append({"id": r.id, "name": r.name, "accounts": [A.path(accounts[i], accounts) for i in json.loads(r.account_ids or "[]") if i in accounts],
                    "account_ids": json.loads(r.account_ids or "[]"), "start": r.start.date().isoformat() if r.start else None,
                    "end": r.end.date().isoformat() if r.end else None, "counts": res.get("counts", {}), "lines": len(json.loads(r.qb_rows or "[]")),
                    "created_by": r.created_by, "created_at": r.created_at})
    return out


@router.get("/reconcile/{rec_id}")
def reconcile_get(rec_id: int, db: Session = Depends(get_db)):
    r = db.get(AcctReconcile, rec_id)
    if not r:
        raise HTTPException(status_code=404, detail="Not found")
    return {"id": r.id, "name": r.name, "account_ids": json.loads(r.account_ids or "[]"), "start": r.start.date().isoformat() if r.start else None,
            "end": r.end.date().isoformat() if r.end else None, "rows": json.loads(r.qb_rows or "[]")}


def _xlsx(sheets: list, filename: str) -> Response:
    """sheets: [(title, header, rows)] -> an Excel download."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    import io as _io
    wb = Workbook()
    wb.remove(wb.active)
    for title, header, rows in sheets:
        ws = wb.create_sheet(str(title)[:31])
        ws.append(header)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in rows:
            ws.append(r)
        ws.column_dimensions["A"].width = 26
        for col in "BCDEFGH":
            ws.column_dimensions[col].width = 16
    buf = _io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{filename}"'})


@router.get("/partner-accounts/export.xlsx")
def partner_accounts_export(start: Optional[str] = None, end: Optional[str] = None, db: Session = Depends(get_db)):
    a, b = _period(db, start, end)
    d = M.partner_accounts(db, a, b, _fy_key(db, a))
    head = ["Period", "Withdrawn", "Invested", "Net", "Lent To Company", "Repaid To Them", "Loan Owed To Them"]
    row = lambda k, t: [k, t["withdrawn"], t["invested"], t["net"], t["lent"], t["repaid"], t["loan_owed"]]
    sheets = [("All Partners", ["Partner"] + head[1:] + ["Share", "Remaining"],
               [[p["name"]] + row("", p["period"])[1:] + [p["share"], p["remaining"]] for p in d["partners"]])]
    for p in d["partners"]:
        sheets.append((p["name"], head, [row(x["fy"], x) for x in p["by_fy"]] + [[]] + [row(x["month"], x) for x in p["by_month"]]))
    return _xlsx(sheets, f"Partner Accounts {a.isoformat()} to {b.isoformat()}.xlsx")


@router.get("/check/export.xlsx")
def check_export(start: Optional[str] = None, end: Optional[str] = None, all_dates: bool = False, db: Session = Depends(get_db)):
    a, b = (None, None) if all_dates else _period(db, start, end)
    d = M.check_books(db, a, b)
    rows = [[f["label"], f["severity"], f["txn"]["date"] if f["txn"] else "", f["txn"]["amount"] if f["txn"] else "",
             (f["txn"]["description"] if f["txn"] else ""), (f["txn"]["account"] or "") if f["txn"] else "", f["why"],
             f["suggest"]["account"] if f["suggest"] else ""] for f in d["findings"]]
    return _xlsx([("Check My Books", ["Check", "How Sure", "Date", "Amount", "Description", "Account", "Why", "Suggested"], rows)], "Check My Books.xlsx")
